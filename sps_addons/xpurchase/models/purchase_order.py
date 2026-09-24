# -*- coding: utf-8 -*-


import base64
import os
import openpyxl
from openpyxl.writer.excel import save_virtual_workbook
from odoo import models, fields, api
from datetime import datetime, date
from odoo.exceptions import UserError
from odoo.tools import float_compare
from openpyxl.styles import NamedStyle, Font, Border, Side
from six import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment
from odoo.addons.num2currency import num999, num2word


class PurchaseOrder(models.Model):
    _inherit = 'purchase.order'

    @api.model
    def search(self, args, offset=0, limit=None, order=None, count=False):
        return super(PurchaseOrder, self.with_context(active_test=False)).search(
            args, offset=offset, limit=limit, order=order, count=count
        )

    date_planned = fields.Datetime(
        string='Receipt Date', index=True, copy=False, store=True, default=False, required=1, tracking=True,
        help="Delivery date promised by vendor. This date is used to determine expected arrival of products.")
    x_state = fields.Selection([
        ('draft', 'Nháp'),
        ('doing', 'Đang thực hiện'),
        ('out_date', 'Đã quá hạn'),
        ('delay', 'Tạm hoãn'),
        ('wait_license', 'Đã giao chờ chứng từ'),
        ('to_late', 'Đã giao hàng muộn'),
        ('done', 'Hoàn thành'),
        ('cancel', 'Hủy')
    ], string='Trạng thái', readonly=True, index=True, copy=False, default='draft', tracking=True)
    x_license = fields.Boolean('Đã có chứng từ')
    x_negotiator_id = fields.Many2one('res.users', 'Người đàm phán')
    x_invisible_button = fields.Selection([
        ('type_1', 'ẩn cả 2'),
        ('type_2', 'ẩn pk hiện iv'),
        ('type_3', 'ẩn iv hiện pk'),
        ('type_4', 'hiện cả 2'),
    ], compute='compute_type', default='type_4', string='tùy chọn')
    x_advanced_count = fields.Integer('Đếm tạm ứng', compute='compute_advanced')
    x_partner_contact = fields.Many2one('res.partner', string="Người liên hệ")
    x_partner_bank = fields.Many2one('res.partner.bank', string="Số tài khoản")
    x_wait_license_date = fields.Date('thời điểm giao đủ hàng')

    def compute_advanced(self):
        for order in self:
            advanceds = self.env['account.advance'].search([('po_id', '=', order.id)])
            order.x_advanced_count = len(advanceds)

    def action_view_advance(self):
        advanceds = self.env['account.advance'].search([('po_id', '=', self.id)])
        return {
            'name': ('Các phiếu tạm ứng'),
            'view_mode': 'tree,form',
            'domain': [('id', 'in', advanceds.ids)],
            'res_model': 'account.advance',
            'type': 'ir.actions.act_window',
        }

    def calculate_all_percent(self):
        purchase_line_ids = self.env['purchase.order.line'].search([('x_payment_ratio', '=', 0)])
        purchase_ids = purchase_line_ids.order_id
        purchase_ids.calculate_line_percent()

    def calculate_line_percent(self):
        for rec in self:
            if rec.amount_total > 0:
                advance_ids = self.env['account.advance'].search([('po_id', '=', rec.id), '|', ('state', '=', 'payment'), ('payment_id.state', '=', 'posted')])
                reconcile_ids = self.env['account.partial.reconcile'].search(
                    [('credit_move_id', 'in', rec.invoice_ids.filtered(lambda p: p.state != 'cancel').line_ids.ids)])
                move_payment_invoice_ids = reconcile_ids.debit_move_id.move_id
                advance_ids = advance_ids.filtered(lambda adv: adv.payment_id.move_id not in move_payment_invoice_ids)
                advance_percent = sum(advance_ids.payment_id.mapped('amount')) / rec.amount_total * 100
                purchase_line_percent = {
                    purchase_line.id: {
                        'purchase_line': purchase_line,
                        'percent': advance_percent,
                    }
                    for purchase_line in rec.order_line
                }
                for invoice_id in rec.invoice_ids.filtered(lambda p: p.state != 'cancel'):
                    if invoice_id.amount_total > 0:
                        reconcile_ids = self.env['account.partial.reconcile'].search(
                            [('credit_move_id', 'in', invoice_id.line_ids.ids)])
                        total_invoice_reconcile = sum(reconcile_ids.mapped('amount'))
                        invoice_percent = total_invoice_reconcile / invoice_id.amount_total * 100
                        for invoice_line in invoice_id.invoice_line_ids.filtered(lambda il: il.purchase_line_id):
                            if invoice_line.purchase_line_id.price_subtotal > 0:
                                purchase_line_invoice_percent = invoice_percent * invoice_line.price_subtotal / invoice_line.purchase_line_id.price_subtotal
                                item = purchase_line_percent.get(invoice_line.purchase_line_id.id)
                                if item:
                                    item['percent'] += purchase_line_invoice_percent
                for line in purchase_line_percent.values():
                    purchase_line = line['purchase_line']
                    x_payment_subtotal = (purchase_line.price_subtotal * line['percent'] / 100)
                    if x_payment_subtotal >= purchase_line.price_subtotal:
                        x_payment_ratio = 100
                    else:
                        x_payment_ratio = (x_payment_subtotal / purchase_line.price_subtotal * 100)
                    line['purchase_line'].write({'x_payment_ratio': x_payment_ratio,
                                                 'x_payment_subtotal': line['purchase_line'].price_subtotal * line['percent'] / 100})

    def compute_type(self):
        for r in self:
            if not r.order_line:
                r.x_invisible_button = 'type_4'
            elif any(line.qty_received < line.product_qty and line.qty_invoiced < line.product_qty for line in r.order_line):
                r.x_invisible_button = 'type_4'
            elif any(line.qty_received == line.product_qty and line.qty_invoiced < line.product_qty for line in r.order_line):
                r.x_invisible_button = 'type_2'
            elif any(line.qty_received < line.product_qty and line.qty_invoiced == line.product_qty for line in r.order_line):
                r.x_invisible_button = 'type_3'
            else:
                r.x_invisible_button = 'type_1'

    def update_lines_paid_amount(self, order_ids):
        self._cr.execute('''
            WITH ratio_data AS (
                SELECT
                    pol.id AS pol_id,

                    COALESCE(
                        (
                            -- Advance %
                            SELECT
                                SUM(ap.amount) / NULLIF(po.amount_total, 0) * 100
                            FROM account_advance aa
                            JOIN account_payment ap
                                ON ap.id = aa.payment_id
                            WHERE aa.po_id = po.id
                              AND (
                                  aa.state = 'payment'
                                  OR ap.state = 'posted'
                              )
                              AND NOT EXISTS (
                                  SELECT 1
                                  FROM account_partial_reconcile pr
                                  JOIN account_move_line debit_line
                                      ON debit_line.id = pr.debit_move_id
                                  WHERE debit_line.move_id = ap.move_id
                              )
                        ),
                        0
                    )
                    +
                    COALESCE(
                        (
                            -- Invoice payment %
                            SELECT SUM(
                                (
                                    COALESCE(
                                        (
                                            SELECT SUM(pr.amount)
                                            FROM account_partial_reconcile pr
                                            WHERE pr.credit_move_id IN (
                                                SELECT aml2.id
                                                FROM account_move_line aml2
                                                WHERE aml2.move_id = am.id
                                            )
                                        ),
                                        0
                                    )
                                    / NULLIF(am.amount_total, 0) * 100
                                )
                                * aml.price_subtotal
                                / NULLIF(pol.price_subtotal, 0)
                            )
                            FROM account_move_line aml
                            JOIN account_move am
                                ON am.id = aml.move_id
                            WHERE aml.purchase_line_id = pol.id
                              AND am.state != 'cancel'
                              AND am.amount_total > 0
                        ),
                        0
                    ) AS payment_ratio

                FROM purchase_order_line pol
                JOIN purchase_order po
                    ON po.id = pol.order_id
                WHERE po.id IN %s
            )

            UPDATE purchase_order_line AS tbl
            SET
                x_payment_subtotal = dt.paid_amount,
                x_payment_ratio = CASE
                    WHEN dt.payment_ratio >= 100 THEN 100
                    ELSE dt.payment_ratio
                END
            FROM (
                SELECT
                    pol.id,
                    SUM(
                        CASE
                            WHEN am.move_type = 'in_invoice'
                            THEN aml.price_subtotal
                            ELSE -aml.price_subtotal
                        END
                    ) AS paid_amount,
                    rd.payment_ratio
                FROM purchase_order po
                INNER JOIN purchase_order_line pol
                    ON po.id = pol.order_id
                LEFT JOIN account_move_line aml
                    ON pol.id = aml.purchase_line_id
                LEFT JOIN account_move am
                    ON aml.move_id = am.id
                INNER JOIN ratio_data rd
                    ON rd.pol_id = pol.id
                WHERE true
                    AND po.id IN %s
                    AND am.payment_state IN ('in_payment', 'paid')
                GROUP BY pol.id, rd.payment_ratio
            ) dt (pol_id, paid_amount, payment_ratio)
            WHERE tbl.id = dt.pol_id
        ''', [
            tuple(order_ids + [0, 0]),
            tuple(order_ids + [0, 0]),
        ])

    def update_lines_paid_amount2(self, order_ids):
        self._cr.execute('''
            WITH ratio_data AS (
                SELECT
                    pol.id AS pol_id,

                    COALESCE(
                        (
                            -- Advance %
                            SELECT
                                SUM(ap.amount) / NULLIF(po.amount_total, 0) * 100
                            FROM account_advance aa
                            JOIN account_payment ap
                                ON ap.id = aa.payment_id
                            WHERE aa.po_id = po.id
                              AND (
                                  aa.state = 'payment'
                                  OR ap.state = 'posted'
                              )
                              AND NOT EXISTS (
                                  SELECT 1
                                  FROM account_partial_reconcile pr
                                  JOIN account_move_line debit_line
                                      ON debit_line.id = pr.debit_move_id
                                  WHERE debit_line.move_id = ap.move_id
                              )
                        ),
                        0
                    )
                    +
                    COALESCE(
                        (
                            -- Invoice payment %
                            SELECT SUM(
                                (
                                    COALESCE(
                                        (
                                            SELECT SUM(pr.amount)
                                            FROM account_partial_reconcile pr
                                            WHERE pr.credit_move_id IN (
                                                SELECT aml2.id
                                                FROM account_move_line aml2
                                                WHERE aml2.move_id = am.id
                                            )
                                        ),
                                        0
                                    )
                                    / NULLIF(am.amount_total, 0) * 100
                                )
                                * aml.price_subtotal
                                / NULLIF(pol.price_subtotal, 0)
                            )
                            FROM account_move_line aml
                            JOIN account_move am
                                ON am.id = aml.move_id
                            WHERE aml.purchase_line_id = pol.id
                              AND am.state != 'cancel'
                              AND am.amount_total > 0
                        ),
                        0
                    ) AS payment_ratio

                FROM purchase_order_line pol
                JOIN purchase_order po
                    ON po.id = pol.order_id
                WHERE po.id IN %s
            )

            UPDATE purchase_order_line AS tbl
            SET
                x_payment_subtotal = dt.paid_amount,
                x_payment_ratio = CASE
                    WHEN dt.payment_ratio >= 100 THEN 100
                    ELSE dt.payment_ratio
                END
            FROM (
                SELECT
                    pol.id,
                    SUM(
                        CASE
                            WHEN am.payment_state IN ('in_payment', 'paid')
                            THEN aml.price_subtotal
                            ELSE 0
                        END
                    ) AS paid_amount,
                    rd.payment_ratio
                FROM purchase_order po
                INNER JOIN purchase_order_line pol
                    ON po.id = pol.order_id
                LEFT JOIN account_move_line aml
                    ON pol.id = aml.purchase_line_id
                LEFT JOIN account_move am
                    ON aml.move_id = am.id
                INNER JOIN ratio_data rd
                    ON rd.pol_id = pol.id
                WHERE true
                    AND po.id IN %s
                    AND am.payment_state NOT IN ('in_payment', 'paid')
                GROUP BY pol.id, rd.payment_ratio
            ) dt (pol_id, paid_amount, payment_ratio)
            WHERE tbl.id = dt.pol_id
        ''', [
            tuple(order_ids + [0, 0]),
            tuple(order_ids + [0, 0]),
        ])

    @api.depends('partner_id')
    @api.onchange('partner_id')
    def onchange_partner(self):
        for r in self:
            if r.partner_id:
                contact = self.env['res.partner'].sudo().search([('parent_id', '=', r.partner_id.id)], limit=1).id
                bank = self.env['res.partner.bank'].sudo().search([('partner_id', '=', r.partner_id.id)], limit=1).id
                r.x_partner_contact = contact
                r.x_partner_bank = bank

    @api.model
    def default_get(self, fields):
        res = super(PurchaseOrder, self).default_get(fields)
        sql = '''
                SELECT
                    pal.ID,
                    pt.name as product_name,
                    pal.product_id,
                    pal.x_cost_classification_id,
                    pal.x_default_code,
                    pal.x_brand,
                    pal.x_origin,
                    pal.x_source,
                    pal.product_uom_id,
                    pal.product_qty,
                    pa.ID AS pa_id,
                    pa.x_purpose,
                    pa.user_id,
                    pa.x_code_project_id,
                    pa.x_name_project,
                    pa.ordering_date,
                    pa.schedule_date,
                    pa.x_cost_other,
                    pal.price_unit,
                    pal.product_qty - pal.qty_ordered AS qty_suggest 
                FROM
                    purchase_requisition_line pal
                    LEFT JOIN purchase_requisition pa ON pal.requisition_id = pa.ID 
                    LEFT JOIN product_product pp ON pal.product_id = pp.ID
                    LEFT JOIN product_template pt ON pp.product_tmpl_id = pt.ID 
                WHERE
                    pa.x_state IN ( 'going', 'official' ) 
                    AND pal.product_qty > COALESCE ( pal.qty_done, 0 ) 
                    AND pal.product_qty > COALESCE (pal.qty_ordered,0) 
        '''
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()
        order_line = []
        for rec in recs:
            order_line.append((0, 0, {
                'name': rec['product_name'],
                'product_id': rec['product_id'],
                'x_cost_type_id': rec['x_cost_classification_id'],
                'x_default_code': rec['x_default_code'],
                'x_brand': rec['x_brand'],
                'x_origin': rec['x_origin'],
                'product_uom': rec['product_uom_id'],
                'x_qty_requisition': rec['product_qty'],
                'x_requisition_id': rec['pa_id'],
                'x_users_id': rec['user_id'],
                'x_purpose': rec['x_purpose'],
                'x_project_id': rec['x_code_project_id'],
                'x_date_order': rec['ordering_date'],
                'x_schedule_date': rec['schedule_date'],
                'date_planned': rec['schedule_date'],
                'x_source': rec['x_source'],
                'x_budget_price': rec['price_unit'],
                'product_qty': rec['qty_suggest'],
                'price_unit': rec['price_unit'],
                'x_pa_lines': [(6, 0, [rec['id']])],
            }))
        res['order_line'] = order_line
        res['notes'] = "1. Quy cách : \n2. Bảo hành: \n3. Người nhận hàng: \n4. Thời gian nhận hàng: \n5. Địa điểm nhận hàng: \n6. Người đặt hàng: \n7. Lịch thanh toán: \n8. Ghi chú:"
        return res

    def action_delay(self):

        for r in self:
            if r.x_state == 'delay':
                continue
            else:
                r.x_state = 'delay'

    def action_continute(self):
        for r in self:
            if r.x_state != 'delay':
                continue
            else:
                time_now = datetime.now()
                if r.date_planned <= time_now:
                    self.x_state = 'doing'
                else:
                    self.x_state = 'out_date'

    def action_create_picking(self, force=False):
        result = super(PurchaseOrder, self).button_approve(force=force)
        return result

    def _check_requisition_qty(self):
        """Kiểm tra tổng số lượng đặt mua cho mỗi dòng yêu cầu không vượt quá số lượng
        cần mua của đề xuất. Bắt trường hợp số lượng cần mua đã được điều chỉnh giảm sau
        khi đơn mua được tạo. Loại trừ đơn mua ở trạng thái nháp và huỷ."""
        self.ensure_one()
        # Gom số lượng đặt mua trong chính đơn này theo từng dòng yêu cầu mua
        qty_by_line = {}
        for ol in self.order_line:
            for rl in ol.x_pa_lines:
                qty_by_line.setdefault(rl, 0.0)
                qty_by_line[rl] += ol.product_qty
        for rl, this_qty in qty_by_line.items():
            rounding = rl.product_uom_id.rounding or 0.01
            # Số lượng đã đặt trên các đơn khác đã xác nhận (loại trừ chính đơn này, nháp & huỷ)
            other_qty = sum(
                pol.product_qty for pol in rl.x_po_lines
                if pol.order_id.id != self.id and pol.order_id.x_state not in ('draft', 'cancel')
            )
            if float_compare(other_qty + this_qty, rl.product_qty, precision_rounding=rounding) > 0:
                raise UserError(
                    'Sản phẩm %s (đề xuất %s): tổng số lượng đặt mua (%s) vượt quá số lượng cần mua '
                    'của đề xuất (%s).\nSố lượng cần mua có thể đã được điều chỉnh giảm. '
                    'Vui lòng sửa lại số lượng trên đơn mua trước khi xác nhận.'
                    % (rl.product_id.display_name or '', rl.requisition_id.name or '',
                       other_qty + this_qty, rl.product_qty)
                )

    def action_confirm(self):
        for r in self:
            if r.x_state != 'draft':
                continue
            else:
                r._check_requisition_qty()
                pa2update = []
                r.x_state = 'doing'
                r.state = 'purchase'
                r._update_requisition_qty_ordered()
                for rec in r.order_line:
                    if not rec.x_project_id.x_order_id:
                        raise UserError("Dự án %s chưa có mã dự án chính" % rec.x_project_id.name)
                    if rec.x_requisition_id.id:
                        pa2update.append(rec.x_requisition_id.id)
                if pa2update:
                    self.env['purchase.requisition'].check_x_state_1(pa2update)

    def action_advanced(self):
        line_ids = []
        amount_advanced = 0
        project_totals = {}
        # Tổng hợp dữ liệu theo x_project_id
        for ol in self.order_line:
            project_id = ol.x_project_id.id
            if project_id not in project_totals:
                project_totals[project_id] = ol.price_total
            else:
                project_totals[project_id] += ol.price_total
            amount_advanced += ol.price_total

        # Chuyển dữ liệu thành danh sách các dictionary cho One2many
        for project_id, advance_detail in project_totals.items():
            line_ids.append((0, 0, {
                'project_id': project_id,
                # 'advance_detail': advance_detail,
                'total_cost': advance_detail,
                'percentage': advance_detail/amount_advanced * 100 if amount_advanced else 0,
            }))

        # Lấy số tiền cần tạm ứng còn lại
        total_advanced = 0
        account_advance_ids = self.env['account.advance'].search(
            [('po_id', '=', self.id), ('state', '!=', 'cancel')])
        # Tính tổng tất cả các phiếu tạm ứng
        if account_advance_ids:
            total_advanced += sum(advance.amount_total for advance in account_advance_ids)
        remain_advanced = self.amount_total - total_advanced

        return {
            'name': 'Tạm ứng cho đơn mua',
            'res_model': 'account.advance.register',
            'view_mode': 'form',
            'context': {
                'po_id': self.id,
                'default_amount_advanced': remain_advanced,
                'default_line_ids': line_ids
            },
            'target': 'new',
            'type': 'ir.actions.act_window',
        }

    @api.model
    def create(self, vals):
        res = super(PurchaseOrder, self).create(vals)
        return res

    def _update_requisition_qty_ordered(self):
        """Tính lại `qty_ordered` trên các dòng yêu cầu mua liên kết với đơn này, dựa trên
        trạng thái hiện tại (x_state) của đơn - gọi sau khi x_state đã được cập nhật."""
        self.order_line.mapped('x_pa_lines').update_qty()

    def button_cancel(self):
        advanced_ids = self.env['account.advance'].search([('po_id', '=', self.id)])
        if any(x.state not in ('draft', 'cancel') for x in advanced_ids):
            raise UserError('Không thể hủy đơn mua này')
        self.x_state = 'cancel'
        self.state = 'cancel'
        self._update_requisition_qty_ordered()
        return super(PurchaseOrder, self).button_cancel()

    @api.constrains('x_license')
    def check_license(self):
        if self.x_state != 'wait_license' or self.x_license == False:
            return
        else:
            if not self.x_wait_license_date:
                return
            if any(x.x_date_received and x.x_schedule_date and x.x_date_received > x.x_schedule_date for x in self.order_line):
                self.update({'x_state' :'to_late', "state" :'done'})
            else:
                 self.update({'x_state': 'done', "state": 'done'})

    # tự động chuyển trạng thái đã quá hạn
    def cron_state_out_date(self):
        sql = '''
        update purchase_order
        set x_state='out_date'
        where id in (
            select distinct po.id
            from purchase_order_line pol
                inner join purchase_order po on po.id=pol.order_id
            where true
                and po.x_state = 'doing'
                and po.date_planned ::date < %s
                and pol.qty_received < pol.product_qty
        )
        '''
        self._cr.execute(sql, [date.today()])

    def state_to_late(self, picking_ids):

        if not picking_ids and len(self) == 1:
            sql = '''
            select sm.picking_id
            from purchase_order po
                inner join purchase_order_line pol on po.id=pol.order_id
                left join stock_move sm on pol.id=sm.purchase_line_id
            where true
                and sm.state='done'
                and po.id=%s
            order by sm.id desc
            limit 1
            '''
            self._cr.execute(sql, [self.id])
            recs = self._cr.fetchone()
            if recs and recs[0]:
                picking_ids = [recs[0]]
        if not picking_ids:
            return
        if isinstance(picking_ids, int):
            picking_ids = [picking_ids]
        sql_check = '''
        select dt.po_id
        from (
            select po.id po_id, pol.id pol_id, sum(sml.qty_done) qty_done, pol.product_qty as qty_request
            from purchase_order po
                inner join purchase_order_line pol on po.id=pol.order_id
                inner join stock_move sm on pol.id=sm.purchase_line_id
                inner join stock_picking sp on sm.picking_id=sp.id
                INNER JOIN stock_move_line sml on sml.move_id = sm.id
            where true
                and po.id in (
                    select distinct po.id 
                    from purchase_order po 
                    inner join purchase_order_line pol on po.id=pol.order_id
                    inner join stock_move sm on pol.id=sm.purchase_line_id
                    inner join stock_picking sp on sm.picking_id=sp.id
                    where sp.id in %s
                    group by po.id
                )
            group by po.id, pol.id
        ) dt where dt.qty_done < dt.qty_request'''
        self._cr.execute(sql_check, [tuple(picking_ids + [0, 0])])
        res = self._cr.dictfetchall()
        po_ids = [r['po_id'] for r in res]
        sql = '''update purchase_order
        set x_state='to_late', "state" ='done'
        where id in (
            select distinct po.id
            from purchase_order po
                inner join purchase_order_line pol on po.id=pol.order_id
                inner join stock_move sm on pol.id=sm.purchase_line_id
                inner join stock_picking sp on sm.picking_id=sp.id
            where true
                and po.id not in %s
                and po.x_state in ('doing','out_date','wait_license')
                and po.x_license = 't'
                and po.date_planned :: date < sp.date_done :: date
                and sp.id in %s
                )'''
        self._cr.execute(sql, [
            tuple(po_ids + [0, 0]),
            tuple(picking_ids + [0, 0])
        ])
        sql_re_state = '''select po.id
                          from purchase_order po
                           where id in (
                                select distinct po.id
                                from purchase_order po
                                    inner join purchase_order_line pol on po.id=pol.order_id
                                    inner join stock_move sm on pol.id=sm.purchase_line_id
                                    inner join stock_picking sp on sm.picking_id=sp.id
                                where true
                                    and po.id not in %s
                                    and po.x_state in ('to_late')
                                    and po.state in ('done')
                                    and po.x_license = 't'
                                    and po.date_planned :: date < sp.date_done :: date
                                    and sp.id in %s
                                    )'''
        self._cr.execute(sql_re_state, [
            tuple(po_ids + [0, 0]),
            tuple(picking_ids + [0, 0])
        ])
        res1 = self._cr.dictfetchall()
        for r in res1:
            po = self.env['purchase.order'].browse(r['id'])
            if any(line.qty_received < line.product_qty for line in po.order_line):
                self._cr.execute(f'''update purchase_order
                       set x_state='doing', "state" ='purchase'
                       where id ={po.id}''')
            else:
                po.update({'x_wait_license_date': fields.date.today()})

    def state_wait_license(self, picking_ids):
        if isinstance(picking_ids, int):
            picking_ids = [picking_ids]
        sql_check = '''
                select dt.po_id
                from (
                    select po.id po_id, pol.id pol_id, sum(sml.qty_done) qty_done, pol.product_qty as qty_request
                    from purchase_order po
                        inner join purchase_order_line pol on po.id=pol.order_id
                        inner join stock_move sm on pol.id=sm.purchase_line_id
                        inner join stock_picking sp on sm.picking_id=sp.id
                        INNER JOIN stock_move_line sml on sml.move_id = sm.id
                    where true
                        and po.id in (
                            select distinct po.id 
                            from purchase_order po 
                            inner join purchase_order_line pol on po.id=pol.order_id
                            inner join stock_move sm on pol.id=sm.purchase_line_id
                            inner join stock_picking sp on sm.picking_id=sp.id
                            where sp.id in %s
                            group by po.id
                        )
                    group by po.id, pol.id
                ) dt where dt.qty_done < dt.qty_request'''
        self._cr.execute(sql_check, [tuple(picking_ids + [0, 0])])
        res = self._cr.dictfetchall()
        po_ids = [r['po_id'] for r in res]
        sql = '''update purchase_order
                set x_state='wait_license'
                where id in (
                    select distinct po.id
                    from purchase_order po
                        inner join purchase_order_line pol on po.id=pol.order_id
                        inner join stock_move sm on pol.id=sm.purchase_line_id
                        inner join stock_picking sp on sm.picking_id=sp.id
                    where true
                        and po.id not in %s
                        and po.x_state in ('doing','out_date')
                        and po.x_license = 'f'
                        and sp.id in %s
                        )'''
        self._cr.execute(sql, [
            tuple(po_ids + [0, 0]),
            tuple(picking_ids + [0, 0])
        ])
        sql_re_state = '''select po.id
                       from purchase_order po
                       where id in (
                           select distinct po.id
                           from purchase_order po
                               inner join purchase_order_line pol on po.id=pol.order_id
                               inner join stock_move sm on pol.id=sm.purchase_line_id
                               inner join stock_picking sp on sm.picking_id=sp.id
                           where true
                               and po.id not in %s
                               and po.x_state in ('wait_license')
                               and po.x_license = 'f'
                               and sp.id in %s
                               )'''
        self._cr.execute(sql_re_state, [
            tuple(po_ids + [0, 0]),
            tuple(picking_ids + [0, 0])
        ])
        res1 = self._cr.dictfetchall()
        for r in res1 :
            po =self.env['purchase.order'].browse(r['id'])
            if any(line.qty_received < line.product_qty for line in po.order_line):
                self._cr.execute(f'''update purchase_order
                set x_state='doing'
                where id ={po.id}''')
            else:
                po.update({'x_wait_license_date': fields.date.today()})

    def state_done(self, picking_ids):
        if not picking_ids and len(self) == 1:
            sql = '''
                  select sm.picking_id
                  from purchase_order po
                      inner join purchase_order_line pol on po.id=pol.order_id
                      left join stock_move sm on pol.id=sm.purchase_line_id
                  where true
                      and sm.state='done'
                      and po.id=%s
                  order by sm.id desc
                  limit 1
                  '''
            self._cr.execute(sql, [self.id])
            recs = self._cr.fetchone()
            if recs and recs[0]:
                picking_ids = [recs[0]]
        if not picking_ids:
            return
        if isinstance(picking_ids, int):
            picking_ids = [picking_ids]
        sql_check = '''
                select dt.po_id
                from (
                    select po.id po_id, pol.id pol_id, sum(sml.qty_done) qty_done, pol.product_qty as qty_request
                    from purchase_order po
                        inner join purchase_order_line pol on po.id=pol.order_id
                        inner join stock_move sm on pol.id=sm.purchase_line_id
                        inner join stock_picking sp on sm.picking_id=sp.id
                        INNER JOIN stock_move_line sml on sml.move_id = sm.id
                    where true
                        and po.id in (
                            select distinct po.id 
                            from purchase_order po 
                            inner join purchase_order_line pol on po.id=pol.order_id
                            inner join stock_move sm on pol.id=sm.purchase_line_id
                            inner join stock_picking sp on sm.picking_id=sp.id
                            where sp.id in %s
                            group by po.id
                        )
                    group by po.id, pol.id
                ) dt where dt.qty_done < dt.qty_request'''
        self._cr.execute(sql_check, [tuple(picking_ids + [0, 0])])
        res = self._cr.dictfetchall()
        po_ids = [r['po_id'] for r in res]
        sql = '''update purchase_order
               set x_state='done', "state" ='done'
               where id in (
                   select distinct po.id
                   from purchase_order po
                       inner join purchase_order_line pol on po.id=pol.order_id
                       inner join stock_move sm on pol.id=sm.purchase_line_id
                       inner join stock_picking sp on sm.picking_id=sp.id
                   where true
                       and po.id not in %s
                       and po.x_state in ('doing','out_date','wait_license')
                       and po.x_license = 't'
                       and po.date_planned :: date >= sp.date_done :: date
                       and sp.id in %s
                       )'''
        self._cr.execute(sql, [
            tuple(po_ids + [0, 0]),
            tuple(picking_ids + [0, 0])
        ])
        sql_re_state = '''select po.id
                                  from purchase_order po
                                   where id in (
                                        select distinct po.id
                                        from purchase_order po
                                            inner join purchase_order_line pol on po.id=pol.order_id
                                            inner join stock_move sm on pol.id=sm.purchase_line_id
                                            inner join stock_picking sp on sm.picking_id=sp.id
                                        where true
                                            and po.id not in %s
                                            and po.x_state in ('done')
                                            and po.state in ('done')
                                            and po.x_license = 't'
                                            and po.date_planned :: date >= sp.date_done :: date
                                            and sp.id in %s
                                            )'''
        self._cr.execute(sql_re_state, [
            tuple(po_ids + [0, 0]),
            tuple(picking_ids + [0, 0])
        ])
        res1 = self._cr.dictfetchall()
        for r in res1:
            po = self.env['purchase.order'].browse(r['id'])
            if any(line.qty_received < line.product_qty for line in po.order_line):
                self._cr.execute(f'''update purchase_order
                               set x_state='doing', "state" ='purchase'
                               where id ={po.id}''')
            else:
                po.update({'x_wait_license_date': fields.date.today()})

    def get_job_description(self):
        self.ensure_one()
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%sprint_purchase.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']
        sql = f'''
                    SELECT
                        rc.company_registry as ten_cty,
                        rc.phone as dien_thoai,
                        rp.website as web,
                        po.name as so,
                        po.date_order as ngay,
                        rc.name as ten_cty_1,
                        rp.name as ten_cty_ban,
                        rp.vat as ms_thue,
                        rp.street as dia_chi_ban,
                        rpb.acc_number as tai_khoan,
                        rp.name as nguoi_lien_he,
                        concat(rp.phone, '  ', rp.email) as dt_email,
                        pt.name as noi_dung,
                        pol.x_brand as nhan_hieu,
                        pol.x_default_code as ma_hang,
                        case when uu.name = 'Units' then 'Cái' 
                            when uu.name = 'Hours' then 'Giờ'
                            when uu.name = 'Set' then 'Bộ' end as don_vi,
                        pol.product_qty as so_luong,
                        pol.price_unit as don_gia,
                        pol.price_total as thanh_tien_pol,
                        po.amount_untaxed as tong,
                        po.amount_tax as vat,
                        po.amount_total as thanh_tien_po,
                        po.notes as ghi_chu
                    FROM purchase_order po
                    INNER JOIN purchase_order_line pol on po.id=pol.order_id
                    left join product_product pp on pol.product_id = pp.id
                    left join product_template pt on pp.product_tmpl_id = pt.id
                    left join uom_uom uu on pol.product_uom = uu.id 
                    LEFT JOIN res_company rc on po.company_id = rc.id
                    left join res_partner rp on po.partner_id = rp.id 
                    LEFT JOIN res_partner_bank rpb on rp.id = rpb.partner_id
                    WHERE po.id = {self.id}
        '''
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        highlight = NamedStyle(name="highlight")
        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dotted', color="000000")
        bd3 = Side(style='none')
        highlight.border = Border(left=bd1, top=bd2, right=bd1, bottom=bd2)

        style_sum1 = NamedStyle(name="style_sum1")
        style_sum1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        style_sum2 = NamedStyle(name="style_sum2")
        style_sum2.font = Font(bold=True)

        style_sum3 = NamedStyle(name="style_sum3")
        style_sum3.number_format = '#,##0'
        style_sum3.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        style_sum4 = NamedStyle(name="style_sum4")
        style_sum4.number_format = '#,##0'

        style_sum5 = NamedStyle(name="style_sum5")
        style_sum5.number_format = '#,##0'
        style_sum5.alignment = Alignment(wrap_text=True)
        style_sum5.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        for r in recs:
            ws.cell(1, 4).value = r['ten_cty']
            ws.cell(2, 4).value = self.env.company.street
            ws.cell(4, 4).value = r['dien_thoai']
            ws.cell(5, 4).value = self.env.company.website
            ws.cell(8, 2).value = r['so']
            ws.cell(8, 7).value = r['ngay'].strftime('%d-%m-%Y')
            ws.cell(10, 3).value = r['ten_cty_1']
            ws.cell(10, 8).value = r['ten_cty_ban']
            ws.cell(11, 3).value = self.env.company.street
            ws.cell(12, 3).value = r['ms_thue']
            ws.cell(11, 8).value = r['dia_chi_ban']
            ws.cell(12, 8).value = r['tai_khoan']
            ws.cell(13, 8).value = r['nguoi_lien_he']
            ws.cell(14, 8).value = r['dt_email']
            break
        row = 18
        index = 1
        for r in recs:
            ws.cell(row, 1).value = index
            ws.cell(row, 2).value, ws.cell(row, 2).style = r['noi_dung'], style_sum3
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            ws.cell(row, 4).value, ws.cell(row, 4).style = r['nhan_hieu'], style_sum3
            ws.cell(row, 5).value, ws.cell(row, 5).style = r['ma_hang'], style_sum3
            ws.cell(row, 6).value, ws.cell(row, 6).style = r['don_vi'], style_sum3
            ws.cell(row, 7).value, ws.cell(row, 7).style = r['so_luong'], style_sum3
            ws.cell(row, 8).value, ws.cell(row, 8).style = r['don_gia'], style_sum3
            ws.cell(row, 9).value, ws.cell(row, 9).style = r['thanh_tien_pol'], style_sum3
            row += 1
            index += 1

        ws.cell(row, 9).value, ws.cell(row, 9).style = r['tong'], style_sum4
        ws.cell(row, 9).value = r['tong']
        ws.cell(row, 5).value, ws.cell(row, 5).style = 'TỔNG (VNĐ)', style_sum2
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
        ws.cell(row + 1, 9).value, ws.cell(row + 1, 9).style = r['vat'], style_sum4
        ws.cell(row + 1, 9).value = r['vat']
        ws.cell(row + 1, 5).value, ws.cell(row + 1, 5).style = 'VAT', style_sum2
        ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=4)
        ws.merge_cells(start_row=row + 1, start_column=6, end_row=row + 1, end_column=8)
        ws.cell(row + 2, 9).value, ws.cell(row + 2, 9).style = r['thanh_tien_po'], style_sum4
        ws.cell(row + 2, 9).value = r['thanh_tien_po']
        ws.cell(row + 2, 5).value, ws.cell(row + 2, 5).style = 'THÀNH TIỀN (VNĐ', style_sum2
        ws.merge_cells(start_row=row + 2, start_column=2, end_row=row + 2, end_column=4)
        ws.merge_cells(start_row=row + 2, start_column=5, end_row=row + 2, end_column=8)
        ws.cell(row + 4, 2).value, ws.cell(row + 4, 2).style = r['ghi_chu'], style_sum5
        ws.cell(row + 3, 2).value, ws.cell(row + 3, 2).style = 'Ghi chú', style_sum2
        ws.merge_cells(start_row=row + 4, start_column=2, end_row=row + 11, end_column=9)
        ws.cell(row + 13, 4).value, ws.cell(row + 13, 4).style = 'BÊN MUA', style_sum2
        ws.cell(row + 13, 8).value, ws.cell(row + 13, 8).style = 'BÊN BÁN', style_sum2

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Mẫu in đơn đặt hàng.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        return attachment_id

    @api.model
    def fields_get(self, fields=None):
        hide = ['state']
        res = super(PurchaseOrder, self).fields_get()
        for field in hide:
            res[field]['searchable'] = False
        return res

    def write(self, vals):
        res = super(PurchaseOrder, self).write(vals)
        today = date.today()
        for r in self:
            if not any(line.qty_received < line.product_qty for line in r.order_line) and r.x_state in ('doing','out_date'):
                #to_late
                if r.x_state in ('doing','out_date','wait_license') and r.x_license and r.date_planned.date() < today:
                    r.write({'state': 'done', 'x_state': 'to_late','x_wait_license_date' : today})
                # wait_license
                if r.x_state in ('doing','out_date') and not r.x_license :
                    r.write({'x_state': 'wait_license','x_wait_license_date' : today})
                #done
                if r.x_state in ('doing','out_date','wait_license')and r.x_license and r.date_planned.date() >= today:
                    r.write({'state': 'done', 'x_state': 'done','x_wait_license_date' : today})
        return res

    def button_draft(self):
        res = super(PurchaseOrder, self).button_draft()
        advanced_ids = self.env['account.advance'].search([('po_id', '=', self.id)])
        if any(x.state not in ('draft', 'cancel') for x in advanced_ids) and not self.env.context.get('skip_validate'):
            raise UserError('Không thể đưa đơn mua này về trạng thái nháp')
        if (self.picking_ids or self.invoice_ids) and not self.env.context.get('skip_validate'):
            raise UserError('Không thể đưa đơn mua này về trạng thái nháp')
        else:
            advanced_ids.action_to_draft()
            advanced_ids.unlink()
            self.write({'x_state': 'draft'})
            self._update_requisition_qty_ordered()
            return res

    def button_draft_skip_check(self):

        return self.with_context(skip_validate=True).button_draft()


    def copy(self, default=None):
        self.ensure_one()
        default = dict(default or {})
        default['date_planned'] = self.date_planned
        res = super(PurchaseOrder, self).copy(default=default)
        return res

    # Khi cập nhật lại ngày giao hàng dự kiến cập nhật lại trạng thái đơn hàng
    @api.constrains('date_planned', 'order_line.x_schedule_date')
    def constrains_x_state(self):
        for rec in self:
            if rec.x_state == 'out_date' and all(x.x_schedule_date == False or rec.date_planned.date() <= x.x_schedule_date for x in rec.order_line):
                rec.x_state = 'doing'


class PurchaseOrderLine(models.Model):
    _inherit = 'purchase.order.line'

    x_default_code = fields.Char('Mã hiệu')
    x_cost_type_id = fields.Many2one('cost.type', 'Phân loại chi phí')
    x_brand = fields.Char('Thương hiệu')
    x_origin = fields.Char('Xuất xứ')
    x_note = fields.Char('Ghi chú')
    x_requisition_id = fields.Many2one('purchase.requisition', string='Số yêu cầu', copy=False, readonly=True)
    x_users_id = fields.Many2one('res.users', 'Người yêu cầu')
    x_purpose = fields.Selection(
        [('project', 'Dự án'), ('other', 'Mục đích khác'), ('replenish_inventory', 'Bổ sung tồn kho')], 'Mục đích')
    x_project_id = fields.Many2one('project.project', 'Mã dự án',required=1)
    x_project_name = fields.Char('Tên dự án', compute="_compute_x_project_name", store=1)
    x_date_order = fields.Date('Ngày yêu cầu')
    x_schedule_date = fields.Date('Ngày cần hàng')
    x_source = fields.Char('Nguồn cấp tham khảo')
    x_qty_requisition = fields.Integer('Số lượng yêu cầu')
    x_qty_done = fields.Integer('Số lượng đã nhận')
    x_qty_debt = fields.Integer('Số lượng đã lên công nợ')
    x_budget_price = fields.Float('ĐG ngân sách')
    x_payment_subtotal = fields.Float('Đã thanh toán')
    x_payment_ratio = fields.Float(string='Tỷ lệ thanh toán')
    x_pa_lines = fields.Many2many('purchase.requisition.line', 'pa_line_po_line_ref', 'po_line_id', 'pa_line_id',
                                  'Dòng yêu cầu mua',copy=False)
    x_date_received = fields.Date('Ngày hàng về')
    x_invoice_subtotal = fields.Float(
        'Giá trị theo hóa đơn', compute='_compute_x_invoice_subtotal',
        help='Tổng giá trị (trước thuế) các dòng HÓA ĐƠN mua hàng (vendor bill) '
             'đã posted gắn với dòng PO này. Có thể khác giá đặt hàng '
             '(price_subtotal) khi hóa đơn thực tế lệch so với đơn mua.')


    # @api.depends('x_payment_subtotal', 'price_subtotal')
    # def _get_payment_ratio(self):
    #     for r in self:
    #         if r.price_subtotal <= r.x_payment_subtotal:
    #             r.x_payment_ratio = 100
    #         else:
    #             r.x_payment_ratio = r.x_payment_subtotal / r.price_subtotal


    @api.depends('invoice_lines.price_subtotal', 'invoice_lines.move_id.state',
                 'invoice_lines.move_id.move_type')
    def _compute_x_invoice_subtotal(self):
        for rec in self:
            inv_lines = rec.invoice_lines.filtered(
                lambda l: l.move_id.move_type == 'in_invoice'
                and l.move_id.state == 'posted')
            rec.x_invoice_subtotal = sum(inv_lines.mapped('price_subtotal'))

    @api.depends('x_project_id', 'x_project_id.label_tasks')
    def _compute_x_project_name(self):
        for rec in self:
            rec.x_project_name = rec.x_project_id.label_tasks

    # @api.constrains('price_unit')
    # def check_price(self):
    #     for line in self:
    #         if line.price_unit == 0:
    #             raise UserError('Đơn giá thực tế của sản phẩm phải lớn hơn 0')
    #         else:
    #             continue

    @api.constrains('qty_received')
    def update_pr_line_qty_done(self):
        for r in self:
            today = date.today()
            r.x_date_received = today
            if not r.x_pa_lines:
                continue
            r.x_pa_lines[0].qty_done = r.qty_received

    def _prepare_account_move_line(self, move=False):
        self.ensure_one()
        aml_currency = move and move.currency_id or self.currency_id
        date = move and move.date or fields.Date.today()
        res = {
            'display_type': self.display_type,
            'sequence': self.sequence,
            'name': '%s: %s' % (self.order_id.name, self.name),
            'product_id': self.product_id.id,
            'product_uom_id': self.product_uom.id,
            'quantity': self.qty_to_invoice,
            'price_unit': self.currency_id._convert(self.price_unit, aml_currency, self.company_id, date, round=False),
            'tax_ids': [(6, 0, self.taxes_id.ids)],
            'analytic_account_id': self.account_analytic_id.id,
            'analytic_tag_ids': [(6, 0, self.analytic_tag_ids.ids)],
            'purchase_line_id': self.id,
            'x_cost_type_id': self.x_cost_type_id or False,
            'x_project_id': self.x_project_id.id
        }
        if not move:
            return res

        if self.currency_id == move.company_id.currency_id:
            currency = False
        else:
            currency = move.currency_id

        res.update({
            'move_id': move.id,
            'currency_id': currency and currency.id or False,
            'date_maturity': move.invoice_date_due,
            'partner_id': move.partner_id.id,
        })
        return res

    def _update_name_project(self):
        po_ids = self.search([])
        for po_id in po_ids:
            po_id._compute_x_project_name()
