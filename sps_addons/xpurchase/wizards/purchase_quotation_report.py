# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side
from openpyxl.writer.excel import save_virtual_workbook
from odoo import api, fields, models


class ReportPurchaseQuotation(models.TransientModel):
    _name = 'purchase.quotation.report'
    _description = 'Báo cáo báo giá mua hàng'

    def purchase_report_data(self):
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(
            dir_path + '%s..%stemplates%spurchase_quotation_report.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Data']



        sql = """
              select pol.id                         as pol,
                     pol.id                         as pol_id,
                     ptemplate.name                 as product_name,
                     ctype.cost_category            as phan_loai_cphi,
                     pol.x_default_code,
                     pol.x_brand,
                     pol.x_origin,
                     uom.name                       as don_vi,
                     pol.product_qty,
                     purchase.name                  as so_don_hang,
                     pol.x_note                 as ghichu,
                     partner.name                   as nguoi_yc,
                     pa.name                  as so_yeu_cau,
                     case pol.x_purpose
                         when 'project' then 'Dự án'
                         when 'other' then 'Mục đích khác'
                         else 'Bổ sung tồn kho' end as x_purpose,
                     project.name                   as ma_du_an,
                     rp.name                        as nha_cung_cap,
                     pol.x_project_name,
                     pol.x_date_order,
                     pol.x_schedule_date,
                     pol.x_source,
                     pol.x_budget_price,
                     pol.x_date_received,
                     pol.price_unit,
                     pol.x_payment_ratio,
                     partner2.name                  as nguoi_dam_phan,
                     purchase.x_state as x_state_base,
                     -- state purchase order
                     case purchase.x_state
                         when 'draft' then 'Nháp'
                         when 'doing' then 'Đang thực hiện'
                         when 'out_date' then 'Đã quá hạn'
                         when 'delay' then 'Tạm hoãn'
                         when 'wait_license' then 'Đã giao chờ chứng từ'
                         when 'to_late' then 'Đã giao hàng muộn'
                         when 'done' then 'Hoàn thành'
                         when 'cancel' then 'Hủy'
                         else '' end                as x_state
              from purchase_order_line pol
                       left join product_product product on product.id = pol.product_id
                       left join purchase_requisition pa on pa.id = pol.x_requisition_id
                       left join product_template ptemplate on ptemplate.id = product.product_tmpl_id
                       left join cost_type as ctype on ctype.id = pol.x_cost_type_id
                       left join uom_uom as uom on uom.id = pol.product_uom
                       left join purchase_order purchase on purchase.id = pol.order_id
                       left join res_partner rp on rp.id = purchase.partner_id
                       left join res_users users on users.id = pol.x_users_id
                       left join res_partner partner on partner.id = users.partner_id
                       left join project_project project on project.id = pol.x_project_id
                       left join res_users users2 on users2.id = purchase.user_id
                       left join res_partner partner2 on partner2.id = users2.partner_id
              order by pol_id asc
              """

        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        highlight = NamedStyle(name="highlight1")
        highlight.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight1 = NamedStyle(name='datetime1', number_format='DD/MM/YYYY')
        highlight1.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight_number = NamedStyle(name='number1', number_format='#,##0')
        highlight_number.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight_number.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight_percent = NamedStyle(name='percent1', number_format='0.0%')
        highlight_percent.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight_percent.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        row = 4
        x = 1

        for r in recs:
            ws.cell(row, 1).value, ws.cell(row, 1).style = x, highlight
            ws.cell(row, 2).value, ws.cell(row, 2).style = r['so_don_hang'], highlight
            ws.cell(row, 3).value, ws.cell(row, 3).style = r['product_name'], highlight
            ws.cell(row, 4).value, ws.cell(row, 4).style = r['phan_loai_cphi'], highlight
            ws.cell(row, 5).value, ws.cell(row, 5).style = r['x_default_code'], highlight
            ws.cell(row, 6).value, ws.cell(row, 6).style = r['x_brand'], highlight
            ws.cell(row, 7).value, ws.cell(row, 7).style = r['x_origin'], highlight
            ws.cell(row, 8).value, ws.cell(row, 8).style = r['don_vi'], highlight
            ws.cell(row, 9).value, ws.cell(row, 9).style = r['product_qty'], highlight
            ws.cell(row, 10).value, ws.cell(row, 10).style = r['nguoi_yc'], highlight
            ws.cell(row, 11).value, ws.cell(row, 11).style = r['so_yeu_cau'], highlight
            ws.cell(row, 12).value, ws.cell(row, 12).style = r['x_purpose'], highlight
            ws.cell(row, 13).value, ws.cell(row, 13).style = r['ma_du_an'], highlight
            ws.cell(row, 14).value, ws.cell(row, 14).style = r['x_project_name'], highlight
            ws.cell(row, 15).value, ws.cell(row, 15).style = r['x_date_order'], highlight1
            ws.cell(row, 16).value, ws.cell(row, 16).style = r['x_schedule_date'], highlight1
            ws.cell(row, 17).value, ws.cell(row, 17).style = r['nha_cung_cap'], highlight
            ws.cell(row, 18).value, ws.cell(row, 18).style = r['x_budget_price'], highlight_number
            ws.cell(row, 19).value, ws.cell(row, 19).style = (r['x_date_received'] if r.get('x_state_base')
                                                                                      in ['done','wait_license','to_late']  else None, highlight1)
            ws.cell(row, 20).value, ws.cell(row, 20).style = r['x_state'], highlight
            ws.cell(row, 21).value, ws.cell(row, 21).style = r['price_unit'], highlight_number
            ws.cell(row, 22).value, ws.cell(row, 22).style = r['nguoi_dam_phan'], highlight
            ws.cell(row, 23).value, ws.cell(row, 23).style = (r['x_payment_ratio'] or 0) / 100.0, highlight_percent  # ghi số phân số + format %, tránh #VALUE! ở sheet Payment Status
            ws.cell(row, 24).value, ws.cell(row, 24).style = r['ghichu'], highlight
            x += 1
            row += 1
        ws.cell(1, 11).value = self.env.user.name
        ws.cell(2, 11).value = date.today()

        # write data sheet PIC
        row_pic = 7
        ws_pic = wb['PIC']
        list_requester = list({item['nguoi_yc'] for item in recs if item['nguoi_yc']})
        line_end = len(list_requester)
        for rec in list_requester:
            ws_pic.cell(row_pic, 1).value, ws_pic.cell(row_pic, 1).style = rec, highlight
            ws_pic.cell(row_pic, 2).value, ws_pic.cell(row_pic, 2).style = f"=COUNTIFS(Data!$J:$J,@$A:$A,Data!$T:$T,B$5)", highlight
            ws_pic.cell(row_pic, 3).value, ws_pic.cell(row_pic, 3).style = f"=COUNTIFS(Data!$J:$J,@$A:$A,Data!$T:$T,C$5)", highlight
            ws_pic.cell(row_pic, 4).value, ws_pic.cell(row_pic, 4).style = f"=COUNTIFS(Data!$J:$J,@$A:$A,Data!$T:$T,D$5)", highlight

            row_pic+=1
        total_row = line_end + 7
        end_data_row = line_end + 6
        ws_pic.cell(total_row, 1).value,ws_pic.cell(total_row, 1).style = f"TỔNG SỐ:" ,highlight
        ws_pic.cell(total_row, 2).value,ws_pic.cell(total_row, 2).style = f"=SUM(B7:B{end_data_row})",highlight
        ws_pic.cell(total_row, 3).value ,ws_pic.cell(total_row, 3).style= f"=SUM(C7:C{end_data_row})" , highlight
        ws_pic.cell(total_row, 4).value,ws_pic.cell(total_row, 4).style = f"=SUM(D7:D{end_data_row})" , highlight

        # write sheet Toplist
        ws_top_list = wb['Toplist']
        
        for row_idx in range(6, ws_top_list.max_row + 1):
            for col_idx in range(1, ws_top_list.max_column + 1):
                cell = ws_top_list.cell(row_idx, col_idx)
                cell.value = None
                cell.style = 'Normal'
        
        len_data = row
        row_supplier = 6
        number_row = 0
        list_supplier = list({item['nha_cung_cap'] for item in recs if item['nha_cung_cap']})
        for supplier in list_supplier:
            number_row += 1
            ws_top_list.cell(row_supplier, 1).value, ws_top_list.cell(row_supplier, 1).style = number_row, highlight
            ws_top_list.cell(row_supplier, 2).value, ws_top_list.cell(row_supplier,
                                                       2).style = supplier, highlight
            ws_top_list.cell(row_supplier, 3).value, ws_top_list.cell(row_supplier,
                                                       3).style = f"=SUMPRODUCT((Data!$I$4:$I${len_data})*(Data!$U$4:$U${len_data})*(Data!$Q$4:$Q${len_data}=B{row_supplier})*(Data!$S$4:$S${len_data}>=$C$2)*(Data!$S$4:$S${len_data}<=$C$3))", highlight_number
            row_supplier+=1

        total_sup_row = len(list_supplier) + 6
        end_data_sup_row = len(list_supplier) + 5
        ws_top_list.cell(total_sup_row, 2).value, ws_top_list.cell(total_sup_row, 2).style = f"TỔNG SỐ:", highlight
        ws_top_list.cell(total_sup_row, 3).value, ws_top_list.cell(total_sup_row, 3).style = f"=SUM(C6:C{end_data_sup_row})", highlight_number

        # write sheet Payment Status
        ws_ps = wb['Payment Status']
        
        columns = ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']
        
        for i, col in enumerate(columns, start=2):
            # Row 7: Amount paid
            formula_7 = f"""=SUMPRODUCT(((Data!$T$4:$T${len_data})<>"Hủy")*(YEAR(Data!$O$4:$O${len_data})={col}$5)*(MONTH(Data!$O$4:$O${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$U$4:$U${len_data})*(Data!$W$4:$W${len_data}))"""
            ws_ps.cell(7, i).value, ws_ps.cell(7, i).style = formula_7, highlight_number
            
            # Row 8: Unpaid amount  
            formula_8 = f"""=SUMPRODUCT((YEAR(Data!$O$4:$O${len_data})={col}$5)*(MONTH(Data!$O$4:$O${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$U$4:$U${len_data}))-{col}7"""
            ws_ps.cell(8, i).value, ws_ps.cell(8, i).style = formula_8, highlight_number
            
            # Row 9: Outstanding amount
            formula_9 = f"""=SUMPRODUCT(((Data!$T$4:$T${len_data})<>"Hủy")*(YEAR(Data!$O$4:$O${len_data})={col}$5)*(MONTH(Data!$O$4:$O${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$R$4:$R${len_data})*(100%-(Data!$W$4:$W${len_data})))"""
            ws_ps.cell(9, i).value, ws_ps.cell(9, i).style = formula_9, highlight_number

        # write sheet Efficiency
        ws_e = wb['Efficiency']

        eff_columns = ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N']
        
        for i, col in enumerate(eff_columns, start=3):
            # Row 8: Efficiency calculation
            formula_eff = f"""=SUMPRODUCT(((Data!$T$4:$T${len_data})<>"Hủy")*(YEAR(Data!$S$4:$S${len_data})={col}$5)*(MONTH(Data!$S$4:$S${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$R$4:$R${len_data}))"""
            ws_e.cell(8, i).value, ws_e.cell(8, i).style = formula_eff, highlight_number

            # Row 9: Purchase amount
            purchase_amount = f"""=SUMPRODUCT(((Data!$T$4:$T${len_data})<>"Hủy")*(YEAR(Data!$S$4:$S${len_data})={col}$5)*(MONTH(Data!$S$4:$S${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$U$4:$U${len_data}))"""
            ws_e.cell(9, i).value, ws_e.cell(9, i).style = purchase_amount, highlight_number

        row_nguoidamphan = 11
        list_nguoidamphan = list(set(item['nguoi_dam_phan'] for item in recs))
        for nguoidamphan in list_nguoidamphan:
            ws_e.cell(row_nguoidamphan, 2).value, ws_e.cell(row_nguoidamphan, 2).style = nguoidamphan if nguoidamphan else 'Others', highlight
            ws_e.cell(row_nguoidamphan, 15).value, ws_e.cell(row_nguoidamphan, 15).style =f"""=SUM(C{row_nguoidamphan}:N{row_nguoidamphan})""", highlight_number

            for i, col in enumerate(eff_columns, start=3):
                value_formula = f"""=SUMPRODUCT(((Data!$T$4:$T${len_data})<>"Hủy")*(YEAR(Data!$S$4:$S${len_data})={col}$5)*(MONTH(Data!$S$4:$S${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$R$4:$R${len_data})*(Data!$V$4:$V${len_data}=$B{row_nguoidamphan}))-SUMPRODUCT((YEAR(Data!$S$4:$S${len_data})={col}$5)*(MONTH(Data!$S$4:$S${len_data})={col}$6)*(Data!$I$4:$I${len_data})*(Data!$U$4:$U${len_data})*(Data!$V$4:$V${len_data}=$B{row_nguoidamphan}))"""
                ws_e.cell(row_nguoidamphan, i).value, ws_e.cell(row_nguoidamphan, i).style = value_formula, highlight_number

            row_nguoidamphan += 1

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Báo cáo báo giá mua hàng.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        # download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/' + str(attachment_id.id) + '?download=true',
            'target': 'new',
        }
