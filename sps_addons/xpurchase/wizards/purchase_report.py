# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side
from openpyxl.writer.excel import save_virtual_workbook
from odoo import api, fields, models


class ReportPurchaseOrder(models.TransientModel):
    _name = 'purchase.data.report'
    _description = 'Báo cáo data mua hàng'

    def purchase_report_data(self):
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%spurchase_report.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']

        sql = '''
                SELECT 
                    pt.name ten_san_pham, 
                    ct.cost_category, 
                    pol.x_default_code, 
                    pol.x_brand, 
                    pol.x_origin,
                    it.value don_vi, 
                    pol.x_qty_requisition,
                    rp2.name nguoi_yc,
                    case when pol.x_purpose = 'project' then 'Dự án'
                    when pol.x_purpose = 'replenish_inventory' then 'Bổ sung tồn kho'
                    else 'Mục đích khác'
                    end as x_purpose, 
                    jj.name ma_du_an, 
                    pol.x_project_name ten_du_an, 
                    pol.x_date_order, 
                    pol.x_schedule_date,
                    rp1.name nha_cc,
                    pol.x_budget_price,
                    case when po.x_state = 'draft' then 'Nháp'
                    when po.x_state = 'out_date' then 'Đã quá hạn'
                    when po.x_state = 'delay' then 'Tạm hoãn'
                    when po.x_state = 'doing' then 'Đang thực hiện'
                    when po.x_state = 'wait_license' then 'Đã giao chờ chứng từ'
                    when po.x_state = 'to_late' then 'Đã giao hàng muộn'
                    when po.x_state = 'done' then 'Hoàn thành'
                    when po.x_state = 'cancel' then 'Hủy'
                    end as state,
                    pol.price_unit,
                    po.x_wait_license_date as ngay_nhan,
                    rp.name as name_user,
                    pol.x_payment_ratio as x_payment_ratio
                FROM purchase_order po 
                LEFT JOIN purchase_order_line pol ON po.id = pol.order_id 
                LEFT JOIN uom_uom uu ON uu.id = pol.product_uom
                left join product_product pp on pp.id = pol.product_id 
                LEFT JOIN product_template pt ON pt.id = pp.product_tmpl_id  
                LEFT JOIN res_users ru ON ru.id = po.user_id
                LEFT JOIN res_partner rp ON ru.partner_id = rp.id
                LEFT JOIN res_partner rp1 ON po.partner_id = rp1.id
                LEFT JOIN cost_type ct ON ct.id = pol.x_cost_type_id
                LEFT JOIN project_project jj ON pol.x_project_id = jj.id
                LEFT JOIN res_users ru1 ON pol.x_users_id = ru1.id
                LEFT JOIN res_partner rp2 ON ru1.partner_id = rp2.id
                left join ir_translation it on it.name = 'uom.uom,name' and it.lang = 'vi_VN' and uu.id = it.res_id 
                left join 
                    (
                        select aax.po_id , sum(ap.amount) as amount_advance_payment
                        from account_advance aax
                        left join account_payment ap on ap.x_origin_advance_id = aax.id
                        left join account_move am2 on am2.payment_id = ap.id 
                        where aax.state in ('payment','completed') and am2.state = 'posted' and aax.po_id is not null
                        group by aax.po_id
                    ) advance_payment on advance_payment.po_id = po.id 
                left join 
                    (
                        select detail_invoice_payment.order_id, coalesce(sum(detail_invoice_payment.amount_total),0) - coalesce(sum(detail_invoice_payment.amount_residual),0) as amount_invoice_payment
                        from
                            (select distinct pol2.order_id , am.id as am_id, am.amount_total , am.amount_residual 
                            from purchase_order_line pol2
                            left join account_move_line aml on aml.purchase_line_id = pol2.id 
                            left join account_move am on am.id = aml.move_id and am.state = 'posted') detail_invoice_payment
                        group by detail_invoice_payment.order_id
                    ) invoice_payment on invoice_payment.order_id = po.id 
                WHERE po.x_state not in ('cancel')
                         '''

        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        highlight = NamedStyle(name="highlight")
        highlight.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight1 = NamedStyle(name='datetime', number_format='DD/MM/YYYY')
        highlight1.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        row = 4
        x = 1

        for r in recs:
            ws.cell(row, 1).value, ws.cell(row, 1).style = x, highlight
            ws.cell(row, 2).value, ws.cell(row, 2).style = r['ten_san_pham'], highlight
            ws.cell(row, 3).value, ws.cell(row, 3).style = r['cost_category'], highlight
            ws.cell(row, 4).value, ws.cell(row, 4).style = r['x_default_code'], highlight
            ws.cell(row, 5).value, ws.cell(row, 5).style = r['x_brand'], highlight
            ws.cell(row, 6).value, ws.cell(row, 6).style = r['x_origin'], highlight
            ws.cell(row, 7).value, ws.cell(row, 7).style = r['don_vi'], highlight
            ws.cell(row, 8).value, ws.cell(row, 8).style = r['x_qty_requisition'], highlight
            ws.cell(row, 9).value, ws.cell(row, 9).style = r['nguoi_yc'], highlight
            ws.cell(row, 10).value, ws.cell(row, 10).style = r['x_purpose'], highlight
            ws.cell(row, 11).value, ws.cell(row, 11).style = r['ma_du_an'], highlight
            ws.cell(row, 12).value, ws.cell(row, 12).style = r['ten_du_an'], highlight
            ws.cell(row, 13).value, ws.cell(row, 13).style = r['x_date_order'], highlight1
            ws.cell(row, 14).value, ws.cell(row, 14).style = r['x_schedule_date'], highlight1
            ws.cell(row, 15).value, ws.cell(row, 15).style = r['nha_cc'], highlight
            ws.cell(row, 16).value, ws.cell(row, 16).style = r['x_budget_price'], highlight
            ws.cell(row, 17).value, ws.cell(row, 17).style = r['ngay_nhan'], highlight1
            ws.cell(row, 18).value, ws.cell(row, 18).style = r['state'], highlight
            ws.cell(row, 19).value, ws.cell(row, 19).style = r['price_unit'], highlight
            ws.cell(row, 20).value, ws.cell(row, 20).style = r['name_user'], highlight
            ws.cell(row, 21).value, ws.cell(row, 21).style = r['x_payment_ratio'], highlight
            x += 1
            row += 1
        ws.cell(1, 11).value, ws.cell(row, 21).style = self.env.user.name, highlight
        ws.cell(2, 11).value, ws.cell(row, 21).style = date.today(), highlight

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Báo cáo data mua hàng.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        # download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/' + str(attachment_id.id) + '?download=true',
            'target': 'new',
        }
