# -*- coding: utf-8 -*-

from odoo import models, fields
import base64
import os
from datetime import date

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import NamedStyle, Font, Side, PatternFill, Border, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from six import BytesIO


class PaymentOrder(models.Model):
    _inherit = 'account.advance'
    _description = 'Đề nghị thanh toán phiếu tạm ứng'

    def get_job_description(self):
        self.ensure_one()
        name_user = self.env.user.name
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%sadvanced_repay.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']

        if self.type == 'internal':
            sql = f'''
                    SELECT
                        hep.NAME nhan_vien,
                        hep.x_bank_account so_tk,
                        hep.x_bank_name ngan_hang,
                        hep.x_living dia_chi,
                        pp.name ma_duan,
                        aal.project_name ten_duan,
                        case when aal.amount is not null then aal.amount else 0 end so_tien,
                        aal.description  noi_dung
                    FROM
                        account_advance aa
                        LEFT JOIN res_users ru ON aa.employee_id = ru.ID 
                        LEFT JOIN hr_employee_public hep ON hep.user_id = ru.id
                        LEFT JOIN account_advance_line aal on aal.advance_id = aa.id
                        LEFT JOIN project_project pp on aal.project_id = pp.id 
                    WHERE
                        aa.id = {self.id}
                  '''
        else:
            sql = f'''
                    SELECT
                            rp1.NAME,
                            rp.NAME name1,
                            aa.amount_total tien,
                            rpb.acc_number so_tk,
                            rb.name ngan_hang,
                            rpb.acc_holder_name as nguoi_th,
                            concat(rp.street, case when rp.street is not null and rp.city is not null then ',' else '' end, rp.city) dia_chi,
                            pt.name ten_san_pham,
                            case when ct.cost_category is not null then ct.cost_category else '' end phan_loai,
                            case when ct.cost_group is not null then ct.cost_group else '' end nhom ,
                            it.value don_vi,
                            pol.product_qty so_luong,
                            pol.price_unit don_gia,
                            pol.price_subtotal bf_tax,
                            pol.price_tax tax,
                            pol.price_total at_tax,
                            pol.x_note ghi_chu,
                            case when ppj.name is not null then ppj.name else '' end ma_du_an,
                            case when pol.x_project_name is not null then pol.x_project_name else '' end ten_du_an
                    FROM
                            account_advance aa
                            LEFT JOIN res_partner rp ON aa.supplier_id = rp.ID 
                            LEFT JOIN res_partner_bank rpb on aa.partner_bank_id = rpb.id
                            LEFT JOIN res_bank rb on rb.id = rpb.bank_id
                            LEFT JOIN purchase_order po on po.id = aa.po_id
                            LEFT JOIN purchase_order_line pol on po.id = pol.order_id
                            LEFT JOIN product_product pp on pol.product_id = pp.id
                            LEFT JOIN product_template pt on pt.id = pp.product_tmpl_id
                            LEFT JOIN cost_type ct on ct.id = pol.x_cost_type_id
                            LEFT JOIN uom_uom uu on uu.id = pol.product_uom
                            LEFT JOIN ir_translation it ON it.NAME = 'uom.uom,name' AND it.lang = 'vi_VN' AND uu.NAME = it.src and uu.id = it.res_id
                            LEFT JOIN project_project ppj on pol.x_project_id = ppj.id
                            LEFT JOIN sale_order so on so.id = ppj.x_order_id
                            LEFT JOIN res_partner rp1 on rp1.id = so.partner_id
                    WHERE 
                            aa.id = {self.id}
                  '''
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dashed', color="000000")
        bd3 = Side(style='mediumDashed', color="0000FF")
        bd4 = Side(style='thick', color="0000FF")

        highlight = NamedStyle(name="highlight")
        highlight.font = Font(name='Times New Roman', size=13, bold=True)
        highlight.border = Border(bottom=bd1, top=bd1)

        highlight1 = NamedStyle(name="highlight1")
        highlight1.font = Font(name='Times New Roman', size=13)
        highlight1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight1.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
        highlight1.number_format = '#,##0'

        highlight2 = NamedStyle(name="highlight2")
        highlight2.font = Font(name='Times New Roman', size=13, bold=True)
        highlight2.border = Border(bottom=bd2)

        highlight3 = NamedStyle(name="highlight3")
        highlight3.font = Font(name='Times New Roman', size=13)
        highlight3.border = Border(bottom=bd1)

        highlight4 = NamedStyle(name="highlight4")
        highlight4.font = Font(name='Times New Roman', size=13)
        highlight4.border = Border(bottom=bd3)

        highlight5 = NamedStyle(name="highlight5")
        highlight5.font = Font(name='Times New Roman', size=13, underline="single", bold=True)
        highlight5.alignment = Alignment(vertical='center')

        highlight6 = NamedStyle(name="highlight6")
        highlight6.font = Font(name='Times New Roman', size=13)
        highlight6.border = Border(bottom=bd4)

        highlight7 = NamedStyle(name="highlight7")
        highlight7.font = Font(name='Times New Roman', size=13)
        highlight7.border = Border(right=bd4)

        highlight8 = NamedStyle(name="highlight8")
        highlight8.font = Font(name='Times New Roman', size=13)
        highlight8.alignment = Alignment(vertical='center', wrap_text=True)

        highlight9 = NamedStyle(name="highlight9")
        highlight9.font = Font(name='Times New Roman', size=13)
        highlight9.alignment = Alignment(vertical='center', horizontal='center')

        highlight10 = NamedStyle(name="highlight10")
        highlight10.font = Font(name='Times New Roman', size=13, bold=True)
        highlight10.alignment = Alignment(vertical='center', horizontal='center')
        highlight10.number_format = '#,##0'

        highlight11 = NamedStyle(name="highlight11")
        highlight11.font = Font(name='Times New Roman', size=13, bold=True)
        highlight11.alignment = Alignment(vertical='center', horizontal='center')
        highlight11.border = Border(bottom=bd1, top=bd1)
        highlight11.number_format = '#,##0'

        highlight12 = NamedStyle(name="highlight12")
        highlight12.font = Font(name='Times New Roman', size=13)
        highlight12.alignment = Alignment(vertical='center')

        highlight13 = NamedStyle(name="highlight13")
        highlight13.font = Font(name='Times New Roman', size=13)
        highlight13.alignment = Alignment(vertical='center')
        highlight13.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight14 = NamedStyle(name="highlight14")
        highlight14.font = Font(name='Times New Roman', size=13)
        highlight14.border = Border(bottom=bd2)

        highlight15 = NamedStyle(name="highlight15")
        highlight15.font = Font(name='Times New Roman', size=13, bold=True)
        highlight15.alignment = Alignment(vertical='center', horizontal='center')

        highlight16 = NamedStyle(name="highlight16")
        highlight16.font = Font(name='Times New Roman', size=13, bold=True, underline="single")
        highlight16.alignment = Alignment(vertical='center')
        highlight16.border = Border(bottom=bd2)

        highlight17 = NamedStyle(name="highlight17")
        highlight17.font = Font(name='Times New Roman', size=13)
        highlight17.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight17.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
        highlight17.number_format = 'DD/MM/YYYY'

        highlight18 = NamedStyle(name="highlight18")
        highlight18.font = Font(name='Times New Roman', size=13)
        highlight18.alignment = Alignment(vertical='center', horizontal='left')
        highlight18.number_format = 'DD/MM/YYYY'

        row = 6
        amount_total = 0
        x = 1

        if self.type == 'supplier':
            bf_tax = 0
            amount_tax = 0
            at_tax = 0
            dict_item = {}
            for rec in recs:
                if rec['nhom'] + rec['ma_du_an'] not in dict_item:
                    dict_item[rec['nhom'] + rec['ma_du_an']] = [rec]
                else:
                    dict_item[rec['nhom'] + rec['ma_du_an']].append(rec)

            ws.cell(4, 3).value, ws.cell(4, 3).style = self.env.user.name, highlight12
            ws.cell(4, 12).value, ws.cell(4, 12).style = fields.Date.today(), highlight18
            # Chỉnh độ rộng của cột
            ws.column_dimensions[get_column_letter(3)].width = 13
            ws.column_dimensions[get_column_letter(4)].width = 37
            ws.column_dimensions[get_column_letter(6)].width = 10
            ws.column_dimensions[get_column_letter(7)].width = 10
            ws.column_dimensions[get_column_letter(10)].width = 13
            ws.column_dimensions[get_column_letter(12)].width = 13
            for name, values in dict_item.items():
                ws.row_dimensions[row].height = 50
                ws.cell(row, 1).value, ws.cell(row, 1).style = values[0]['nhom'], highlight5
                ws.cell(row, 5).value, ws.cell(row, 5).style = 'MÃ DỰ ÁN:', highlight9
                ws.merge_cells(start_row=row, start_column=6, end_row=row,
                               end_column=7)
                ws.cell(row, 6).value, ws.cell(row, 6).style = values[0]['ma_du_an'], highlight8
                ws.cell(row, 8).value, ws.cell(row, 8).style = 'TÊN DỰ ÁN:', highlight9
                ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=12)
                ws.cell(row, 9).value, ws.cell(row, 9).style = values[0]['ten_du_an'], highlight8
                row += 1
                ws.cell(row, 2).value, ws.cell(row, 2).style = 'KHÁCH HÀNG:', highlight9
                ws.merge_cells(start_row=row, start_column=3, end_row=row,
                               end_column=4)
                ws.cell(row, 3).value, ws.cell(row, 3).style = values[0]['name'], highlight8
                ws.cell(row, 8).value, ws.cell(row, 8).style = 'TỔNG SỐ', highlight9
                ws.cell(row, 9).value, ws.cell(row, 9).style = sum([total['bf_tax'] for total in values]), highlight10
                ws.cell(row, 10).value, ws.cell(row, 10).style = sum([total['tax'] for total in values]), highlight10
                ws.cell(row, 11).value, ws.cell(row, 11).style = sum([total['at_tax'] for total in values]), highlight10
                row += 1
                for value in values:
                    ws.cell(row, 1).value, ws.cell(row, 1).style = x, highlight1
                    ws.cell(row, 2).style = highlight1
                    ws.cell(row, 3).style = highlight17
                    ws.cell(row, 4).value, ws.cell(row, 4).style = value['ten_san_pham'], highlight13
                    ws.cell(row, 5).value, ws.cell(row, 5).style = value['phan_loai'], highlight1
                    ws.cell(row, 6).value, ws.cell(row, 6).style = value['don_vi'], highlight1
                    ws.cell(row, 7).value, ws.cell(row, 7).style = value['so_luong'], highlight1
                    ws.cell(row, 8).value, ws.cell(row, 8).style = value['don_gia'], highlight1
                    ws.cell(row, 9).value, ws.cell(row, 9).style = value['bf_tax'], highlight1
                    ws.cell(row, 10).value, ws.cell(row, 10).style = value['tax'], highlight1
                    ws.cell(row, 11).value, ws.cell(row, 11).style = value['at_tax'], highlight1
                    ws.cell(row, 12).value, ws.cell(row, 12).style = value['ghi_chu'], highlight1
                    row += 1
                    x += 1
                    bf_tax += value['bf_tax']
                    amount_tax += value['tax']
                    at_tax += value['at_tax']
                x = 1
            temp_total_amount = 0
            for r in self.search([('po_id', '=', self.po_id.id)]):
                amount_total += r.amount_total

            for r in self.search([('po_id', '=', self.po_id.id)]):
                temp_total_amount += r.amount_total
                if r.id == self.id:
                    break
            amount_total -= temp_total_amount
            self._fill_border_bottom_solid(ws, row, highlight)
            ws.cell(row, 1).value = 'Tổng số tiền (VNĐ)'
            ws.cell(row, 9).value, ws.cell(row, 9).style = bf_tax, highlight11
            ws.cell(row, 10).value, ws.cell(row, 10).style = amount_tax, highlight11
            ws.cell(row, 11).value, ws.cell(row, 11).style = at_tax, highlight11
            row += 1
            self._fill_border_bottom_solid(ws, row, highlight)
            ws.cell(row, 1).value = 'Số tiền đã tạm ứng (VNĐ)'
            ws.cell(row, 11).value, ws.cell(row, 11).style = amount_total, highlight11
            row += 1
            self._fill_border_bottom_solid(ws, row, highlight)
            ws.cell(row, 1).value = 'Số tiền công ty phải thanh toán'
            ws.cell(row, 11).value, ws.cell(row, 11).style = recs[0].get('tien', 0), highlight11
            row += 1
            self._fill_border_bottom_solid(ws, row, highlight2)
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'Phương thức thanh toán:', highlight16
            ws.cell(row, 4).value, ws.cell(row, 4).style = '☐ ' + 'Tiền mặt / Chuyển khoản cá nhân', highlight14
            ws.cell(row, 5).value, ws.cell(row, 5).style = '☐ ' + "CK Techcombank", highlight14
            ws.cell(row, 7).value, ws.cell(row, 7).style = '☐ ' + "CK Vietcombank", highlight14
            ws.cell(row, 9).value, ws.cell(row, 9).style = '☐ ' + "CK VPBank", highlight14
            ws.cell(row, 11).value, ws.cell(row, 11).style = '☐ ' + "CK BIDV", highlight14
            row += 1
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'Thông tin người nhận tiền:', highlight5
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Số tài khoản:',highlight12
            ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('so_tk', None),highlight12
            row += 1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Tên nhà cung cấp: ',highlight12
            ws.cell(row, 3).value, ws.cell(row, 3).style = recs[0].get('name1', None),highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Người thụ hưởng:',highlight12
            ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('nguoi_th', None),highlight12
            row += 1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Địa chỉ: ',highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Ngân hàng',highlight12
            ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('ngan_hang', None),highlight12
            self._fill_border_bottom_solid(ws, row, highlight3)
        else:
            ws.cell(4, 3).value, ws.cell(4, 3).style = recs[0]['nhan_vien'], highlight12
            ws.cell(4, 12).value, ws.cell(4, 12).style = fields.Date.today(), highlight18
            dict_item = {}
            for rec in recs:
                if rec['ma_duan'] not in dict_item:
                    dict_item[rec['ma_duan']] = [rec]
                else:
                    dict_item[rec['ma_duan']].append(rec)
            # Chỉnh độ rộng của cột
            ws.column_dimensions[get_column_letter(3)].width = 13
            ws.column_dimensions[get_column_letter(4)].width = 37
            ws.column_dimensions[get_column_letter(6)].width = 10
            ws.column_dimensions[get_column_letter(7)].width = 10
            ws.column_dimensions[get_column_letter(10)].width = 13
            ws.column_dimensions[get_column_letter(12)].width = 13
            for key, values in dict_item.items():
                ws.row_dimensions[row].height = 50
                ws.cell(row, 5).value, ws.cell(row, 5).style = 'MÃ DỰ ÁN:', highlight9
                ws.merge_cells(start_row=row, start_column=6, end_row=row,
                               end_column=7)
                ws.cell(row, 6).value, ws.cell(row, 6).style = values[0]['ma_duan'], highlight8
                ws.cell(row, 8).value, ws.cell(row, 8).style = 'TÊN DỰ ÁN:', highlight9
                ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=12)
                ws.cell(row, 9).value, ws.cell(row, 9).style = values[0]['ten_duan'], highlight8
                row += 1
                ws.cell(row, 8).value, ws.cell(row, 8).style = 'TỔNG SỐ', highlight9
                ws.cell(row, 9).value, ws.cell(row, 9).style = sum([total['so_tien'] for total in values]), highlight10
                ws.cell(row, 10).value, ws.cell(row, 10).style = 0, highlight10
                ws.cell(row, 11).value, ws.cell(row, 11).style = sum([total['so_tien'] for total in values]), highlight10
                row += 1
                for value in values:
                    ws.cell(row, 1).style = highlight13
                    ws.cell(row, 2).style = highlight13
                    ws.cell(row, 3).style = highlight13
                    ws.cell(row, 4).value, ws.cell(row, 4).style = value.get('noi_dung', None), highlight13
                    ws.cell(row, 5).style = highlight13
                    ws.cell(row, 6).style = highlight13
                    ws.cell(row, 7).style = highlight13
                    ws.cell(row, 8).style = highlight13
                    ws.cell(row, 9).value, ws.cell(row, 9).style = value.get('so_tien', 0), highlight1
                    ws.cell(row, 10).value, ws.cell(row, 10).style = 0, highlight1
                    ws.cell(row, 11).value, ws.cell(row, 11).style = value.get('so_tien', 0), highlight1
                    ws.cell(row, 12).style = highlight13
                    amount_total += value.get('so_tien', 0)
                    row += 1
                    x += 1
                x = 1
            self._fill_border_bottom_solid(ws, row, highlight)
            ws.cell(row, 1).value = 'Tổng số tiền (VNĐ)'
            ws.cell(row, 9).value, ws.cell(row, 9).style = amount_total, highlight11
            ws.cell(row, 10).value, ws.cell(row, 10).style = 0, highlight11
            ws.cell(row, 11).value, ws.cell(row, 11).style = amount_total, highlight11
            row += 1
            self._fill_border_bottom_solid(ws, row, highlight)
            ws.cell(row, 1).value = 'Số tiền đã tạm ứng (VNĐ)'
            ws.cell(row, 11).value, ws.cell(row, 11).style = 0, highlight11
            if self.type == 'internal':
                if self.license_type == 'internal':
                    row += 1
                    self._fill_border_bottom_solid(ws, row, highlight)
                    ws.cell(row, 1).value = 'Số tiền công ty phải thanh toán'
                    ws.cell(row, 11).value, ws.cell(row, 11).style = amount_total, highlight11
                else:
                    row += 1
                    self._fill_border_bottom_solid(ws, row, highlight)
                    ws.cell(row, 1).value = 'Số tiền công ty phải thanh toán.'
                    ws.cell(row, 11).value, ws.cell(row, 11).style = amount_total, highlight11
                row += 1
                self._fill_border_bottom_solid(ws, row, highlight2)
                ws.cell(row, 1).value, ws.cell(row, 1).style = 'Phương thức thanh toán:', highlight16
                ws.cell(row, 4).value, ws.cell(row, 4).style = '☐ ' + 'Tiền mặt / Chuyển khoản cá nhân', highlight14
                ws.cell(row, 5).value, ws.cell(row, 5).style = '☐ ' + "CK Techcombank", highlight14
                ws.cell(row, 7).value, ws.cell(row, 7).style = '☐ ' + "CK Vietcombank", highlight14
                ws.cell(row, 9).value, ws.cell(row, 9).style = '☐ ' + "CK VPBank", highlight14
                ws.cell(row, 11).value, ws.cell(row, 11).style = '☐ ' + "CK BIDV", highlight14
                row += 1
                ws.cell(row, 1).value, ws.cell(row, 1).style = 'Thông tin người nhận tiền:', highlight5
                ws.cell(row, 7).value, ws.cell(row, 7).style = 'Số tài khoản:',highlight12
                ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('so_tk', None),highlight12
                row += 1
                ws.cell(row, 2).value, ws.cell(row, 2).style = 'Tên nhân viên:',highlight12
                ws.cell(row, 3).value, ws.cell(row, 3).style = recs[0].get('nhan_vien', None),highlight12
                ws.cell(row, 7).value, ws.cell(row, 7).style = 'Người thụ hưởng:',highlight12
                ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('nhan_vien', None),highlight12
                row += 1
                ws.cell(row, 2).value, ws.cell(row, 2).style = 'Địa chỉ:',highlight12

                ws.cell(row, 7).value, ws.cell(row, 7).style = 'Ngân hàng',highlight12
                ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('ngan_hang', None),highlight12
                self._fill_border_bottom_solid(ws, row, highlight3)
        row += 1
        ws.cell(row, 2).value, ws.cell(row, 2).style = 'Người đề nghị', highlight15
        ws.cell(row, 4).value, ws.cell(row, 4).style = 'Trưởng bộ phận', highlight15
        ws.cell(row, 7).value, ws.cell(row, 7).style = 'Phụ trách kế toán', highlight15
        ws.cell(row, 11).value, ws.cell(row, 11).style = 'Tổng Giám Đốc', highlight15

        row += 4
        ws.cell(row, 2).value, ws.cell(row, 2).style = self.env.user.name, highlight12
        self._fill_border_bottom_solid(ws, row, highlight3)

        if self.license_type != 'tax':
            row += 2
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'Người nhận tiền xác nhận', highlight5

            row += 2
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Số tiền đã nhận: ',highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Chữ kí người nhận tiền:',highlight12

            row += 2
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Ngày nhận tiền:',highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Họ và tên người nhận tiền:',highlight12

        if self.license_type == 'tax':
            ws.cell(3, 1).value = 'GIẤY ĐỀ NGHỊ THANH TOÁN '
        else:
            ws.cell(3, 1).value = 'GIẤY ĐỀ NGHỊ THANH TOÁN / TẠM ỨNG KIÊM PHIẾU CHI (NỘI BỘ)'
        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Đề nghị thanh toán phiếu tạm ứng.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        return attachment_id

    def _fill_border_bottom_solid(self, ws, row, highlight):
        for i in range(1, 13):
            ws.cell(row, i).style = highlight
