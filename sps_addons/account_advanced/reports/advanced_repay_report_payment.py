# -*- coding: utf-8 -*-

from odoo import models, fields
import base64
import os
from datetime import date

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import NamedStyle, Font, Side, PatternFill, Border, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from six import BytesIO


class PaymentOrder(models.Model):
    _inherit = 'account.advance.repay'
    _description = 'Đề nghị thanh toán phiếu hoàn ứng'

    def get_job_payment_description(self):
        self.ensure_one()
        name_user = self.env.user.name
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%sadvanced_repay.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']
        if self.license_type == 'internal':
            ws.cell(3, 1).value = 'GIẤY ĐỀ NGHỊ THANH TOÁN / TẠM ỨNG KIÊM PHIẾU CHI (NỘI BỘ)'
        else:
            ws.cell(3, 1).value = 'GIẤY ĐỀ NGHỊ THANH TOÁN'
        sql = f'''
                    SELECT
                        rp2.name as ncc,
                        hep.name as ncc_1,
                        aar.beneficiary as nhan_vien,
                        aar.date_suggest::date ngay,
                        pp.NAME ma_duan,
                        arl.project_name ten_duan,
                        rp1.NAME ten_kh,
                        arl.content noi_dung,
                        ct.cost_category phan_loai,
                        coalesce(ct.cost_group, '') nhom,
                        it.VALUE don_vi ,
                        arl.qty so_luong,
                        arl.invoice_number so_hd,
                        arl.date_invoice ngay_hd,
                        case when arl.price_unit is not null then arl.price_unit else 0 end don_gia,
                        case when arl.amount_untaxed is not null then arl.amount_untaxed else 0 end bf_tax,
                        case when arl.amount_tax - arl.amount_untaxed is not null then arl.amount_tax - arl.amount_untaxed else 0 end tax,
                        case when arl.amount_tax is not null then arl.amount_tax else 0 end at_tax,
                        arl.note ghi_chu,
                        case when aar.amount_advance is not null then aar.amount_advance else 0 end tam_ung,
                        aar.still_pay con_phai_thanh_toan,
                        rpb.acc_number so_tk,
                        rb.name ngan_hang,
                        case when aar.repay_type = 'employee' then hep.x_living 
                        else rp2.street
                        end as dia_chi
                    FROM
                        account_advance_repay aar
                        LEFT JOIN res_users ru ON aar.employee_id = ru.ID 
                        LEFT JOIN hr_employee_public hep ON hep.user_id = ru.id
                        LEFT JOIN account_repay_line arl ON arl.account_repay_id = aar.ID
                        LEFT JOIN project_project pp ON pp.ID = arl.project_id
                        LEFT JOIN res_partner rp1 ON rp1.ID = pp.partner_id
                        LEFT JOIN cost_type ct ON ct.ID = arl.cost_type_id
                        LEFT JOIN uom_uom uu ON uu.ID = arl.uom_id
                        LEFT JOIN ir_translation it ON it.NAME = 'uom.uom,name' AND it.lang = 'vi_VN' AND uu.NAME = it.src and uu.id = it.res_id
                        LEFT JOIN res_partner rp2 on rp2.id = aar.supplier_id
                        left join res_partner_bank rpb on rpb.id = aar.partner_bank_id
                        left join res_bank rb on rb.id = rpb.bank_id
                    WHERE
                        aar.ID = {self.id}
                    ORDER BY ct.cost_group ,  arl.project_name
                    '''
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dashed', color="000000")
        bd3 = Side(style='mediumDashed', color="0000FF")
        bd4 = Side(style='thick', color="0000FF")

        highlight = NamedStyle(name="highlight")
        highlight.font = Font(name='Times New Roman', size=13, bold=True)
        highlight.border = Border(bottom=bd1, top=bd1)

        highlight1 = NamedStyle(name="highlight1")
        highlight1.font = Font(name='Times New Roman', size=13)
        highlight1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight1.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
        highlight1.number_format = '#,##0'

        highlight2 = NamedStyle(name="highlight2")
        highlight2.font = Font(name='Times New Roman', size=13, bold=True)
        highlight2.border = Border(bottom=bd2)

        highlight3 = NamedStyle(name="highlight3")
        highlight3.font = Font(name='Times New Roman', size=13)
        highlight3.border = Border(bottom=bd1)

        highlight4 = NamedStyle(name="highlight4")
        highlight4.font = Font(name='Times New Roman', size=13)
        highlight4.border = Border(bottom=bd3)

        highlight5 = NamedStyle(name="highlight5")
        highlight5.font = Font(name='Times New Roman', size=13, underline="single", bold=True)
        highlight5.alignment = Alignment(vertical='center')

        highlight6 = NamedStyle(name="highlight6")
        highlight6.font = Font(name='Times New Roman', size=13)
        highlight6.border = Border(bottom=bd4)

        highlight7 = NamedStyle(name="highlight7")
        highlight7.font = Font(name='Times New Roman', size=13)
        highlight7.border = Border(right=bd4)

        highlight8 = NamedStyle(name="highlight8")
        highlight8.font = Font(name='Times New Roman', size=13)
        highlight8.alignment = Alignment(vertical='center', wrap_text=True)

        highlight9 = NamedStyle(name="highlight9")
        highlight9.font = Font(name='Times New Roman', size=13)
        highlight9.alignment = Alignment(vertical='center', horizontal='center')

        highlight10 = NamedStyle(name="highlight10")
        highlight10.font = Font(name='Times New Roman', size=13, bold=True)
        highlight10.alignment = Alignment(vertical='center', horizontal='center')
        highlight10.number_format = '#,##0'

        highlight11 = NamedStyle(name="highlight11")
        highlight11.font = Font(name='Times New Roman', size=13, bold=True)
        highlight11.alignment = Alignment(vertical='center', horizontal='center')
        highlight11.border = Border(bottom=bd1, top=bd1)
        highlight11.number_format = '#,##0'

        highlight12 = NamedStyle(name="highlight12")
        highlight12.font = Font(name='Times New Roman', size=13)
        highlight12.alignment = Alignment(vertical='center')

        highlight13 = NamedStyle(name="highlight13")
        highlight13.font = Font(name='Times New Roman', size=13)
        highlight13.alignment = Alignment(vertical='center',wrap_text=True)
        highlight13.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight14 = NamedStyle(name="highlight14")
        highlight14.font = Font(name='Times New Roman', size=13)
        highlight14.border = Border(bottom=bd2)

        highlight15 = NamedStyle(name="highlight15")
        highlight15.font = Font(name='Times New Roman', size=13, bold=True)
        highlight15.alignment = Alignment(vertical='center', horizontal='center')

        highlight16 = NamedStyle(name="highlight16")
        highlight16.font = Font(name='Times New Roman', size=13, bold=True, underline="single")
        highlight16.alignment = Alignment(vertical='center')
        highlight16.border = Border(bottom=bd2)

        highlight17 = NamedStyle(name="highlight17")
        highlight17.font = Font(name='Times New Roman', size=13)
        highlight17.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight17.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
        highlight17.number_format = 'DD/MM/YYYY'

        highlight18 = NamedStyle(name="highlight18")
        highlight18.font = Font(name='Times New Roman', size=13)
        highlight18.alignment = Alignment(vertical='center', horizontal='left')
        highlight18.number_format = 'DD/MM/YYYY'

        dict_item = {}
        for rec in recs:
            if rec['nhom'] + rec['ma_duan'] not in dict_item:
                dict_item[rec['nhom'] + rec['ma_duan']] = [rec]
            else:
                dict_item[rec['nhom'] + rec['ma_duan']].append(rec)

        row = 6
        x = 1
        bf_tax = 0
        amount_tax = 0
        at_tax = 0
        # Chỉnh độ rộng của cột
        ws.column_dimensions[get_column_letter(3)].width = 13
        ws.column_dimensions[get_column_letter(4)].width = 37
        ws.column_dimensions[get_column_letter(6)].width = 10
        ws.column_dimensions[get_column_letter(7)].width = 10
        ws.column_dimensions[get_column_letter(10)].width = 10
        ws.column_dimensions[get_column_letter(12)].width = 13
        for name, values in dict_item.items():
            ws.row_dimensions[row].height = 50
            ws.cell(row, 1).value, ws.cell(row, 1).style = values[0]['nhom'], highlight5
            ws.cell(row, 5).value, ws.cell(row, 5).style = 'MÃ DỰ ÁN:', highlight9
            ws.merge_cells(start_row=row, start_column=6, end_row=row,
                           end_column=7)
            ws.cell(row, 6).value, ws.cell(row, 6).style = values[0]['ma_duan'], highlight8
            ws.cell(row, 8).value, ws.cell(row, 8).style = 'TÊN DỰ ÁN:', highlight9
            ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=12)
            ws.cell(row, 9).value, ws.cell(row, 9).style = values[0]['ten_duan'], highlight8
            row += 1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'KHÁCH HÀNG:', highlight9
            ws.merge_cells(start_row=row, start_column=3, end_row=row,
                           end_column=4)
            ws.cell(row, 3).value, ws.cell(row, 3).style = values[0]['ten_kh'], highlight8
            ws.cell(row, 8).value, ws.cell(row, 8).style = 'TỔNG SỐ', highlight9
            ws.cell(row, 9).value, ws.cell(row, 9).style = sum([total['bf_tax'] for total in values]), highlight10
            ws.cell(row, 10).value, ws.cell(row, 10).style = sum([total['tax'] for total in values]), highlight10
            ws.cell(row, 11).value, ws.cell(row, 11).style = sum([total['at_tax'] for total in values]), highlight10
            row += 1
            for value in values:
                ws.cell(row, 1).value, ws.cell(row, 1).style = x, highlight1
                ws.cell(row, 2).value, ws.cell(row, 2).style = value['so_hd'], highlight1
                ws.cell(row, 3).value, ws.cell(row, 3).style = value['ngay_hd'], highlight17
                ws.cell(row, 4).value, ws.cell(row, 4).style = value['noi_dung'], highlight13
                ws.cell(row, 5).value, ws.cell(row, 5).style = value['phan_loai'], highlight1
                ws.cell(row, 6).value, ws.cell(row, 6).style = value['don_vi'], highlight1
                ws.cell(row, 7).value, ws.cell(row, 7).style = value['so_luong'], highlight1
                ws.cell(row, 8).value, ws.cell(row, 8).style = value['don_gia'], highlight1
                ws.cell(row, 9).value, ws.cell(row, 9).style = value['bf_tax'], highlight1
                ws.cell(row, 10).value, ws.cell(row, 10).style = value['tax'], highlight1
                ws.cell(row, 11).value, ws.cell(row, 11).style = value['at_tax'], highlight1
                ws.cell(row, 12).value, ws.cell(row, 12).style = value['ghi_chu'], highlight1
                row += 1
                x += 1
                bf_tax += value['bf_tax']
                amount_tax += value['tax']
                at_tax += value['at_tax']
            x = 1
        ws.cell(4, 3).value, ws.cell(4, 3).style = self.create_uid.partner_id.name, highlight12
        ws.cell(4, 12).value, ws.cell(4, 12).style = fields.Date.today(), highlight18
        self._fill_border_bottom_solid(ws, row, highlight)
        ws.cell(row, 1).value = 'Tổng số tiền (VNĐ)'
        ws.cell(row, 9).value, ws.cell(row, 9).style = bf_tax, highlight11
        ws.cell(row, 10).value, ws.cell(row, 10).style = amount_tax, highlight11
        ws.cell(row, 11).value, ws.cell(row, 11).style = at_tax, highlight11

        row += 1
        self._fill_border_bottom_solid(ws, row, highlight)
        ws.cell(row, 1).value = 'Số tiền đã tạm ứng (VNĐ)'
        ws.cell(row, 11).value, ws.cell(row, 11).style = recs[0].get('tam_ung', 0), highlight11

        row += 1
        self._fill_border_bottom_solid(ws, row, highlight)
        if self.still_pay > 0:
            ws.cell(row, 1).value = 'Số tiền cá nhân hoàn lại'
            ws.cell(row, 11).value, ws.cell(row, 11).style = abs(at_tax - recs[0].get('tam_ung', 0)), highlight11
        else:
            ws.cell(row, 1).value = 'Số tiền công ty phải thanh toán'
            ws.cell(row, 11).value, ws.cell(row, 11).style = abs(at_tax - recs[0].get('tam_ung', 0)), highlight11

        row += 1
        self._fill_border_bottom_solid(ws, row, highlight2)
        ws.cell(row, 1).value, ws.cell(row, 1).style = 'Phương thức thanh toán:', highlight16
        ws.cell(row, 4).value, ws.cell(row, 4).style = '☐ ' + 'Tiền mặt / Chuyển khoản cá nhân', highlight14
        ws.cell(row, 5).value, ws.cell(row, 5).style = '☐ ' + "CK Techcombank", highlight14
        ws.cell(row, 7).value, ws.cell(row, 7).style = '☐ ' + "CK Vietcombank", highlight14
        ws.cell(row, 9).value, ws.cell(row, 9).style = '☐ ' + "CK VPBank", highlight14
        ws.cell(row, 11).value, ws.cell(row, 11).style = '☐ ' + "CK BIDV", highlight14

        row += 1
        ws.cell(row, 1).value, ws.cell(row, 1).style = 'Thông tin người nhận tiền:', highlight5
        ws.cell(row, 7).value, ws.cell(row, 7).style = 'Số tài khoản:',highlight12
        ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('so_tk', None),highlight12

        row += 1
        ws.cell(row, 2).value, ws.cell(row, 2).style = 'Tên nhà cung cấp:',highlight12
        ws.cell(row, 3).value, ws.cell(row, 3).style = recs[0].get('ncc', None) if self.repay_type == 'supplier' else recs[0].get('ncc_1', None),highlight12
        ws.cell(row, 7).value, ws.cell(row, 7).style = 'Người thụ hưởng:',highlight12
        ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('nhan_vien', None),highlight12

        row += 1
        ws.cell(row, 2).value, ws.cell(row, 2).style = 'Địa chỉ',highlight12
        ws.cell(row, 7).value, ws.cell(row, 7).style = 'Ngân hàng',highlight12
        ws.cell(row, 8).value, ws.cell(row, 8).style = recs[0].get('ngan_hang', None),highlight12
        self._fill_border_bottom_solid(ws, row, highlight3)

        row += 1
        ws.cell(row, 2).value, ws.cell(row, 2).style = 'Người đề nghị', highlight15
        ws.cell(row, 4).value, ws.cell(row, 4).style = 'Trưởng bộ phận', highlight15
        ws.cell(row, 7).value, ws.cell(row, 7).style = 'Phụ trách kế toán', highlight15
        ws.cell(row, 11).value, ws.cell(row, 11).style = 'Tổng Giám Đốc', highlight15

        row += 4
        ws.cell(row, 2).value, ws.cell(row, 2).style = self.env.user.name, highlight12
        self._fill_border_bottom_solid(ws, row, highlight3)

        if self.license_type == 'internal':
            row += 2
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'Người nhận tiền xác nhận', highlight5

            row += 2
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Số tiền đã nhận: ',highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Chữ kí người nhận tiền:',highlight12

            row += 2
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Ngày nhận tiền:',highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Họ và tên người nhận tiền:',highlight12

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Đề nghị thanh toán.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        return attachment_id

    def _fill_border_bottom_solid(self, ws, row, highlight):
        for i in range(1, 13):
            ws.cell(row, i).style = highlight
