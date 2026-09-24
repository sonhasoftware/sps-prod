import datetime
from datetime import timedelta, date
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side
from openpyxl.utils import get_column_letter
from lxml import etree
from openpyxl.writer.excel import save_virtual_workbook
from odoo.exceptions import ValidationError
from dateutil.relativedelta import relativedelta
from odoo import api, fields, models


class ManhourRate(models.Model):
    _name = 'manhour.rate'
    _description = 'Lương chi tiết'

    year = fields.Selection([(str(x), str(x)) for x in range(2000, 2050)], 'Năm',
                            default=str(date.today().year), required=1)

    def print_reports(self):
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%smanhour_rate_template.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']

        sql = '''
        with all_day_year as
                    (
                        SELECT
                            cast(generate_series as date) as ngay_trong_nam,
                            extract (dow from generate_series) as ngay_trong_tuan
                        FROM
                            GENERATE_SERIES
                                (
                                    cast(concat(cast({year} as varchar),'-01-01') as date),
                                    cast(concat(cast({year} as varchar),'-12-31') as date),
                                    interval '1 day'
                                )
                    ),
                company_other_expenses as
                    (
                        select 
                            oex.hr_job_id,
                            oex.expense_type,
                            oex.gender,
                            sum(coalesce(oex.expense_price,0)) as price
                        from 
                            (
                                select 
                                    hj.id as hr_job_id,
                                    oe."type" as expense_type,
                                    oe.price as expense_price,
                                    oe.gender 
                                from hr_job hj 
                                left join hr_job_other_expenses_rel hjoer on hjoer.hr_job_id = hj.id 
                                left join other_expenses oe on oe.id = hjoer.other_expenses_id
                                where {year} between oe.start_year and oe.end_year and oe.gender is not null 
                                
                                union all 
                                
                                select
                                    hj2.id as hr_job_id,
                                    oe2."type" as expense_type,
                                    oe2.price as expense_price,
                                    oe2.gender 
                                from hr_job hj2 
                                left join other_expenses oe2 on (select count(*) from hr_job_other_expenses_rel hjoer2 where hjoer2.other_expenses_id = oe2.id) = 0
                                where {year} between oe2.start_year and oe2.end_year and oe2.gender is not null
                                
                                union all 
                                
                                select 
                                    hj.id as hr_job_id,
                                    oe."type" as expense_type,
                                    oe.price as expense_price,
                                    'male' as gender 
                                from hr_job hj 
                                left join hr_job_other_expenses_rel hjoer on hjoer.hr_job_id = hj.id 
                                left join other_expenses oe on oe.id = hjoer.other_expenses_id
                                where {year} between oe.start_year and oe.end_year and oe.gender is null 
                                
                                union all 
                                
                                select
                                    hj2.id as hr_job_id,
                                    oe2."type" as expense_type,
                                    oe2.price as expense_price,
                                    'male' as gender 
                                from hr_job hj2 
                                left join other_expenses oe2 on (select count(*) from hr_job_other_expenses_rel hjoer2 where hjoer2.other_expenses_id = oe2.id) = 0
                                where {year} between oe2.start_year and oe2.end_year and oe2.gender is null
                                
                                union all 
                                
                                select 
                                    hj.id as hr_job_id,
                                    oe."type" as expense_type,
                                    oe.price as expense_price,
                                    'female' as gender 
                                from hr_job hj 
                                left join hr_job_other_expenses_rel hjoer on hjoer.hr_job_id = hj.id 
                                left join other_expenses oe on oe.id = hjoer.other_expenses_id
                                where {year} between oe.start_year and oe.end_year and oe.gender is null 
                                
                                union all 
                                
                                select
                                    hj2.id as hr_job_id,
                                    oe2."type" as expense_type,
                                    oe2.price as expense_price,
                                    'female' as gender 
                                from hr_job hj2 
                                left join other_expenses oe2 on (select count(*) from hr_job_other_expenses_rel hjoer2 where hjoer2.other_expenses_id = oe2.id) = 0
                                where {year} between oe2.start_year and oe2.end_year and oe2.gender is null
                                
                                union all 
                                
                                select 
                                    hj.id as hr_job_id,
                                    oe."type" as expense_type,
                                    oe.price as expense_price,
                                    'other' as gender 
                                from hr_job hj 
                                left join hr_job_other_expenses_rel hjoer on hjoer.hr_job_id = hj.id 
                                left join other_expenses oe on oe.id = hjoer.other_expenses_id
                                where {year} between oe.start_year and oe.end_year and oe.gender is null 
                                
                                union all 
                                
                                select
                                    hj2.id as hr_job_id,
                                    oe2."type" as expense_type,
                                    oe2.price as expense_price,
                                    'other' as gender 
                                from hr_job hj2 
                                left join other_expenses oe2 on (select count(*) from hr_job_other_expenses_rel hjoer2 where hjoer2.other_expenses_id = oe2.id) = 0
                                where {year} between oe2.start_year and oe2.end_year and oe2.gender is null
                            ) oex
                        group by oex.hr_job_id, oex.expense_type, oex.gender
                    ),
                detail_data as
                    (
                        select 
                            hdb."name" as block,
                            he."name" as ten_nhan_vien,
                            hj."name" as chuc_danh,
                            he.x_resource_type as vi_tri,
                            case when hpst.wage_type = 'hourly' then 0 else coalesce (hcx.x_insurance_wage,0) end as luong_chinh,
                            case when hpst.wage_type = 'hourly' then 0 else coalesce (hcx.kpi_norm,0) end as luong_an_trua,
                            0::numeric as luong_nha_o,
                            0::numeric as du_kien_tang_luong,
                            case when hpst.wage_type = 'hourly' then coalesce(hcx.hourly_wage)*208 else coalesce (hcx.wage,0) end as luong_tong,
                            case when hpst.wage_type = 'hourly' then 0 else coalesce (hcx.x_insurance_wage,0) * 21.5 / 100 end as bao_hiem,
                            coalesce (hcx.x_allowance_phone,0) as dien_thoai,
                            coalesce (hcx.x_gasoline_standard,0) * coalesce(hpop.price,0) as xang_xe,
                            case when hpst.wage_type = 'hourly' then coalesce(hcx.hourly_wage)*208/12 else coalesce (hcx.wage,0) / 12 end as thuong_13,
                            coalesce (coe.price,0)/12 as dong_phuc,
                            coalesce (coe2.price,0)/12 as kham_sk,
                            coalesce (coe3.price,0)/12 as thuong_le,
                            coalesce (coe4.price,0)/12 as du_lich,
                            case when hpst.wage_type = 'hourly' then 0 else coalesce (hcx.x_insurance_wage,0) * 2 / 100 end as cong_doan,
                            (
                                (select count(ady.ngay_trong_nam) from all_day_year ady) * 8 - (12+11) * 8 --so ngay trong nam tru so ngay nghi phep tru so ngay nghi le
                                - (case 
                                    when rc.full_time_required_hours = 44 then (select count(ady2.ngay_trong_nam) from all_day_year ady2 where ady2.ngay_trong_tuan = 6) * 4 + (select count(ady3.ngay_trong_nam) from all_day_year ady3 where ady3.ngay_trong_tuan = 0) * 8
                                    when rc.full_time_required_hours = 48 then (select count(ady4.ngay_trong_nam) from all_day_year ady4 where ady4.ngay_trong_tuan = 0) * 8
                                    else (select count(ady2.ngay_trong_nam) from all_day_year ady2 where ady2.ngay_trong_tuan = 6) * 4 + (select count(ady3.ngay_trong_nam) from all_day_year ady3 where ady3.ngay_trong_tuan = 0) * 8
                                end) -- tru so ngay thu 7 chu nhat
                            ) / 12 as gio_lv_thang,
                            coalesce(hj.x_labor_efficiency::numeric,0)::numeric as hs_lao_dong,
                            case when hcx.x_type_employee in ('1','3') then 1 else 0 end as is_meal_eligible
                        from hr_employee he
                        left join 
                            (
                                select 
                                    row_number () over (partition by hc.employee_id order by hc.date_start desc) as stt,
                                    * 
                                from hr_contract hc 
                                where hc.state not in ('draft','cancel') and {year} between date_part('year',hc.date_start) and date_part('year',hc.date_end)
                            ) as hcx on hcx.stt = 1 and hcx.employee_id = he.id 
                        left join hr_department hd on hd.id = he.department_id
                        left join hr_department_block hdb on hdb.id = hd.x_block_id 
                        left join hr_job hj on hj.id = he.job_id 
                        left join hr_payroll_structure_type hpst on hpst.id = hcx.structure_type_id  
                        left join 
                            (
                                select 
                                    row_number () over (order by hpopx."month" desc) as stt,
                                    * 
                                from hr_payroll_oil_price hpopx 
                                where hpopx."year"::int = {year}
                            ) as hpop on hpop.stt = 1
                        left join resource_calendar rc on rc.id = hcx.resource_calendar_id 
                        left join company_other_expenses coe on coe.hr_job_id = hj.id and coe.gender = he.gender and coe.expense_type = 'uniform'
                        left join company_other_expenses coe2 on coe2.hr_job_id = hj.id and coe2.gender = he.gender  and coe2.expense_type = 'health'
                        left join company_other_expenses coe3 on coe3.hr_job_id = hj.id and coe3.gender = he.gender  and coe3.expense_type = 'holiday'
                        left join company_other_expenses coe4 on coe4.hr_job_id = hj.id and coe4.gender = he.gender  and coe4.expense_type = 'travel'
                        where 1=1
                            and hcx.id is not null 
                            and he.active is true
                    )
                select
                    dd.block,
                    dd.ten_nhan_vien,
                    dd.chuc_danh,
                    case 
                        when dd.vi_tri = 'seniorengsub' then 'Kỹ sư / Giám sát cao cấp'
                        when dd.vi_tri = 'engsub' then 'Kỹ sư / Giám sát'
                        when dd.vi_tri = 'teamleader' then 'Trưởng nhóm'
                        when dd.vi_tri = 'technician' then 'Kỹ thuật viên'
                        when dd.vi_tri = 'internship' then 'Thực tập sinh'
                    end as vi_tri,
                    dd.luong_chinh,
                    dd.luong_an_trua,
                    dd.luong_nha_o,
                    dd.du_kien_tang_luong,
                    dd.luong_tong,
                    dd.bao_hiem,
                    dd.dien_thoai,
                    dd.xang_xe,
                    dd.thuong_13,
                    dd.dong_phuc,
                    dd.kham_sk,
                    dd.thuong_le,
                    dd.du_lich,
                    dd.cong_doan,
                    dd.bao_hiem+dd.dien_thoai+dd.xang_xe+dd.thuong_13+dd.dong_phuc+dd.kham_sk+dd.thuong_le+dd.du_lich+dd.cong_doan as tong_cp_cty,
                    dd.bao_hiem+dd.dien_thoai+dd.xang_xe+dd.thuong_13+dd.dong_phuc+dd.kham_sk+dd.thuong_le+dd.du_lich+dd.cong_doan+dd.luong_tong
                        + (case when dd.is_meal_eligible = 1 then dd.gio_lv_thang/8*40000 else 0 end) as tong_cong_cp_nv,
                    dd.gio_lv_thang,
                    dd.hs_lao_dong,
                    dd.gio_lv_thang as gio_lv_thang_tt,
                    dd.hs_lao_dong as hs_lao_dong_tt,
                    case
                        when coalesce(dd.gio_lv_thang,0) = 0 or coalesce(dd.hs_lao_dong,0) = 0 then 0
                        else (dd.bao_hiem+dd.dien_thoai+dd.xang_xe+dd.thuong_13+dd.dong_phuc+dd.kham_sk+dd.thuong_le+dd.du_lich+dd.cong_doan+dd.luong_tong
                              + (case when dd.is_meal_eligible = 1 then dd.gio_lv_thang/8*40000 else 0 end)) / (dd.gio_lv_thang*dd.hs_lao_dong)
                    end as don_gia_tb
                from detail_data dd


        '''.format(year=self.year)
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        highlight = NamedStyle(name="highlight")
        highlight.font = Font(bold=True, size=13)
        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dotted', color="000000")
        bd3 = Side(style=None)
        highlight.border = Border(left=bd1, top=bd2, right=bd1, bottom=bd2)

        style_sum1 = NamedStyle(name="style_sum1")
        style_sum1.font = Font(bold=True, size=13)
        style_sum1.border = Border(left=bd3, top=bd1, right=bd3, bottom=bd1)

        style_sum2 = NamedStyle(name="style_sum2")
        style_sum2.font = Font(bold=True, size=13)
        style_sum2.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        style_sum3 = NamedStyle(name="style_sum3")
        style_sum3.font = Font(bold=False, size=13)
        style_sum3.number_format = '#,##0'
        style_sum3.border = Border(left=bd1, top=bd2, right=bd1, bottom=bd2)

        general = []
        facility = []
        maintenenance = []
        for r in recs:
            if r['block'] == 'General':
                general.append(r)
            if r['block'] == 'Facility Management':
                facility.append(r)
            if r['block'] == 'Maintenenance & Services':
                maintenenance.append(r)

        row = 4

        if general:
            row += 1
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'I', style_sum1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'General', style_sum1
            ws.cell(row, 3).style = style_sum1
            ws.cell(row, 4).style = style_sum1
            ws.cell(row, 22).style = style_sum2
            ws.cell(row, 24).style = style_sum2
            for column in range(5, 26):
                if column in [22, 24]:
                    continue
                ws.cell(row, column).value, ws.cell(row, column).style, ws.cell(row, column).number_format = '=SUBTOTAL(9,%s6:%s%s)' % (
                    get_column_letter(column),
                    get_column_letter(column),
                    row + len(general)
                ), style_sum2, '#,##0'
            i = 0
            for r in general:
                row += 1
                i += 1
                ws.cell(row, 1).value = i
                ws.cell(row, 2).value, ws.cell(row, 2).style = r['ten_nhan_vien'] or '', style_sum3
                ws.cell(row, 3).value, ws.cell(row, 3).style = r['chuc_danh'] or '', style_sum3
                ws.cell(row, 4).value, ws.cell(row, 4).style = r['vi_tri'] or '', style_sum3
                ws.cell(row, 5).value, ws.cell(row, 5).style = r['luong_chinh'] or '', style_sum3
                ws.cell(row, 6).value, ws.cell(row, 6).style = r['luong_an_trua'] or '', style_sum3
                ws.cell(row, 7).value, ws.cell(row, 7).style = r['luong_nha_o'] or '', style_sum3
                ws.cell(row, 8).value, ws.cell(row, 8).style = r['du_kien_tang_luong'] or '', style_sum3
                ws.cell(row, 9).value, ws.cell(row, 9).style = r['luong_tong'] or '', style_sum3
                ws.cell(row, 10).value, ws.cell(row, 10).style = r['bao_hiem'] or '', style_sum3
                ws.cell(row, 11).value, ws.cell(row, 11).style = r['dien_thoai'] or '', style_sum3
                ws.cell(row, 12).value, ws.cell(row, 12).style = r['xang_xe'] or '', style_sum3
                ws.cell(row, 13).value, ws.cell(row, 13).style = r['thuong_13'] or '', style_sum3
                ws.cell(row, 14).value, ws.cell(row, 14).style = r['dong_phuc'] or '', style_sum3
                ws.cell(row, 15).value, ws.cell(row, 15).style = r['kham_sk'] or '', style_sum3
                ws.cell(row, 16).value, ws.cell(row, 16).style = r['thuong_le'] or '', style_sum3
                ws.cell(row, 17).value, ws.cell(row, 17).style = r['du_lich'] or '', style_sum3
                ws.cell(row, 18).value, ws.cell(row, 18).style = r['cong_doan'] or '', style_sum3
                ws.cell(row, 19).value, ws.cell(row, 19).style = r['tong_cp_cty'] or '', style_sum3
                ws.cell(row, 20).value, ws.cell(row, 20).style = r['tong_cong_cp_nv'] or '', style_sum3
                ws.cell(row, 21).value, ws.cell(row, 21).style = r['gio_lv_thang'] or '', style_sum3
                ws.cell(row, 22).value, ws.cell(row, 22).style = r['hs_lao_dong'] or '', style_sum3
                ws.cell(row, 23).value, ws.cell(row, 23).style = r['gio_lv_thang_tt'] or '', style_sum3
                ws.cell(row, 24).value, ws.cell(row, 24).style = r['hs_lao_dong_tt'] or '', style_sum3
                ws.cell(row, 25).value, ws.cell(row, 25).style = r['don_gia_tb'] or '', highlight
        ws.cell(2, 11).value = self.year
        if facility:
            row += 1
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'II', style_sum1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Facility Management', style_sum1
            ws.cell(row, 3).style = style_sum1
            ws.cell(row, 4).style = style_sum1
            ws.cell(row, 22).style = style_sum2
            ws.cell(row, 24).style = style_sum2
            for column in range(5, 26):
                if column in [22, 24]:
                    continue
                ws.cell(row, column).value, ws.cell(row, column).style, ws.cell(row, column).number_format = '=SUBTOTAL(9,%s%s:%s%s)' % (
                    get_column_letter(column),
                    row + 1,
                    get_column_letter(column),
                    row + len(facility)
                ), style_sum2, '#,##0'
            i = 0
            for r in facility:
                row += 1
                i += 1
                ws.cell(row, 1).value = i
                ws.cell(row, 2).value, ws.cell(row, 2).style = r['ten_nhan_vien'] or '', style_sum3
                ws.cell(row, 3).value, ws.cell(row, 3).style = r['chuc_danh'] or '', style_sum3
                ws.cell(row, 4).value, ws.cell(row, 4).style = r['vi_tri'] or '', style_sum3
                ws.cell(row, 5).value, ws.cell(row, 5).style = r['luong_chinh'] or '', style_sum3
                ws.cell(row, 6).value, ws.cell(row, 6).style = r['luong_an_trua'] or '', style_sum3
                ws.cell(row, 7).value, ws.cell(row, 7).style = r['luong_nha_o'] or '', style_sum3
                ws.cell(row, 8).value, ws.cell(row, 8).style = r['du_kien_tang_luong'] or '', style_sum3
                ws.cell(row, 9).value, ws.cell(row, 9).style = r['luong_tong'] or '', style_sum3
                ws.cell(row, 10).value, ws.cell(row, 10).style = r['bao_hiem'] or '', style_sum3
                ws.cell(row, 11).value, ws.cell(row, 11).style = r['dien_thoai'] or '', style_sum3
                ws.cell(row, 12).value, ws.cell(row, 12).style = r['xang_xe'] or '', style_sum3
                ws.cell(row, 13).value, ws.cell(row, 13).style = r['thuong_13'] or '', style_sum3
                ws.cell(row, 14).value, ws.cell(row, 14).style = r['dong_phuc'] or '', style_sum3
                ws.cell(row, 15).value, ws.cell(row, 15).style = r['kham_sk'] or '', style_sum3
                ws.cell(row, 16).value, ws.cell(row, 16).style = r['thuong_le'] or '', style_sum3
                ws.cell(row, 17).value, ws.cell(row, 17).style = r['du_lich'] or '', style_sum3
                ws.cell(row, 18).value, ws.cell(row, 18).style = r['cong_doan'] or '', style_sum3
                ws.cell(row, 19).value, ws.cell(row, 19).style = r['tong_cp_cty'] or '', style_sum3
                ws.cell(row, 20).value, ws.cell(row, 20).style = r['tong_cong_cp_nv'] or '', style_sum3
                ws.cell(row, 21).value, ws.cell(row, 21).style = r['gio_lv_thang'] or '', style_sum3
                ws.cell(row, 22).value, ws.cell(row, 22).style = r['hs_lao_dong'] or '', style_sum3
                ws.cell(row, 23).value, ws.cell(row, 23).style = r['gio_lv_thang_tt'] or '', style_sum3
                ws.cell(row, 24).value, ws.cell(row, 24).style = r['hs_lao_dong_tt'] or '', style_sum3
                ws.cell(row, 25).value, ws.cell(row, 25).style = r['don_gia_tb'] or '', style_sum3

        if maintenenance:
            row += 1
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'III', style_sum1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Maintenenance & Services', style_sum1
            ws.cell(row, 3).style = style_sum1
            ws.cell(row, 4).style = style_sum1
            ws.cell(row, 22).style = style_sum2
            ws.cell(row, 24).style = style_sum2
            for column in range(5, 26):
                if column in [22, 24]:
                    continue
                ws.cell(row, column).value, ws.cell(row, column).style, ws.cell(row, column).number_format = '=SUBTOTAL(9,%s%s:%s%s)' % (
                    get_column_letter(column),
                    row + 1,
                    get_column_letter(column),
                    row + len(maintenenance)
                ), style_sum2, '#,##0'
            i = 0
            for r in maintenenance:
                row += 1
                i += 1
                ws.cell(row, 1).value = i
                ws.cell(row, 2).value, ws.cell(row, 2).style = r['ten_nhan_vien'] or '', style_sum3
                ws.cell(row, 3).value, ws.cell(row, 3).style = r['chuc_danh'] or '', style_sum3
                ws.cell(row, 4).value, ws.cell(row, 4).style = r['vi_tri'] or '', style_sum3
                ws.cell(row, 5).value, ws.cell(row, 5).style = r['luong_chinh'] or '', style_sum3
                ws.cell(row, 6).value, ws.cell(row, 6).style = r['luong_an_trua'] or '', style_sum3
                ws.cell(row, 7).value, ws.cell(row, 7).style = r['luong_nha_o'] or '', style_sum3
                ws.cell(row, 8).value, ws.cell(row, 8).style = r['du_kien_tang_luong'] or '', style_sum3
                ws.cell(row, 9).value, ws.cell(row, 9).style = r['luong_tong'] or '', style_sum3
                ws.cell(row, 10).value, ws.cell(row, 10).style = r['bao_hiem'] or '', style_sum3
                ws.cell(row, 11).value, ws.cell(row, 11).style = r['dien_thoai'] or '', style_sum3
                ws.cell(row, 12).value, ws.cell(row, 12).style = r['xang_xe'] or '', style_sum3
                ws.cell(row, 13).value, ws.cell(row, 13).style = r['thuong_13'] or '', style_sum3
                ws.cell(row, 14).value, ws.cell(row, 14).style = r['dong_phuc'] or '', style_sum3
                ws.cell(row, 15).value, ws.cell(row, 15).style = r['kham_sk'] or '', style_sum3
                ws.cell(row, 16).value, ws.cell(row, 16).style = r['thuong_le'] or '', style_sum3
                ws.cell(row, 17).value, ws.cell(row, 17).style = r['du_lich'] or '', style_sum3
                ws.cell(row, 18).value, ws.cell(row, 18).style = r['cong_doan'] or '', style_sum3
                ws.cell(row, 19).value, ws.cell(row, 19).style = r['tong_cp_cty'] or '', style_sum3
                ws.cell(row, 20).value, ws.cell(row, 20).style = r['tong_cong_cp_nv'] or '', style_sum3
                ws.cell(row, 21).value, ws.cell(row, 21).style = r['gio_lv_thang'] or '', style_sum3
                ws.cell(row, 22).value, ws.cell(row, 22).style = r['hs_lao_dong'] or '', style_sum3
                ws.cell(row, 23).value, ws.cell(row, 23).style = r['gio_lv_thang_tt'] or '', style_sum3
                ws.cell(row, 24).value, ws.cell(row, 24).style = r['hs_lao_dong_tt'] or '', style_sum3
                ws.cell(row, 25).value, ws.cell(row, 25).style = r['don_gia_tb'] or '', style_sum3

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Lương chi tiết nhân viên.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        # download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/' + str(attachment_id.id) + '?download=true',
            'target': 'new',
        }
