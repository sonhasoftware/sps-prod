# -*- coding: utf-8 -*-

from odoo import models, fields
import base64
import os
from datetime import date

import openpyxl
from openpyxl.styles import NamedStyle, Font, Side, PatternFill, Border, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from six import BytesIO


class PaymentOrder(models.Model):
    _inherit = 'account.move'
    _description = 'Đề nghị thanh toán'

    def get_job_description(self):
        self.ensure_one()
        name_user = self.env.user.name
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%spayment_order.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']
        sql = f'''
                SELECT
                    coalesce(prj.name, so.name) ma_duan,
                    coalesce(prj.label_tasks, so.work_content) ten_duan,
                    rp.NAME khach_hang,
                    aml.x_number_bill so_hd,
                    aml.x_date_bill ngay_hd,
                    pt.name noi_dung,
                    it.value don_vi,
                    aml.quantity so_luong,
                    aml.price_unit don_gia,
                    case when aml.price_subtotal is not null then aml.price_subtotal else 0 end bf_tax,
                    aml.price_total - aml.price_subtotal amount_tax,
                    case when aml.price_total is not null then aml.price_total else 0 end at_tax,
                    am.x_note ghi_chu,
                    case when coalesce(am.amount_total,0)-coalesce(am.amount_residual,0) != 0 then coalesce(am.amount_total)-coalesce(am.amount_residual) else 0 end tam_ung,
                    rp.NAME ten_ncc,
	                concat(rp.street, case when rp.street is not null and rp.city is not null then '-' else '' end ,rp.city, case when rp.city is not null and rcs.name is not null then '-' else '' end ,rcs.name) dia_chi,
                    rpb.acc_number stk,
                    rpb.acc_holder_name nguoi_thu_huong,
                    rb.NAME ngan_hang,
                    ct.cost_category phan_loai,
                    ct.cost_group nhom
                FROM
                    account_move am
                    LEFT JOIN account_move_line aml ON am.ID = aml.move_id
                    LEFT JOIN sale_order so ON aml.x_sale_project_id = so.ID
                    LEFT JOIN project_project prj ON aml.x_project_id = prj.ID
                    LEFT JOIN res_partner rp ON rp.id = am.partner_id
                    LEFT JOIN uom_uom uu ON aml.product_uom_id = uu.ID 
                    -- LEFT JOIN res_partner rp1 ON am.partner_id = rp1.ID 
                    LEFT JOIN res_partner_bank rpb ON rpb.id = am.partner_bank_id 
                    LEFT JOIN res_bank rb ON rb.ID = rpb.bank_id
                    LEFT JOIN cost_type ct ON aml.x_cost_type_id = ct.ID
                    LEFT JOIN res_country_state rcs on rp.state_id = rcs.id
                    left join product_product pp on pp.id = aml.product_id 
                    left join product_template pt on pt.id = pp.product_tmpl_id 
                    left join ir_translation it on it.name = 'uom.uom,name' and it.lang = 'vi_VN' and it.res_id = aml.product_uom_id 
                WHERE
                    am.ID = {self.id}
                    AND aml.exclude_from_invoice_tab != 't'
                ORDER BY ct.cost_group , coalesce(prj.name, so.name)
            '''
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dashed', color="000000")
        bd3 = Side(style='mediumDashed', color="0000FF")
        bd4 = Side(style='thick', color="0000FF")

        highlight = NamedStyle(name="highlight")
        highlight.font = Font(name='Times New Roman', size=13, bold=True)
        highlight.border = Border(bottom=bd1, top=bd1)

        highlight1 = NamedStyle(name="highlight1")
        highlight1.font = Font(name='Times New Roman', size=13)
        highlight1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight1.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
        highlight1.number_format = '#,##0'

        highlight2 = NamedStyle(name="highlight2")
        highlight2.font = Font(name='Times New Roman', size=13, bold=True)
        highlight2.border = Border(bottom=bd2)

        highlight3 = NamedStyle(name="highlight3")
        highlight3.font = Font(name='Times New Roman', size=13)
        highlight3.border = Border(bottom=bd1)

        highlight4 = NamedStyle(name="highlight4")
        highlight4.font = Font(name='Times New Roman', size=13)
        highlight4.border = Border(bottom=bd3)

        highlight5 = NamedStyle(name="highlight5")
        highlight5.font = Font(name='Times New Roman', size=13, underline="single", bold=True)
        highlight5.alignment = Alignment(vertical='center')

        highlight6 = NamedStyle(name="highlight6")
        highlight6.font = Font(name='Times New Roman', size=13)
        highlight6.border = Border(bottom=bd4)

        highlight7 = NamedStyle(name="highlight7")
        highlight7.font = Font(name='Times New Roman', size=13)
        highlight7.border = Border(right=bd4)

        highlight8 = NamedStyle(name="highlight8")
        highlight8.font = Font(name='Times New Roman', size=13)
        highlight8.alignment = Alignment(vertical='center', wrap_text=True)

        highlight9 = NamedStyle(name="highlight9")
        highlight9.font = Font(name='Times New Roman', size=13)
        highlight9.alignment = Alignment(vertical='center', horizontal='center')

        highlight10 = NamedStyle(name="highlight10")
        highlight10.font = Font(name='Times New Roman', size=13, bold=True)
        highlight10.alignment = Alignment(vertical='center', horizontal='center')
        highlight10.number_format = '#,##0'

        highlight11 = NamedStyle(name="highlight11")
        highlight11.font = Font(name='Times New Roman', size=13, bold=True)
        highlight11.alignment = Alignment(vertical='center', horizontal='center')
        highlight11.border = Border(bottom=bd1, top=bd1)
        highlight11.number_format = '#,##0'

        highlight12 = NamedStyle(name="highlight12")
        highlight12.font = Font(name='Times New Roman', size=13)

        highlight13 = NamedStyle(name="highlight13")
        highlight13.font = Font(name='Times New Roman', size=13)
        highlight13.alignment = Alignment(vertical='center')
        highlight13.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight14 = NamedStyle(name="highlight14")
        highlight14.font = Font(name='Times New Roman', size=13)
        highlight14.border = Border(bottom=bd2)

        highlight17 = NamedStyle(name="highlight17")
        highlight17.font = Font(name='Times New Roman', size=13)
        highlight17.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight17.alignment = Alignment(horizontal='center', vertical='center', wrapText=True)
        highlight17.number_format = 'DD/MM/YYYY'
        highlight15 = NamedStyle(name="highlight15")
        highlight15.font = Font(name='Times New Roman', size=13, bold=True)
        highlight15.alignment = Alignment(vertical='center', horizontal='center')

        highlight16 = NamedStyle(name="highlight16")
        highlight16.font = Font(name='Times New Roman', size=13, bold=True, underline="single")
        highlight16.alignment = Alignment(vertical='center')
        highlight16.border = Border(bottom=bd2)

        dict_item = {}
        for rec in recs:
            key = rec['nhom'] + rec['ma_duan'] if rec['nhom'] else rec['ma_duan']
            if key not in dict_item:
                dict_item[key] = [rec]
            else:
                dict_item[key].append(rec)

        row = 6
        x = 1
        bf_tax = 0
        amount_tax = 0
        at_tax = 0
        tax_dispose = 0
        at_dispose = 0
        for name, values in dict_item.items():
            ws.row_dimensions[row].height = 30
            ws.cell(row, 1).value, ws.cell(row, 1).style = values[0]['nhom'] if values[0]['nhom'] else 'CHI PHÍ CHƯA PHÂN LOẠI', highlight5
            ws.cell(row, 5).value, ws.cell(row, 5).style = 'MÃ DỰ ÁN:', highlight9
            ws.cell(row, 6).value, ws.cell(row, 6).style = values[0]['ma_duan'], highlight8
            ws.cell(row, 8).value, ws.cell(row, 8).style = 'TÊN DỰ ÁN:', highlight9
            ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=12)
            ws.cell(row, 9).value, ws.cell(row, 9).style = values[0]['ten_duan'], highlight8
            tax_dispose = self.amount_tax - sum([total['amount_tax'] for total in values]) if len(dict_item)<=1 else 0
            at_dispose = self.amount_total - sum([total['at_tax'] for total in values]) if len(dict_item)<=1 else 0
            row += 1
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'KHÁCH HÀNG:', highlight9
            ws.cell(row, 3).value, ws.cell(row, 3).style = values[0]['khach_hang'], highlight8
            ws.cell(row, 8).value, ws.cell(row, 8).style = 'TỔNG SỐ', highlight9
            ws.cell(row, 9).value, ws.cell(row, 9).style = sum([total['bf_tax'] for total in values]), highlight10
            ws.cell(row, 10).value, ws.cell(row, 10).style = sum([total['amount_tax'] for total in values]) if tax_dispose<= 0 else self.amount_tax, highlight10
            ws.cell(row, 11).value, ws.cell(row, 11).style = sum([total['at_tax'] for total in values]) if at_dispose <= 0 else self.amount_total, highlight10
            row += 1
            for value in values:
                ws.cell(row, 1).value, ws.cell(row, 1).style = x, highlight1
                ws.cell(row, 2).value, ws.cell(row, 2).style = value['so_hd'], highlight1
                ws.cell(row, 3).value, ws.cell(row, 3).style = value['ngay_hd'], highlight17
                ws.cell(row, 4).value, ws.cell(row, 4).style = value['noi_dung'], highlight13
                ws.cell(row, 5).value, ws.cell(row, 5).style = value['phan_loai'], highlight1
                ws.cell(row, 6).value, ws.cell(row, 6).style = value['don_vi'], highlight1
                ws.cell(row, 7).value, ws.cell(row, 7).style = value['so_luong'], highlight1
                ws.cell(row, 8).value, ws.cell(row, 8).style = value['don_gia'], highlight1
                ws.cell(row, 9).value, ws.cell(row, 9).style = value['bf_tax'], highlight1
                ws.cell(row, 10).value, ws.cell(row, 10).style = value['amount_tax']if x != values.__len__() else value['amount_tax']+ tax_dispose, highlight1
                ws.cell(row, 11).value, ws.cell(row, 11).style = value['at_tax'] if x != values.__len__() else value['at_tax']+ at_dispose, highlight1
                ws.cell(row, 12).value, ws.cell(row, 12).style = value['ghi_chu'], highlight1
                bf_tax += value['bf_tax']
                amount_tax += value['amount_tax'] if x != values.__len__() else value['amount_tax']+ tax_dispose
                at_tax += value['at_tax'] if x != values.__len__() else value['at_tax']+ at_dispose
                row += 1
                x += 1
            x = 1
        ws.cell(4, 3).value, ws.cell(4, 3).style = self.env.user.name, highlight12
        self._fill_border_bottom_solid(ws, row, highlight)
        ws.cell(row, 1).value = 'Tổng số tiền (VNĐ)'
        ws.cell(row, 9).value, ws.cell(row, 9).style = bf_tax, highlight11
        ws.cell(row, 10).value, ws.cell(row, 10).style = amount_tax, highlight11
        ws.cell(row, 11).value, ws.cell(row, 11).style = at_tax, highlight11

        row += 1
        self._fill_border_bottom_solid(ws, row, highlight)
        ws.cell(row, 1).value = 'Số tiền đã tạm ứng (VNĐ)'
        ws.cell(row, 11).value, ws.cell(row, 11).style = recs[0].get('tam_ung', 0), highlight11

        row += 1
        self._fill_border_bottom_solid(ws, row, highlight)
        ws.cell(row, 1).value = 'Số tiền công ty phải thanh toán'
        ws.cell(row, 11).value, ws.cell(row, 11).style = at_tax - recs[0].get('tam_ung', 0), highlight11

        row += 1
        self._fill_border_bottom_solid(ws, row, highlight2)
        ws.cell(row, 1).value, ws.cell(row, 1).style = 'Phương thức thanh toán:', highlight16
        ws.cell(row, 4).value, ws.cell(row, 4).style = '☐ ' + 'Tiền mặt / Chuyển khoản cá nhân', highlight14
        ws.cell(row, 5).value, ws.cell(row, 5).style = '☐ ' + "CK Techcombank", highlight14
        ws.cell(row, 7).value, ws.cell(row, 7).style = '☐ ' + "CK Vietcombank", highlight14
        ws.cell(row, 9).value, ws.cell(row, 9).style = '☐ ' + "CK VPBank", highlight14
        ws.cell(row, 11).value, ws.cell(row, 11).style = '☐ ' + "CK BIDV", highlight14

        row += 1
        ws.cell(row, 1).value, ws.cell(row, 1).style = 'Thông tin người nhận tiền:', highlight5
        ws.cell(row, 7).value,ws.cell(row, 7).style = 'Số tài khoản:',highlight12
        ws.cell(row, 8).value,ws.cell(row, 8).style = recs[0].get('stk', None),highlight12

        row += 1
        ws.cell(row, 1).value,ws.cell(row, 1).style = 'Tên nhà cung cấp:',highlight12
        ws.cell(row, 3).value,ws.cell(row, 3).style = recs[0].get('ten_ncc', None),highlight12
        ws.cell(row, 7).value,ws.cell(row, 7).style = 'Người thụ hưởng:',highlight12
        ws.cell(row, 8).value,ws.cell(row, 8).style = recs[0].get('nguoi_thu_huong', None),highlight12

        row += 1
        ws.cell(row, 1).value,ws.cell(row, 1).style = 'Địa chỉ',highlight12

        ws.cell(row, 7).value,ws.cell(row, 7).style = 'Ngân hàng',highlight12
        ws.cell(row, 8).value,ws.cell(row, 8).style = recs[0].get('ngan_hang', None),highlight12
        self._fill_border_bottom_solid(ws, row, highlight3)

        row += 1
        ws.cell(row, 2).value, ws.cell(row, 2).style = 'Người đề nghị', highlight15
        ws.cell(row, 4).value, ws.cell(row, 4).style = 'Trưởng bộ phận', highlight15
        ws.cell(row, 7).value, ws.cell(row, 7).style = 'Phụ trách kế toán', highlight15
        ws.cell(row, 11).value, ws.cell(row, 11).style = 'Tổng Giám Đốc', highlight15

        row += 4
        ws.cell(row, 2).value, ws.cell(row, 2).style = self.env.user.name, highlight12
        self._fill_border_bottom_solid(ws, row, highlight3)

        if self.x_license_type == 'internal':
            row += 2
            ws.cell(row, 1).value, ws.cell(row, 1).style = 'Người nhận tiền xác nhận', highlight5

            row += 2
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Số tiền đã nhận: ', highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Chữ kí người nhận tiền:', highlight12

            row += 2
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'Ngày nhận tiền:', highlight12
            ws.cell(row, 7).value, ws.cell(row, 7).style = 'Họ và tên người nhận tiền:', highlight12

        if self.x_license_type == 'tax':
            ws.cell(3, 1).value = 'GIẤY ĐỀ NGHỊ THANH TOÁN '
        else:
            ws.cell(3, 1).value = 'GIẤY ĐỀ NGHỊ THANH TOÁN / TẠM ỨNG KIÊM PHIẾU CHI (NỘI BỘ)'

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Đề nghị thanh toán.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        return attachment_id

    def _fill_border_bottom_solid(self, ws, row, highlight):
        for i in range(1, 13):
            ws.cell(row, i).style = highlight
