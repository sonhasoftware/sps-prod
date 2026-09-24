# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, PatternFill
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.worksheet import filters
from odoo import api, fields, models
from openpyxl.styles import Alignment


class CashFlowReport(models.Model):
    _name = 'cash.flow.popup'
    _description = "Báo cáo dòng tiền"

    # year = fields.Integer('Year', default=fields.Date.today().year)

    def action_cash_flow(self):
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%scash-flow-statement.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['DATA']
        ws2 = wb['Saving Cash']
        sql = '''
            select 
                ngay,
                nguoi_nhan,
                ten_khach_hang,
                dien_giai,
                so3.name as ma_du_an,
                chung_tu,
                ma_khoan_tien,
                kieu_thanh_toan,
                chi,
                thu,
                ghi_chu,
                so_chung_tu
            from 
                (select 
                    amx.date as ngay,
                    case when ru.id is not null then coalesce(rp.name,'') else coalesce(ap.x_receiver,'') end as nguoi_nhan,
                    case when ru.id is not null then coalesce(rc.name,'') else coalesce(rp.name,rc.name) end as ten_khach_hang,
                    coalesce(amx.ref,'') as dien_giai,
                    case	
                        when aa.id is not null and pd.id is not null then pd.code_project  
                        when aa.id is not null and pd.id is null and aa."type" = 'supplier' then po_rate.so_id
                        when aar.id is not null then prd.main_project 
                        when am.id is not null and pd.id is not null then pd.code_project 
                        when am.id is not null and pd.id is null then am.x_order_id 
                        when so.id is not null then so.id 
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and pd.id is not null and ap.is_internal_transfer is false and ap.x_pay_cost is false then pd.code_project 
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and ap.is_internal_transfer is false and ap.x_pay_cost is true then pl.x_sale_project_id  
                        else null
                    end as ma_du_an,
                    coalesce(ap.x_license_type,'') as chung_tu,
                    coalesce(cm.code_money,'') as ma_khoan_tien,
                    coalesce(aj.name,'') as kieu_thanh_toan,
                    case	
                        when aa.id is not null and pd.id is not null then (case when pdc.pd_count = 1 and ap.amount > pd.amount_total then ap.amount else pd.amount_total end) 
                        when aa.id is not null and pd.id is null and aa."type" = 'supplier' then po_rate.pol_so_rate*ap.amount  
                        when aa.id is not null and pd.id is null and aa."type" = 'internal' then ap.amount 
                        when aar.id is not null and ap.payment_type = 'outbound' then prd.amount_repay 
                        when am.id is not null and pd.id is not null and ap.payment_type = 'outbound' then (case when pdc.pd_count = 1 and ap.amount > pd.amount_total then ap.amount else pd.amount_total end)
                        when am.id is not null and pd.id is null and ap.payment_type = 'outbound' then ap.amount  
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and pd.id is not null and ap.x_pay_cost is false and ap.payment_type = 'outbound' then (case when pdc.pd_count = 1 and ap.amount > pd.amount_total then ap.amount else pd.amount_total end)
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and pd.id is null and ap.x_pay_cost is false and ap.payment_type = 'outbound' then ap.amount 
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and ap.x_pay_cost is true and ap.payment_type = 'outbound' then pl.amount
                        else 0
                    end as chi,
                    case	
                        when aar.id is not null and ap.payment_type = 'inbound' then prd.amount_repay 
                        when am.id is not null and pd.id is not null and ap.payment_type = 'inbound' then (case when pdc.pd_count = 1 and ap.amount > pd.amount_total then ap.amount else pd.amount_total end)
                        when am.id is not null and pd.id is null and ap.payment_type = 'inbound' then ap.amount 
                        when so.id is not null then ap.amount 
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and pd.id is not null and ap.x_pay_cost is false and ap.payment_type = 'inbound' then (case when pdc.pd_count = 1 and ap.amount > pd.amount_total then ap.amount else pd.amount_total end)
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and pd.id is null and ap.x_pay_cost is false and ap.payment_type = 'inbound' then ap.amount 
                        when aa.id is null and aar.id is null and am.id is null and so.id is null and ap.x_pay_cost is true and ap.payment_type = 'inbound' then pl.amount 
                        else 0
                    end as thu,
                    coalesce(ap.x_note,'') as ghi_chu,
                    coalesce(amx.name,'') as so_chung_tu
                from account_payment ap 
                left join account_move amx on amx.payment_id = ap.id 
                left join account_journal aj on aj.id = amx.journal_id 
                left join res_partner rp on ap.partner_id = rp.id
                left join res_users ru on ap.partner_id = ru.partner_id
                left join res_company rc on 1 = 1 and rc.id=1
                left join code_money cm on cm.id = ap.x_code_money 
                left join project_detail pd on pd.payment_id = ap.id and ap.is_internal_transfer is false and ap.x_pay_cost is false
                left join (
                    select payment_id, count(*) as pd_count
                    from project_detail
                    where payment_id is not null
                    group by payment_id
                ) as pdc on pdc.payment_id = ap.id
                left join account_advance aa on aa.id = ap.x_origin_advance_id 
                left join 
                    (
                        select 
                            po.id as po_id,
                            so2.id as so_id,
                            case when po.amount_total = 0 then 0 else sum(pol.price_total)/po.amount_total end as pol_so_rate
                        from purchase_order po 
                        left join purchase_order_line pol on pol.order_id = po.id
                        left join project_project pp on pp.id = pol.x_project_id 
                        left join sale_order so2 on so2.id = pp.x_order_id 
                        group by po.id,so2.id,po.amount_total
                    ) as po_rate on ap.x_origin_advance_id is not null and pd.id is null and aa.type = 'supplier' and aa.po_id = po_rate.po_id
                left join account_advance_repay aar on aar.id = ap.x_origin_repay_id 
                left join project_repay_detail prd on prd.payment_id = ap.id and ap.x_origin_repay_id is not null
                left join account_move am on am.id = ap.x_origin_move_id 
                left join sale_order so on so.id = ap.x_origin_so_id 
                left join payment_line pl on pl.payment_id = ap.id and ap.x_pay_cost is true and aa.id is null and aar.id is null and am.id is null and so.id is null
                where amx.state = 'posted'
                )X
            left join sale_order so3 on so3.id = X.ma_du_an
            order by X.ngay,X.so_chung_tu
        '''
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()

        highlight = NamedStyle(name="highlight")
        highlight.font = Font(bold=True, size=13)
        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dotted', color="000000")
        bd3 = Side(style=None)
        highlight.border = Border(left=bd1, top=bd2, right=bd1, bottom=bd2)

        style_sum1 = NamedStyle(name="style_sum1")
        style_sum1.font = Font(bold=True, size=13)
        style_sum1.border = Border(left=bd3, top=bd1, right=bd3, bottom=bd1)

        style_sum2 = NamedStyle(name="style_sum2")
        style_sum2.font = Font(bold=True, size=13)
        style_sum2.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight17 = NamedStyle(name="highlight17")
        highlight17.font = Font(name='Times New Roman', size=10)
        highlight17.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        highlight17.alignment = Alignment(horizontal='center', vertical='center')
        highlight17.number_format = '[$-en-US]dd-mmm-yyyy;@'

        style_sum3 = NamedStyle(name="style_sum3")
        style_sum3.font = Font(bold=False, size=10)
        style_sum3.number_format = '#,##0'
        style_sum3.alignment = Alignment(wrap_text=True)
        style_sum3.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        row = 5

        for r in recs:
            ws.cell(row, 1).value, ws.cell(row, 1).style = r['ngay'], highlight17
            ws.cell(row, 2).value, ws.cell(row, 2).style = r['nguoi_nhan'], style_sum3
            ws.cell(row, 3).value, ws.cell(row, 3).style = r['ten_khach_hang'], style_sum3
            ws.cell(row, 4).value, ws.cell(row, 4).style = r['dien_giai'], style_sum3
            ws.cell(row, 5).value, ws.cell(row, 5).style = r['ma_du_an'], style_sum3
            ws.cell(row, 6).value, ws.cell(row, 6).style = r['chung_tu'], style_sum3
            ws.cell(row, 7).value, ws.cell(row, 7).style = r['ma_khoan_tien'], style_sum3
            ws.cell(row, 8).value, ws.cell(row, 8).style = r['kieu_thanh_toan'], style_sum3
            ws.cell(row, 9).value, ws.cell(row, 9).style = r['thu'], style_sum3
            ws.cell(row, 10).value, ws.cell(row, 10).style = r['chi'], style_sum3
            ws.cell(row, 11).value, ws.cell(row, 11).style = r['ghi_chu'], style_sum3
            ws.cell(row, 12).value, ws.cell(row, 12).style = r['so_chung_tu'], style_sum3

            row += 1

        last_data_row = row - 1
        ws_total_font = Font(bold=True, name='Times New Roman', size=10)
        ws_total_border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        ws.cell(row, 9).value = '=SUBTOTAL(9,I5:I%d)' % last_data_row
        ws.cell(row, 9).font = ws_total_font
        ws.cell(row, 9).border = ws_total_border
        ws.cell(row, 9).number_format = '#,##0'
        ws.cell(row, 10).value = '=SUBTOTAL(9,J5:J%d)' % last_data_row
        ws.cell(row, 10).font = ws_total_font
        ws.cell(row, 10).border = ws_total_border
        ws.cell(row, 10).number_format = '#,##0'

        sql2 = '''
            SELECT av.date_sent,
                av.content,
                av.number_contract,
                av.value_deposits,
                av.interest_rate,
                av.date_settlement,
                av.amount_settlement,
                rb.name as ngan_hang,
                av.note
            FROM account_saving av
            LEFT JOIN res_bank rb on av.bank_id = rb.id
            WHERE av.state != 'draft'
            ORDER BY av.date_sent
            '''
        self._cr.execute(sql2)
        recs2 = self._cr.dictfetchall()

        row2 = 5

        for r in recs2:
            ws2.cell(row2, 1).value, ws2.cell(row2, 1).style = r['date_sent'], highlight17
            ws2.cell(row2, 2).value, ws2.cell(row2, 2).style = r['content'], style_sum3
            ws2.cell(row2, 3).value, ws2.cell(row2, 3).style = r['number_contract'], style_sum3
            ws2.cell(row2, 4).value, ws2.cell(row2, 4).style = r['value_deposits'], style_sum3
            ws2.cell(row2, 5).value, ws2.cell(row2, 5).style = r['interest_rate'], style_sum3
            ws2.cell(row2, 5).number_format = '#,##0.00'
            ws2.cell(row2, 6).value, ws2.cell(row2, 6).style = r['date_settlement'], highlight17
            ws2.cell(row2, 7).value, ws2.cell(row2, 7).style = r['amount_settlement'], style_sum3
            ws2.cell(row2, 8).value, ws2.cell(row2, 8).style = r['ngan_hang'], style_sum3
            ws2.cell(row2, 9).value, ws2.cell(row2, 9).style = r['note'], style_sum3

            row2 += 1

        total_font = Font(color='FF0000', bold=True, name='Times New Roman', size=10)
        total_fill = PatternFill(start_color='ADD8E6', end_color='ADD8E6', fill_type='solid')
        total_border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        for col in range(1, 10):
            cell = ws2.cell(row2, col)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = total_border
        ws2.cell(row2, 2).value = 'CỘNG'
        ws2.cell(row2, 4).value = '=SUMIF(F:F,">="&TODAY(),D:D)'
        ws2.cell(row2, 4).number_format = '#,##0'
        ws2.cell(row2, 7).value = '=SUMIF(F:F,">="&TODAY(),G:G)'
        ws2.cell(row2, 7).number_format = '#,##0'

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Báo cáo dòng tiền Cash flow statement.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        # download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/' + str(attachment_id.id) + '?download=true',
            'target': 'new',
        }

    def action_view_cash_flow(self):
        action = self.env["ir.actions.act_window"]._for_xml_id('account_saving.action_cash_flow_report')
        return action