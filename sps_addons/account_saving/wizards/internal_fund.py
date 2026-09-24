# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from odoo import models, fields
from dateutil.relativedelta import relativedelta


class InternalFundPopup(models.Model):
    _name = 'internal.fund.popup'
    _description = 'Báo cáo thu chi quỹ nội bộ'

    account_id = fields.Many2one('account.account.type', string='Tài khoản',
                                 domain="[('type','=','liquidity'),('internal_group','=','asset')]")
    master_key = fields.Integer('Master Key', default=lambda self: self.env.uid)

    def action_internal_fund_popup(self):
        domain_acc = ''
        if self.account_id:
            domain_acc = 'and aat.id = ' + str(self.account_id.id)
        else:
            domain_acc = ''
        year = date.today()
        last_year = year - relativedelta(years=1)

        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%sso_quy_tien_mat.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Data']
        sql = '''
                SELECT
                    am.NAME so_hieu,
                    am.DATE AS ngay,
                    ap.x_receiver nguoi_nhan_nop,
                    rp.NAME ten,
                    am.REF AS dien_giai,
                    pp.NAME ma_du_an,
                    ap.x_license_type AS phan_loai_chung_tu,
                    cm.code_money ma_khoan_tien,
                    aj.code,
                    case when ap.payment_type ='inbound' then ap.amount else 0 end AS thu,
                    case when ap.payment_type ='outbound' then ap.amount else 0 end AS chi
                FROM
                    account_payment ap
                    LEFT JOIN account_move am ON am.ID = ap.move_id
                    LEFT JOIN account_journal aj ON am.journal_id = aj.
                    ID LEFT JOIN account_move_line aml ON aml.move_id = am.
                    ID LEFT JOIN project_project pp ON aml.x_sale_project_id = pp.
                    ID LEFT JOIN account_account aa ON aa.ID = aml.account_id
                    LEFT JOIN res_partner rp ON rp.ID = ap.partner_id
                    LEFT JOIN code_money cm on cm.id = ap.x_code_money
                    INNER JOIN account_account_type aat ON aat.ID = aa.user_type_id 
                            and aat.type = 'liquidity' 
                            and aat.internal_group = 'asset'
                WHERE
                    ap.is_internal_transfer = 't' 
                    AND am.DATE BETWEEN '{last_year}' 
                    AND '{year}' 
                    {domain_acc}
                GROUP BY
                    am.DATE,
                    am.NAME,
                    aj.code,
                    ap.partner_type,
                    ap.amount,
                    am.REF,
                    pp.NAME,
                    rp.NAME,
                    ap.x_license_type,
                    ap.payment_type,
                    ap.x_receiver,
                    cm.code_money
        '''.format(year=year, last_year=last_year, domain_acc=domain_acc)
        self._cr.execute(sql)
        recs = self._cr.dictfetchall()
        highlight = NamedStyle(name="highlight")
        highlight.font = Font(size=13)
        bd1 = Side(style='thin', color="000000")
        highlight.alignment = Alignment(wrap_text=True)
        highlight.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        highlight1 = NamedStyle(name='datetime', number_format='dd/Mon/yy')
        highlight1.font = Font(size=13)
        highlight1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        style_sum1 = NamedStyle(name="style_sum1")
        style_sum1.font = Font(size=15, bold=True)
        style_sum1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)

        amount_thu = 0
        amount_chi = 0
        row = 5
        for r in recs:
            ws.cell(row, 1).value, ws.cell(row, 1).style = r['so_hieu'], highlight
            ws.cell(row, 2).value, ws.cell(row, 2).style = r['ngay'], highlight1
            ws.cell(row, 3).value, ws.cell(row, 3).style = r['nguoi_nhan_nop'], highlight
            ws.cell(row, 4).value, ws.cell(row, 4).style = r['ten'], highlight
            ws.cell(row, 5).value, ws.cell(row, 5).style = r['dien_giai'], highlight
            ws.cell(row, 6).value, ws.cell(row, 6).style = r['ma_du_an'], highlight
            ws.cell(row, 7).value, ws.cell(row, 7).style = r['phan_loai_chung_tu'], highlight
            ws.cell(row, 8).value, ws.cell(row, 8).style = r['ma_khoan_tien'], highlight
            ws.cell(row, 9).value, ws.cell(row, 9).style = r['code'], highlight
            ws.cell(row, 10).value, ws.cell(row, 10).style = r['thu'], highlight
            amount_thu += r['thu']
            ws.cell(row, 11).value, ws.cell(row, 11).style = r['chi'], highlight
            amount_chi += r['chi']
            row += 1
        ws.cell(1, 11).value = 'Prepared by '
        ws.cell(1, 12).value = self.env.user.name
        if len(recs) > 0:
            ws.cell(row, 2).value, ws.cell(row, 2).style = 'TỔNG CỘNG', style_sum1
            ws.cell(row, 10).value, ws.cell(row, 10).style = amount_thu, style_sum1
            ws.cell(row, 11).value, ws.cell(row, 11).style = amount_chi, style_sum1
            ws.cell(row, 12).value, ws.cell(row, 12).style = amount_thu - amount_chi, style_sum1

        ws1 = wb['Report']
        ws1.cell(1, 16).value = 'Prepared by '
        ws1.cell(1, 17).value = self.env.user.name

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Báo cáo thu chi quỹ nội bộ.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        # download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/' + str(attachment_id.id) + '?download=true',
            'target': 'new',
        }
