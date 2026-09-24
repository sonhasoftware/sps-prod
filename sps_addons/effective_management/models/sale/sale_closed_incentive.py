# -*- coding: utf-8 -*-
import os
from io import BytesIO
import openpyxl
from odoo import api, fields, models
from openpyxl.writer.excel import save_virtual_workbook
from odoo.modules import get_module_path
from odoo.http import Response, content_disposition

class SaleClosedIncentive(models.Model):
    _name = 'sale.closed.incentive'
    _description = 'Thống kê báo giá đã ký'

    master_key = fields.Integer('Master Key', index=True)
    classification = fields.Char('Phân loại')
    year = fields.Float()
    t1 = fields.Float('Jan')
    t2 = fields.Float('Feb')
    t3 = fields.Float('Mar')
    t4 = fields.Float('Apr')
    t5 = fields.Float('May')
    t6 = fields.Float('Jun')
    t7 = fields.Float('Jul')
    t8 = fields.Float('Aug')
    t9 = fields.Float('Sep')
    t10 = fields.Float('Oct')
    t11 = fields.Float('Nov')
    t12 = fields.Float('Dec')
    # total được LƯU (stored) và ghi trực tiếp bởi wizard action_report, vì với
    # dòng "Giá trị thưởng dự kiến", Total = tổng đại số CẢ NĂM × tỉ lệ, KHÁC
    # tổng 12 tháng (tháng lỗ bị chặn về 0). Không còn tính = t1+...+t12.
    total = fields.Float('Total')
    light = fields.Boolean('bôi đậm', compute='compute_light')

    @api.depends('classification')
    def compute_light(self):
        value = 'Giá trị thưởng dự kiến'
        for r in self:
            if value in r.classification :
                r.light = True
            else:
                r.light = False

    def export_excel_report(self, domain=None, context=None):
        domain = domain or []
        records = self.search(domain, order='id')
        year_value = records and records[0].year or fields.Date.today().year
        try:
            year = int(year_value)
        except (TypeError, ValueError):
            year = fields.Date.today().year

        # Chọn SO hoàn thành trong năm — NGUỒN CHUNG (đồng nhất báo cáo hiệu
        # quả dự án chung + thưởng ký). Trả [{sale_order_id, sale_order_name,
        # completion_date}].
        rows = self.env['sale.order']._get_completed_orders_with_date(year)

        order_ids = [r['sale_order_id'] for r in rows]
        sale_orders = self.env['sale.order'].sudo().browse(order_ids).exists()
        sale_order_map = {so.id: so for so in sale_orders}

        profit_by_order = {so.id: so.profit_after_tax or 0.0 for so in sale_orders}
        for row in rows:
            so_id = row['sale_order_id']
            row['profit_after_tax_total'] = profit_by_order.get(so_id, 0.0)

        user_ids = {
            uid
            for so in sale_orders
            for uid in (so.user_id.id, so.solution_maker.id if so.solution_maker else False)
            if uid
        }
        bod_user_ids = set()
        if user_ids:
            bod_employees = self.env['hr.employee'].sudo().search([
                ('user_id', 'in', list(user_ids)),
                ('department_id', '=', 4),
            ])
            bod_user_ids = {emp.user_id.id for emp in bod_employees if emp.user_id}

        module_path = get_module_path('effective_management')
        excel_path = os.path.join(module_path, 'templates', 'template_so_tinh_thuong.xlsx')
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        ws = wb['Value']

        start_row = 2
        write_row = start_row
        for data in rows:
            sale_order = sale_order_map.get(data['sale_order_id'])
            if not sale_order:
                continue
            solution_maker = sale_order.solution_maker
            key_account = sale_order.user_id
            if (
                solution_maker and solution_maker.id in bod_user_ids
                and key_account and key_account.id in bod_user_ids
            ):
                continue

            row_index = write_row
            ws.cell(row=row_index, column=1, value=data.get('sale_order_name') or '')
            ws.cell(row=row_index, column=2, value=data.get('profit_after_tax_total') or 0.0)
            ws.cell(row=row_index, column=3, value=solution_maker.name if solution_maker else '')
            ws.cell(row=row_index, column=4, value=key_account.name if key_account else '')
            completion_date = data.get('completion_date')
            if completion_date:
                completion_str = fields.Date.to_string(completion_date)
            else:
                completion_str = ''
            ws.cell(row=row_index, column=5, value=sale_order.project_type or '')
            ws.cell(row=row_index, column=6, value=completion_str)
            write_row += 1

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'Danh sách BG tính thưởng {year}.xlsx'

        # Return response
        response = Response(
            output.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename))
            ]
        )

        return response
