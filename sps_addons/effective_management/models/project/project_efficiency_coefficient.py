# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from docutils.nodes import target
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side
from openpyxl.writer.excel import save_virtual_workbook
from odoo import api, fields, models
import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, PatternFill, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.worksheet import filters
from odoo.modules import get_module_path
from odoo.http import content_disposition, dispatch_rpc, request, serialize_exception as _serialize_exception, Response


class EfficiencyCoefficient(models.Model):
    _name = 'project.efficiency.coefficient'
    _description = 'Báo cáo hiệu xuất lao động'

    master_key = fields.Integer('Master Key', index=True)
    classification = fields.Char('Phân loại')
    resource_type = fields.Char('Loại nguồn lực')
    t1 = fields.Float('Jan (%)')
    t2 = fields.Float('Feb (%)')
    t3 = fields.Float('Mar (%)')
    t4 = fields.Float('Apr (%)')
    t5 = fields.Float('May (%)')
    t6 = fields.Float('Jun (%)')
    t7 = fields.Float('Jul (%)')
    t8 = fields.Float('Aug (%)')
    t9 = fields.Float('Sep (%)')
    t10 = fields.Float('Oct (%)')
    t11 = fields.Float('Nov (%)')
    t12 = fields.Float('Dec (%)')
    total = fields.Float('Tổng cộng (%)')
    hight_light = fields.Boolean('bôi đậm', compute = 'compute_hight_light')
    employee_id = fields.Many2one('hr.employee')
    target = fields.Float('Mục tiêu (%)')
    ratio = fields.Float(compute = 'compute_ratio',string = "Tỷ lệ(%)")
    year = fields.Integer()
    # def compute_amount(self):
    #     for r in self:
    #         r.total = r.t1 + r.t2 + r.t3 + r.t4 + r.t5 + r.t6 + r.t7 + r.t8 + r.t9 + r.t10 + r.t11 + r.t12
    def compute_ratio(self):
        for rec in self:
            rec.ratio = 100 - rec.target + rec.total


    @api.depends('classification')
    def compute_hight_light(self):
        for r in self :
            if r.classification in ('Tổng giá trị ngân sách','Tổng giá trị đã mua','Tổng hiệu quả mua hàng','Người lập dự toán','Người đàm phán'):
                r.hight_light = True
            else:
                r.hight_light =False

    def export_project_efficiency_coefficient(self,domain,context):
        module_path = get_module_path('effective_management')
        excel_path = os.path.join(module_path, 'templates', 'Template_bao_cao_hieu_suat_lao_dong.xlsx')
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        ws = wb['Coefficient']
        project_efficiency_coefficient_ids = self.search(domain, order="resource_type")

        start_row = 5
        current_row = start_row
        bold_font = Font(bold=True)
        grouped_data = {}
        if project_efficiency_coefficient_ids:
            ws.cell(row=3, column=7, value=project_efficiency_coefficient_ids[0].year)
        for record in project_efficiency_coefficient_ids:
            if record.resource_type not in grouped_data:
                grouped_data[record.resource_type] = []
            grouped_data[record.resource_type].append(record)
        
        for resource_type, records in grouped_data.items():
            resource_type_cell = ws.cell(row=current_row, column=1, value=resource_type or '')
            resource_type_cell.font = bold_font
            current_row += 1
            
            for record in records:
                ws.cell(row=current_row, column=1, value=record.classification or '')
                ws.cell(row=current_row, column=2, value=record.t1/100 if record.t1 else 0)
                ws.cell(row=current_row, column=3, value=record.t2/100 if record.t2 else 0)
                ws.cell(row=current_row, column=4, value=record.t3/100 if record.t3 else 0)
                ws.cell(row=current_row, column=5, value=record.t4/100 if record.t4 else 0)
                ws.cell(row=current_row, column=6, value=record.t5/100 if record.t5 else 0)
                ws.cell(row=current_row, column=7, value=record.t6/100 if record.t6 else 0)
                ws.cell(row=current_row, column=8, value=record.t7/100 if record.t7 else 0)
                ws.cell(row=current_row, column=9, value=record.t8/100 if record.t8 else 0)
                ws.cell(row=current_row, column=10, value=record.t9/100 if record.t9 else 0)
                ws.cell(row=current_row, column=11, value=record.t10/100 if record.t10 else 0)
                ws.cell(row=current_row, column=12, value=record.t11/100 if record.t11 else 0)
                ws.cell(row=current_row, column=13, value=record.t12/100 if record.t12 else 0)
                ws.cell(row=current_row, column=14, value=record.total/100 if record.total else 0)
                ws.cell(row=current_row, column=15, value=record.target/100)
                ws.cell(row=current_row, column=17, value=record.ratio/100)


                current_row += 1

        end_row = current_row - 1
        standard_font = Font(name='Calibri', size=11, color="000000")  # Font chữ chung
        center_alignment = Alignment(horizontal='center', vertical='center')  # Căn giữa
        
        # Apply format cho tất cả các cells từ start_row đến end_row
        for row in range(start_row, end_row + 1):
            for col in range(1, 15):  # Từ cột A đến N
                cell = ws.cell(row=row, column=col)
                cell.font = standard_font
                cell.alignment = center_alignment
        
        # Giữ lại font đậm cho resource_type headers
        current_row_temp = start_row
        for resource_type, records in grouped_data.items():
            resource_type_cell = ws.cell(row=current_row_temp, column=1)
            resource_type_cell.font = Font(name='Calibri', size=11, color="000000", bold=True)
            current_row_temp += len(records) + 1

        # Lưu workbook vào BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Tạo filename
        filename = f'Bao_cao_hieu_suat_lao_dong_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        
        # Return response
        response = Response(
            output.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename))
            ]
        )
        
        return response
