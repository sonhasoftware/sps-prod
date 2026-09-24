# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from odoo import api, fields, models
from odoo.modules import get_module_path
from odoo.http import Response, content_disposition

class SumValueLabor(models.Model):
    _name = 'project.sum.value.labor'
    _description = 'BÁO CÁO THỐNG KÊ THEO TỔNG GIÁ TRỊ NGÂN SÁCH'

    master_key = fields.Integer('Master Key', index=True)
    classification = fields.Char('Quản lý dự án')
    t1 = fields.Float('Jan')
    t2 = fields.Float('Feb')
    t3 = fields.Float('Mar')
    t4 = fields.Float('Apr')
    t5 = fields.Float('May')
    t6 = fields.Float('Jun')
    t7 = fields.Float('Jul')
    t8 = fields.Float('Aug')
    t9 = fields.Float('Sep')
    t10 = fields.Float('Oct')
    t11 = fields.Float('Nov')
    t12 = fields.Float('Dec')
    total = fields.Float('Tổng cộng', compute='compute_amount')
    hight_light = fields.Boolean('bôi đậm', compute = 'compute_hight_light')
    year = fields.Integer('Năm')

    def compute_amount(self):
        for r in self:
            r.total = r.t1 + r.t2 + r.t3 + r.t4 + r.t5 + r.t6 + r.t7 + r.t8 + r.t9 + r.t10 + r.t11 + r.t12

    @api.depends('classification')
    def compute_hight_light(self):
        for r in self :
            if r.classification in ('Tổng giá trị ngân sách','Tổng giá trị đã mua','Tổng hiệu quả mua hàng','Người lập dự toán','Người đàm phán'):
                r.hight_light = True
            else:
                r.hight_light =False

    def export_excel_report(self, domain, context):
        module_path = get_module_path('effective_management')
        excel_path = os.path.join(module_path, 'templates', 'template_tong_gia_tri_ngan_sach_nhan_luc.xlsx')
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        ws = wb['Value']
        value_labor_ids = self.search(domain)

        ws.cell(row=2, column=8, value=value_labor_ids[0].year or 0)

        start_row = 5
        current_row = start_row
        stt = 0

        # Tạo style viền và font
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        tahoma_font = Font(name='Tahoma', size=11)

        for record in value_labor_ids:
            stt += 1
            # Căn giữa cột 1
            cell_col1 = ws.cell(row=current_row, column=1, value=stt or '')
            cell_col1.alignment = Alignment(horizontal='center')
            cell_col1.border = thin_border
            cell_col1.font = tahoma_font

            cell = ws.cell(row=current_row, column=2, value=record.classification or '')
            cell.border = thin_border
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=3, value=record.t1 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=4, value=record.t2 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=5, value=record.t3 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=6, value=record.t4 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=7, value=record.t5 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=8, value=record.t6 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=9, value=record.t7 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=10, value=record.t8 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=11, value=record.t9 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=12, value=record.t10 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=13, value=record.t11 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            cell = ws.cell(row=current_row, column=14, value=record.t12 or 0)
            cell.border = thin_border
            cell.number_format = '#,##0'
            cell.font = tahoma_font

            # Tô đậm cột 15
            cell_col15 = ws.cell(row=current_row, column=15, value=record.total or 0)
            cell_col15.font = Font(name='Tahoma', size=11, bold=True)
            cell_col15.border = thin_border
            cell_col15.number_format = '#,##0'

            current_row += 1

        # Thêm dòng tổng cộng với công thức SUM
        tahoma_font_bold = Font(name='Tahoma', size=11, bold=True)

        cell = ws.cell(row=current_row, column=1, value='')
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=2, value='Tổng cộng')
        cell.border = thin_border
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=3, value=f'=SUM(C{start_row}:C{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=4, value=f'=SUM(D{start_row}:D{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=5, value=f'=SUM(E{start_row}:E{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=6, value=f'=SUM(F{start_row}:F{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=7, value=f'=SUM(G{start_row}:G{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=8, value=f'=SUM(H{start_row}:H{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=9, value=f'=SUM(I{start_row}:I{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=10, value=f'=SUM(J{start_row}:J{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=11, value=f'=SUM(K{start_row}:K{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=12, value=f'=SUM(L{start_row}:L{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=13, value=f'=SUM(M{start_row}:M{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell = ws.cell(row=current_row, column=14, value=f'=SUM(N{start_row}:N{current_row-1})')
        cell.border = thin_border
        cell.number_format = '#,##0'
        cell.font = tahoma_font_bold

        cell_total = ws.cell(row=current_row, column=15, value=f'=SUM(O{start_row}:O{current_row-1})')
        cell_total.font = tahoma_font_bold
        cell_total.border = thin_border
        cell_total.number_format = '#,##0'

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # Tạo filename
        filename = f'Tổng giá trị ngân sách nhân lực.xlsx'

        # Return response
        response = Response(
            output.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename))
            ]
        )

        return response