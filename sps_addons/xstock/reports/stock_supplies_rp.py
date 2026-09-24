# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from odoo import api, fields, models
from odoo.modules import get_module_path
from odoo.http import content_disposition, dispatch_rpc, request, serialize_exception as _serialize_exception, Response


class StockSupplies(models.Model):
    _name = 'stock.supplies.rp'
    _description = 'Báo cáo hiệu suất tái sử dụng vật tư'

    master_key = fields.Integer('Master Key', index=True)
    index = fields.Integer('STT')
    login = fields.Char('Short name')
    name = fields.Char('Full name')
    t1 = fields.Integer('Tháng 1')
    t2 = fields.Integer('Tháng 2')
    t3 = fields.Integer('Tháng 3')
    t4 = fields.Integer('Tháng 4')
    t5 = fields.Integer('Tháng 5')
    t6 = fields.Integer('Tháng 6')
    t7 = fields.Integer('Tháng 7')
    t8 = fields.Integer('Tháng 8')
    t9 = fields.Integer('Tháng 9')
    t10 = fields.Integer('Tháng 10')
    t11 = fields.Integer('Tháng 11')
    t12 = fields.Integer('Tháng 12')
    total = fields.Integer('Tổng cộng', compute='compute_amount')

    def compute_amount(self):
        for r in self:
            r.total = r.t1 + r.t2 + r.t3 + r.t4 + r.t5 + r.t6 + r.t7 + r.t8 + r.t9 + r.t10 + r.t11 + r.t12

    def export_excel_report(self, domain, context):
        module_path = get_module_path('xstock')
        excel_path = os.path.join(module_path, 'templates', 'reuse.xlsx')
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        ws = wb['Reuse']
        resue_ids = self.search(domain)
        re = resue_ids[2]
        resue_ids -= resue_ids[2]
        resue_ids += re
        # Define styles for center alignment, border, and formatting
        center_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        bold_font = Font(bold=True)
        number_format = '#,##0'  # Number format with thousands separator

        start_row = 5
        current_row = start_row
        stt = 0
        ws.cell(row=3, column=3, value=2025)

        total_records = len(resue_ids)
        for idx, record in enumerate(resue_ids):
            is_last_row = (idx == total_records - 1)

            if not is_last_row:
                stt += 1

            # Apply formatting to each cell (excluding columns 2 and 3 from center alignment)
            for col in range(1, 17):  # Columns 1 to 16
                cell = ws.cell(row=current_row, column=col)
                if col not in [2, 3]:  # Don't center align columns 2 and 3
                    cell.alignment = center_alignment
                if col >= 4:  # Apply number format to columns 4-16
                    cell.number_format = number_format
                cell.border = thin_border

                # Make last row bold
                if is_last_row:
                    cell.font = bold_font

            # For last row, don't show STT
            if is_last_row:
                ws.cell(row=current_row, column=1, value='')
            else:
                ws.cell(row=current_row, column=1, value=stt or '')

            ws.cell(row=current_row, column=2, value=record.login or '')
            ws.cell(row=current_row, column=3, value=record.name or '')
            ws.cell(row=current_row, column=4, value=record.t1 or '')
            ws.cell(row=current_row, column=5, value=record.t2 or '')
            ws.cell(row=current_row, column=6, value=record.t3 or '')
            ws.cell(row=current_row, column=7, value=record.t4 or '')
            ws.cell(row=current_row, column=8, value=record.t5 or '')
            ws.cell(row=current_row, column=9, value=record.t6 or '0')
            ws.cell(row=current_row, column=10, value=record.t7 or '0')
            ws.cell(row=current_row, column=11, value=record.t8 or '0')
            ws.cell(row=current_row, column=12, value=record.t9 or '0')
            ws.cell(row=current_row, column=13, value=record.t10 or '0')
            ws.cell(row=current_row, column=14, value=record.t11 or '0')
            ws.cell(row=current_row, column=15, value=record.t12 or '0')
            ws.cell(row=current_row, column=16, value=f'=sum(D{current_row}:O{current_row})')
            current_row += 1



        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # Tạo filename
        filename = f'Bao_cao_tai_su_dung.xlsx'

        # Return response
        response = Response(
            output.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename))
            ]
        )

        return response