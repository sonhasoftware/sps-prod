# -*- coding: utf-8 -*-
import math
import base64
from datetime import datetime, date
from dateutil.relativedelta import relativedelta

from odoo import api, fields, models, _
from odoo.exceptions import ValidationError
from odoo.exceptions import UserError
from odoo.osv import expression
import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.worksheet import filters
from openpyxl.styles import Alignment
from odoo.modules import get_module_path



class StockPicking(models.Model):
    _inherit = 'stock.picking'

    x_type_5_state = fields.Selection([
        ('wait_payer', 'Chờ người bàn giao xác nhận'),
        ('wait_receiver', 'Chờ người nhận xác nhận'),
        ('confirmed', 'Đã được xác nhận'),
    ], string='Trạng thái xác nhận',track_visibility="always",readonly=True,default = 'wait_payer')

    x_receiver_id = fields.Many2one('res.users', string='Người nhận/Người được bàn giao',
                                    default=lambda self: self.env.user)
    x_payer_id = fields.Many2one('res.users', string='Người trả/Người bàn giao', default=lambda self: self.env.user)
    x_paid = fields.Boolean('Đã trả', default=False)
    x_picking_type = fields.Selection([
        ('type_1', 'Nhập kho từ đơn mua'),
        ('type_2', 'Nhập lại vật tư tiêu chuẩn'),
        ('type_3', 'Xuất kho'),
        ('type_4', 'Mượn công cụ dụng cụ'),
        ('type_5', 'Chuyển giao công cụ dụng cụ'),
        ('type_6', 'Thu hồi/Hoàn trả công cụ dụng cụ'),
        ('type_7', 'Chuyển kho hàng hóa'),
        ('type_8', 'Thu hồi vật tư thừa'),
        ('type_9', 'Báo hỏng, hủy'),
    ], string='Kiểu hoạt động', related='picking_type_id.x_type')
    is_receiver = fields.Boolean('là người nhận', compute='compute_nonstore_data')
    x_check = fields.Boolean("check", compute='compute_nonstore_data')
    x_to_the_project_id = fields.Many2one('project.project', string='Tới dự án')
    x_type_code = fields.Selection([('incoming', 'Receipt'), ('outgoing', 'Delivery'), ('internal', 'Internal Transfer')], 'Type of Operation', reaonly=True ,related='picking_type_id.code')
    x_valuation_layer = fields.Float('Định giá', compute='_compute_valuation_layer')

    @api.depends('move_ids_without_package.x_valuation_layer')
    def _compute_valuation_layer(self):
        for rec in self:
            rec.x_valuation_layer = sum(abs(x.x_valuation_layer) for x in rec.move_ids_without_package)


    def compute_nonstore_data(self):
        user = self.env.uid
        for r in self:
            if r.x_picking_type == 'type_5' and r.x_receiver_id.id == user:
                r.is_receiver = True
            else:
                r.is_receiver = False
            if self.user_has_groups('stock.group_stock_manager'):
                r.x_check = True
            else:
                r.x_check = False

    def act_payer_confirm(self):
        for r in self:
            if r.x_payer_id.id != self.env.uid:
                raise UserError('Bạn không phải là người bàn giao nên không đc phép xác nhận phiếu này !')
            if r.state == 'assigned':
                r.state = 'confirmed'
                r.x_type_5_state = 'wait_receiver'
            else:
                continue

    def act_receiver_confirm(self):
        for r in self:
            if r.state == 'confirmed':
                r.x_type_5_state = 'confirmed'
            else:
                continue

    def act_refuse(self):
        for r in self:
            if r.state in ['confirmed', 'assigned']:
                r.state = 'assigned'
                r.x_type_5_state = 'wait_payer'
            else:
                continue
        return

    def unlink(self):
        for r in self:
            if r.state == "draft":
                rtn = super(StockPicking, self).unlink()
                return rtn
            else:
                raise UserError("Chỉ có thể xoá bản ghi ở trạng thái nháp")

    def _ccdc_affected_projects(self):
        """Các dự án bị ảnh hưởng bởi phiếu CCDC này (để rebuild lại khấu hao)."""
        self.ensure_one()
        projects = self.env['project.project']
        for line in self.move_line_ids_without_package:
            if line.x_to_the_project_id:
                projects |= line.x_to_the_project_id
            if line.x_from_the_project_id:
                projects |= line.x_from_the_project_id
        return projects

    def _check_damage_requirements(self):
        """Chặn validate phiếu Báo hỏng, mất CCDC (sequence_code='HM') nếu có
        dòng CCDC THẬT SỰ hỏng/mất (đích kho hỏng/hủy, payer khác "HM - Hao
        mòn", chưa bị trả lại) mà chưa chọn "Từ dự án" — không có nơi nào
        khác ghi nhận chi phí này nếu thiếu (project.x_damage_cost_total tính
        LIVE trực tiếp từ sml.x_from_the_project_id, không có bước sync)."""
        self.ensure_one()
        if self.picking_type_id.sequence_code != 'HM':
            return
        if self.x_payer_id and self.x_payer_id.login == 'HM':
            return  # hao mòn tự nhiên, không cần quy về dự án
        for line in self.move_line_ids_without_package:
            if line.state != 'done' or not line.qty_done:
                continue
            tmpl = line.product_id.product_tmpl_id
            if tmpl.x_type != 'product' or tmpl.x_product_type != 'tools':
                continue  # vật tư, không phải CCDC
            if line.move_id.origin_returned_move_id:
                continue  # chính nó là dòng trả lại
            if self.env['stock.move'].search(
                    [('origin_returned_move_id', '=', line.move_id.id)], limit=1):
                continue  # đã bị trả lại sau đó
            if (line.location_dest_id.x_name or '') != 'Kho hàng hỏng, hủy':
                continue  # không thực sự nhập kho hỏng/hủy
            if not line.x_from_the_project_id:
                raise UserError(_(
                    'Sản phẩm "%s" báo hỏng/mất nhưng chưa chọn "Từ dự án" để '
                    'ghi nhận chi phí.') % line.product_id.display_name)

    def button_validate(self):
        for r in self:
            if r.x_type_code in ('outgoing', 'internal'):
                if any(round(l.x_qty_available,1) < l.qty_done for l in r.move_line_ids_without_package):
                    raise UserError(' Có sản phẩm không đủ số lượng tồn kho , vui lòng kiểm tra lại')
        res = super(StockPicking, self).button_validate()
        for r in self:
            if r.x_picking_type not in ['type_4', 'type_5', 'type_6'] or r.state != 'done':
                continue
            # mượn CCDC
            elif r.x_picking_type == 'type_4':
                for line in r.move_line_ids_without_package:
                    if any(or_line.id != line.id and line.product_id.id == or_line.product_id.id and line.lot_id.id == or_line.lot_id.id and
                           line.x_to_the_project_id.id == or_line.x_to_the_project_id.id for or_line in r.move_line_ids_without_package ):
                        raise UserError("Không thể tạo các dòng giống nhau , vui lòng kiểm tra lại")
                    sql = f''' select * from stock_borrow_tools sbt where
                         sbt.product_id = {line.product_id.id} and
                         (sbt.lot_id = {line.lot_id.id if line.lot_id.id else 'NULL'} or sbt.lot_id  is NULL) and
                         sbt.receiver_id = {r.x_receiver_id.id} and
                         sbt.project_id = {line.x_to_the_project_id.id}'''
                    self._cr.execute(sql)
                    res = self._cr.dictfetchall()

                    if len(res) == 0:

                        insert = f'''INSERT INTO stock_borrow_tools (product_id, lot_id, receiver_id, project_id, amount, value)
                                        VALUES ({line.product_id.id},{line.lot_id.id if line.lot_id.id else 'NULL'},
                                                {r.x_receiver_id.id},{line.x_to_the_project_id.id},
                                                {line.qty_done},{line.qty_done * line.move_id.price_unit})'''
                        self._cr.execute(insert)
                    else:
                        update = f'''UPDATE stock_borrow_tools
                                    SET amount = amount + {line.qty_done},
                                    value = value + {line.qty_done * line.move_id.price_unit}
                                    WHERE product_id = {line.product_id.id} and receiver_id = {r.x_receiver_id.id}
                                    and project_id = {line.x_to_the_project_id.id}'''
                        if line.lot_id.id:
                            update += f''' and lot_id = {line.lot_id.id}'''
                        self._cr.execute(update)
                    # Mở dòng khấu hao CCDC theo dự án nhận (1 dòng/lot, gộp qty cho lot rỗng).
                    self.env['project.depreciation.line'].open_line(
                        project=line.x_to_the_project_id,
                        product=line.product_id,
                        lot=line.lot_id,
                        qty=line.qty_done,
                        receiver=r.x_receiver_id,
                        picking=r,
                        date_borrow=r.date_done or fields.Datetime.now(),
                    )
            # thu hồi trả CCDC
            elif r.x_picking_type == 'type_6':
                for line in r.move_line_ids_without_package:
                    if any(or_line.id != line.id and line.product_id.id == or_line.product_id.id and line.lot_id.id == or_line.lot_id.id and
                           line.x_from_the_project_id.id == or_line.x_from_the_project_id.id for or_line in r.move_line_ids_without_package ):
                        raise UserError("Không thể tạo các dòng giống nhau , vui lòng kiểm tra lại")
                    sql = f''' select * from stock_borrow_tools sbt where
                         sbt.product_id = {line.product_id.id} and 
                         (sbt.lot_id = {line.lot_id.id if line.lot_id.id else 'NULL'} or sbt.lot_id  is NULL) and 
                         sbt.receiver_id = {r.x_payer_id.id} and
                         sbt.project_id = {line.x_from_the_project_id.id}'''
                    self._cr.execute(sql)
                    res = self._cr.dictfetchall()

                    if len(res) == 0:
                        raise UserError('Bạn không có dụng cụ để trả / thu hồi')
                    elif r.picking_type_id.default_location_src_id == r.location_id and r.picking_type_id.default_location_dest_id == r.location_dest_id:
                        update = f'''UPDATE stock_borrow_tools
                                    SET amount = amount - {line.qty_done},
                                    value = value - {line.qty_done * line.move_id.price_unit}
                                    WHERE product_id = {line.product_id.id}
                                    and receiver_id = {r.x_payer_id.id}
                                    and project_id = {line.x_from_the_project_id.id}  '''
                        if line.lot_id.id:
                            update += f'''and lot_id = {line.lot_id.id}'''
                        self._cr.execute(update)
                        # Ghi nhận trả khấu hao CCDC tương ứng (thu hồi/trả/báo hỏng-mất).
                        self.env['project.depreciation.line'].register_return(
                            project=line.x_from_the_project_id,
                            product=line.product_id,
                            lot=line.lot_id,
                            qty=line.qty_done,
                            picking_in=r,
                            date_return=r.date_done or fields.Datetime.now(),
                        )
                    elif r.picking_type_id.default_location_src_id == r.location_dest_id and r.picking_type_id.default_location_dest_id == r.location_id:
                        update = f'''UPDATE stock_borrow_tools
                                                            SET amount = amount + {line.qty_done},
                                                            value = value + {line.qty_done * line.move_id.price_unit}
                                                            WHERE product_id = {line.product_id.id}
                                                            and receiver_id = {r.x_payer_id.id}
                                                            and project_id = {line.x_to_the_project_id.id}  '''
                        if line.lot_id.id:
                            update += f'''and lot_id = {line.lot_id.id}'''
                        self._cr.execute(update)
                        # Cấp lại ra dự án (tăng số dư) → mở dòng khấu hao tương ứng.
                        self.env['project.depreciation.line'].open_line(
                            project=line.x_to_the_project_id,
                            product=line.product_id,
                            lot=line.lot_id,
                            qty=line.qty_done,
                            receiver=r.x_payer_id,
                            picking=r,
                            date_borrow=r.date_done or fields.Datetime.now(),
                        )


            # chuyển giao CCDC
            elif r.x_picking_type == 'type_5':
                for line in r.move_line_ids_without_package:
                    if any(
                            or_line.id != line.id and line.product_id.id == or_line.product_id.id and line.lot_id.id == or_line.lot_id.id and
                            line.x_from_the_project_id.id == or_line.x_from_the_project_id.id and line.x_to_the_project_id.id == or_line.x_to_the_project_id.id for or_line in
                            r.move_line_ids_without_package):
                        raise UserError("Không thể tạo các dòng giống nhau , vui lòng kiểm tra lại")
                    sql_tra = f''' select * from stock_borrow_tools sbt where
                             sbt.product_id = {line.product_id.id} and 
                             (sbt.lot_id = {line.lot_id.id if line.lot_id.id else 'NULL'} or sbt.lot_id  is NULL) and 
                             sbt.receiver_id = {r.x_payer_id.id} and
                             sbt.project_id = {line.x_from_the_project_id.id}'''

                    self._cr.execute(sql_tra)
                    res = self._cr.dictfetchall()

                    if len(res) == 0:
                        raise UserError('Bạn không có dụng cụ để bàn giao')
                    else:
                        update_tra = f'''UPDATE stock_borrow_tools 
                               SET amount = amount - {line.qty_done}, value = value - {line.qty_done * line.move_id.price_unit}
                               WHERE product_id = {line.product_id.id} 
                                     and receiver_id = {r.x_payer_id.id} 
                                     and project_id = {line.x_from_the_project_id.id} 
                                     and (lot_id = {line.lot_id.id if line.lot_id.id else 'NULL'} or lot_id  is NULL)'''
                        self._cr.execute(update_tra)

                        sql_nhan = f''' select * from stock_borrow_tools sbt where
                                     sbt.product_id = {line.product_id.id} and 
                                     (sbt.lot_id = {line.lot_id.id if line.lot_id.id else 'NULL'} or sbt.lot_id  is NULL) and 
                                     sbt.receiver_id = {r.x_receiver_id.id} and
                                     sbt.project_id = {line.x_to_the_project_id.id}'''

                        self._cr.execute(sql_nhan)
                        res = self._cr.dictfetchall()

                    if len(res) == 0:

                        insert = f'''INSERT INTO stock_borrow_tools (product_id, lot_id, receiver_id, project_id, amount, value)
                                                            VALUES ({line.product_id.id},
                                                            {line.lot_id.id if line.lot_id.id else 'NULL'},
                                                            {r.x_receiver_id.id},
                                                            {line.x_to_the_project_id.id},
                                                            {line.qty_done},
                                                            {line.qty_done * line.move_id.price_unit})'''
                        self._cr.execute(insert)
                    else:
                        update_nhan = f'''UPDATE stock_borrow_tools
                                                        SET amount = amount + {line.qty_done},
                                                        value = value + {line.qty_done * line.move_id.price_unit}
                                                        WHERE product_id = {line.product_id.id}
                                                        and receiver_id = {r.x_receiver_id.id}
                                                        and project_id = {line.x_to_the_project_id.id} '''
                        if line.lot_id.id:
                            update_nhan += f'''and lot_id = {line.lot_id.id}'''
                        self._cr.execute(update_nhan)
                    # Khấu hao CCDC: ghi nhận trả ở dự án cũ, mở dòng ở dự án mới.
                    DL = self.env['project.depreciation.line']
                    DL.register_return(
                        project=line.x_from_the_project_id,
                        product=line.product_id,
                        lot=line.lot_id,
                        qty=line.qty_done,
                        picking_in=r,
                        date_return=r.date_done or fields.Datetime.now(),
                    )
                    DL.open_line(
                        project=line.x_to_the_project_id,
                        product=line.product_id,
                        lot=line.lot_id,
                        qty=line.qty_done,
                        receiver=r.x_receiver_id,
                        picking=r,
                        date_borrow=r.date_done or fields.Datetime.now(),
                    )
        # Chi phí hỏng/mất CCDC (project.x_damage_cost_total) tính LIVE trực
        # tiếp từ lịch sử phiếu HM, không materialize gì ở đây — chỉ cần chặn
        # validate nếu thiếu "Từ dự án" cho dòng CCDC thật sự hỏng/mất.
        for r in self:
            if r.state == 'done':
                r._check_damage_requirements()
        return res

    @api.onchange('x_payer_id')
    def onchange_payer(self):
        for rec in self:
            rec.move_line_ids_without_package = False

    @api.onchange('move_line_ids_without_package')
    def onchange_project(self):
        self.ensure_one()
        if self.move_line_ids_without_package and self.x_picking_type in ('type_3', 'type_4'):
            line = self.move_line_ids_without_package[len(self.move_line_ids_without_package) - 1]
            if line.x_to_the_project_id:
                self.x_to_the_project_id = line.x_to_the_project_id
        else:
            self.x_to_the_project_id = None

    def write(self, vals):
        # Đồng bộ ngày khấu hao CCDC khi SỬA date_done của phiếu CCDC đã hoàn tất.
        # Thu thập dự án liên quan TRƯỚC super (để so ngày cũ), rebuild SAU super.
        # Bỏ qua lần set đầu lúc validate (date_done: False -> giá trị) để không
        # rebuild trùng với open_line/register_return trong button_validate.
        sync_projects = self.env['project.project']
        if 'date_done' in vals:
            new_dd = fields.Datetime.to_datetime(vals.get('date_done'))
            for r in self:
                if (r.x_picking_type in ('type_4', 'type_5', 'type_6')
                        and r.state == 'done' and r.date_done
                        and r.date_done != new_dd):
                    sync_projects |= r._ccdc_affected_projects()
        res = super(StockPicking, self).write(vals)
        for r in self:
            if r.state in ('confirmed', 'assigned') and not len(r.move_line_ids_without_package):
                self.write({'state': 'draft'})
            if r.state == 'done' and not self.env.user.has_group('stock.group_stock_manager'):
                raise UserError('Phiếu đã được hoàn thành ! Bạn không được chỉnh sửa phiếu này')
        if sync_projects:
            DL = self.env['project.depreciation.line']
            for project in sync_projects:
                DL._rebuild_for_project(project, preserve_snapshot=True)
        # Chi phí hỏng/mất CCDC tính LIVE — đổi date_done hay huỷ phiếu tự động
        # phản ánh đúng vào lần đọc project.x_damage_cost_total tiếp theo,
        # không cần đồng bộ gì ở đây.
        return res

    def action_report_payment_receipt_warehouse(self):
        """Generate Excel report for stock picking"""
        module_path = get_module_path('xstock')
        excel_path = os.path.join(module_path, 'templates', 'Phieu_in.xlsx')
        
        def safe_write_cell(worksheet, row, col, value):
            """Safely write to a cell, handling merged cells"""
            try:
                cell = worksheet.cell(row=row, column=col)
                # Check if it's a merged cell
                for merged_range in worksheet.merged_cells.ranges:
                    if merged_range.min_row <= row <= merged_range.max_row and \
                       merged_range.min_col <= col <= merged_range.max_col:
                        # Write to the top-left cell of the merged range
                        top_left_cell = worksheet.cell(row=merged_range.min_row, column=merged_range.min_col)
                        top_left_cell.value = value
                        return
                # If not merged, write directly
                cell.value = value
            except Exception:
                pass  # Skip if can't write
        
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=False)
            ws_a4 = wb['A4']
            ws_a5 = wb['A5']
            move_lines = self.move_line_ids_without_package

            start_row_a4 = 7
            start_row_a5 = 6

            for index, line in enumerate(move_lines):
                row_a4 = start_row_a4 + index
                row_a5 = start_row_a5 + index
                project_name = line.x_to_the_project_id.name or ''
                product_name = line.product_id.name or ''
                if line.lot_id:
                    product_name += f" (S/N: {line.lot_id.name})"
                date_done = self.date_done.strftime('%d/%m/%Y') if self.date_done else ''

                # Ghi an toàn vào sheet A4
                safe_write_cell(ws_a4, row_a4, 1, index + 1)
                safe_write_cell(ws_a4, row_a4, 3, project_name)
                safe_write_cell(ws_a4, row_a4, 4, product_name)
                safe_write_cell(ws_a4, row_a4, 5, line.qty_done)
                safe_write_cell(ws_a4, row_a4, 6, line.product_uom_id.name or '')
                safe_write_cell(ws_a4, row_a4, 7, date_done)

                # Ghi an toàn vào sheet A5
                safe_write_cell(ws_a5, row_a5, 1, index + 1)
                safe_write_cell(ws_a5, row_a5, 3, project_name)
                safe_write_cell(ws_a5, row_a5, 4, product_name)
                safe_write_cell(ws_a5, row_a5, 5, line.qty_done)
                safe_write_cell(ws_a5, row_a5, 6, line.product_uom_id.name or '')
                safe_write_cell(ws_a5, row_a5, 7, date_done)



            output = BytesIO()
            wb.save(output)
            output.seek(0)
            
            filename = f"Bao_cao_phieu_xuat_tra_kho{self.name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            attachment = self.env['ir.attachment'].create({
                'name': filename,
                'type': 'binary',
                'datas': base64.b64encode(output.getvalue()),
                'res_model': 'stock.picking',
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            })
            
            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{attachment.id}?download=true',
                'target': 'self',
            }
            
        except Exception as e:
            raise UserError(f"Lỗi khi tạo báo cáo: {str(e)}")


class StockMoveLine(models.Model):
    _inherit = 'stock.move.line'

    x_note = fields.Char(string='Ghi chú')
    x_location_stock = fields.Char(string='Vị trí trong kho SPS')
    x_from_the_project_id = fields.Many2one('project.project', string='Từ dự án')
    x_to_the_project_id = fields.Many2one('project.project', string='Tới dự án')
    x_qty_delivered = fields.Float(string='SL đang giữ', readonly='True')
    x_qty_available = fields.Float(string='SL tồn kho', readonly='True', compute='compute_qty_available')
    x_picking_type = fields.Selection(related='picking_id.x_picking_type', string='Loại giao nhận', store=1)
    x_type_code = fields.Selection(related='picking_id.x_type_code', string='Kiểu giao nhận', store=1)
    x_payer_id = fields.Many2one('res.users', string='Người trả/Người bàn giao', related='picking_id.x_payer_id', store=1)
    is_return = fields.Boolean(string='Là phiếu trả hàng', compute='_compute_is_return')

    @api.depends('move_id.origin_returned_move_id', 'picking_id.move_lines.origin_returned_move_id')
    def _compute_is_return(self):
        for rec in self:
            rec.is_return = bool(rec.move_id.origin_returned_move_id) or any(
                m.origin_returned_move_id for m in rec.picking_id.move_lines)
    x_qty_done_t1 = fields.Float(related='qty_done', readonly=False, string='SL nhập kho', digits='Product Unit of Measure')
    x_qty_done_t2 = fields.Float(related='qty_done', readonly=False, string='SL nhập lại', digits='Product Unit of Measure')
    x_qty_done_t3 = fields.Float(related='qty_done', readonly=False, string='SL xuất', digits='Product Unit of Measure')
    x_qty_done_t4 = fields.Float(related='qty_done', readonly=False, string='SL mượn', digits='Product Unit of Measure')
    x_qty_done_t5 = fields.Float(related='qty_done', readonly=False, string='SL chuyển giao', digits='Product Unit of Measure')
    x_qty_done_t6 = fields.Float(related='qty_done', readonly=False, string='SL thu hồi/hoàn', digits='Product Unit of Measure')
    x_qty_done_t7 = fields.Float(related='qty_done', readonly=False, string='SL chuyển', digits='Product Unit of Measure')
    x_qty_done_t8 = fields.Float(related='qty_done', readonly=False, string='SL thu hồi', digits='Product Unit of Measure')
    x_qty_done_t9 = fields.Float(related='qty_done', readonly=False, string='SL báo hỏng/huỷ', digits='Product Unit of Measure')

    _X_QTY_DONE_KEYS = tuple(f'x_qty_done_t{i}' for i in range(1, 10))

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            chosen = None
            for k in self._X_QTY_DONE_KEYS:
                if k in vals and vals.get(k):
                    chosen = vals[k]
                    break
            if chosen is not None:
                vals['qty_done'] = chosen
            for k in self._X_QTY_DONE_KEYS:
                vals.pop(k, None)
        return super().create(vals_list)

    def write(self, vals):
        chosen = None
        for k in self._X_QTY_DONE_KEYS:
            if k in vals and vals.get(k):
                chosen = vals[k]
                break
        if chosen is not None:
            vals['qty_done'] = chosen
        for k in self._X_QTY_DONE_KEYS:
            vals.pop(k, None)
        return super().write(vals)

    def compute_qty_available(self):
        for r in self :
            lot_query = ""
            if not r.product_id:
                continue
            if r.lot_id :
                lot_query = f''' and lot_id = {r.lot_id.id}'''
            sql = f'''select quantity from stock_quant where location_id = {r.location_id.id} and product_id = {r.product_id.id} {lot_query}'''
            self.env.cr.execute(sql)
            data = self.env.cr.fetchall()
            if data:
                r.x_qty_available = data[0][0]
            else:
                r.x_qty_available = 0

    @api.onchange('x_from_the_project_id')
    def onchange_from_project(self):
        if self.picking_id.x_picking_type in ('type_5', 'type_6'):
            for rec in self:
                sql_tra = ''' select sbt.amount from stock_borrow_tools sbt where
                                             (sbt.product_id = {product_id} or sbt.product_id  is NULL) and 
                                             (sbt.lot_id = {lot_id} or sbt.lot_id  is NULL) and 
                                             sbt.receiver_id = {user_id} and
                                             (sbt.project_id = {project_id} or sbt.product_id  is NULL)
                                             and sbt.amount > 0 '''.format(
                    product_id=rec.product_id.id if rec.product_id.id else 'NULL',
                    lot_id=rec.lot_id.id if rec.lot_id.id else 'NULL',
                    user_id=rec.picking_id.x_payer_id.id,
                    project_id=rec.x_from_the_project_id.id if rec.x_from_the_project_id.id else 'NULL')
                self._cr.execute(sql_tra)
                res = self._cr.dictfetchall()
                if res:
                    rec.x_qty_delivered = res[0]['amount']
                else:
                    rec.x_qty_delivered = 0

    @api.onchange('product_id','lot_id')
    def onchange_product_id(self):
        for rec in self:
            if rec.product_id:
                if rec._context.get('picking_type') not in ('type_3', 'type_4'):
                    rec.x_from_the_project_id = False
                    rec.x_to_the_project_id = False
                    rec.qty_done = False
                    print('hehehehehehehe')
                rec.x_location_stock = rec.product_id.x_location_stock
                rec.compute_qty_available()
                
    @api.constrains('qty_done')
    def onchange_qty_done(self):
        if self.picking_id.x_picking_type == 'type_5':
            for r in self:
                if r.qty_done > r.x_qty_delivered:
                    raise UserError('Bạn không thể chuyển số hàng nhiều hơn số hàng đang giữ')

    @api.model
    def default_get(self, fields_list):
        defaults = super(StockMoveLine, self).default_get(fields_list)
        if self._context.get('default_project'):
            project = self.env['project.project'].sudo().browse(self._context.get('default_project'))
            defaults['x_to_the_project_id'] = project.id
        return defaults


class StockMove(models.Model):
    _inherit = 'stock.move'

    x_valuation_layer = fields.Float('Định giá', compute='_compute_valuation_layer')

    @api.depends('stock_valuation_layer_ids', 'stock_valuation_layer_ids.value')
    def _compute_valuation_layer(self):
        for rec in self:
            valuation_layer_ids = self.env['stock.valuation.layer'].search([('stock_move_id', '=', rec.id)])
            rec.x_valuation_layer = sum(abs(x.value) for x in valuation_layer_ids)

    def _get_price_unit(self):
        self.ensure_one()
        if self.picking_id.x_picking_type == 'type_2':
            self.price_unit = self.product_id.standard_price
        return super(StockMove, self)._get_price_unit()


class StockPickingType(models.Model):
    _inherit = 'stock.picking.type'

    x_type = fields.Selection([
        ('type_1', 'Nhập kho từ đơn mua'),
        ('type_2', 'Nhập lại vật tư tiêu chuẩn'),
        ('type_3', 'Xuất kho'),
        ('type_4', 'Mượn công cụ dụng cụ'),
        ('type_5', 'Chuyển giao công cụ dụng cụ'),
        ('type_6', 'Thu hồi/Hoàn trả công cụ dụng cụ'),
        ('type_7', 'Chuyển kho hàng hóa'),
        ('type_8', 'Thu hồi vật tư thừa'),
        ('type_9', 'Báo hỏng, hủy'),
    ], string='Kiểu hoạt động')


class StockProductLot(models.Model):
    _inherit = 'stock.production.lot'

    @api.model
    def _name_search(self, name, args=None, operator='ilike', limit=100, name_get_uid=None):

        if self._context.get('x_project_by_payer', False) and self._context.get('type') in ('type_5', 'type_6'):
            args = args or []
            domain = []
            payer_id = self._context.get('payer_id')
            product_id = self._context.get('product_id')
            if not payer_id:
                raise UserError(_('Chưa chọn người trả'))
            if not product_id:
                raise UserError(_('Chưa chọn sản phẩm'))
            records = self.env['stock.borrow.tools'].sudo().search([
                ('product_id', '=', product_id),
                ('receiver_id', '=', payer_id),
                ('amount', '>', 0),
            ])
            domain = [('id', 'in', [x.lot_id.id for x in records])]

            return self._search(expression.AND([domain, args]), limit=limit, access_rights_uid=name_get_uid)
        elif self._context.get('code_type') in ('outgoing', 'internal') and self._context.get('type') not in ('type_5', 'type_6'):
            args = args or []
            domain = []
            loaction_id = self._context.get('location_id')
            product_id = self._context.get('product_id')
            if not product_id:
                raise UserError(_('Chưa chọn sản phẩm'))
            records = self.env['stock.quant'].sudo().search([
                ('product_id', '=', product_id),
                ('location_id', '=', loaction_id),
                ('quantity', '>', 0),
            ])
            domain = [('id', 'in', [x.lot_id.id for x in records])]
            return self._search(expression.AND([domain, args]), limit=limit, access_rights_uid=name_get_uid)
        return super()._name_search(name, args, operator, limit, name_get_uid)

    @api.model
    def search_read(self, domain=None, fields=None, offset=0, limit=None, order=None):
        if self._context.get('x_project_by_payer', False) and self._context.get('type') in ('type_5', 'type_6'):
            payer_id = self._context.get('payer_id')
            product_id = self._context.get('product_id')
            if not payer_id:
                raise UserError(_('Chưa chọn người trả'))
            if not product_id:
                raise UserError(_('Chưa chọn sản phẩm'))
            records = self.env['stock.borrow.tools'].sudo().search([
                ('product_id', '=', product_id),
                ('receiver_id', '=', payer_id),
                ('amount', '>', 0),
            ])
            domain.append(('id', 'in', [x.lot_id.id for x in records]))
        elif self._context.get('code_type') in ('outgoing', 'internal') and self._context.get('type') not in ('type_5', 'type_6'):
            loaction_id = self._context.get('location_id')
            product_id = self._context.get('product_id')
            if not product_id:
                raise UserError(_('Chưa chọn sản phẩm'))
            records = self.env['stock.quant'].sudo().search([
                ('product_id', '=', product_id),
                ('location_id', '=', loaction_id),
                ('quantity', '>', 0),
            ])
            domain = [('id', 'in', [x.lot_id.id for x in records])]
        return super(StockProductLot, self).search_read(domain, fields, offset, limit, order)
