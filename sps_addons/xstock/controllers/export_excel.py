from odoo.addons.web.controllers.main import ExcelExport,serialize_exception
from odoo.http import content_disposition, dispatch_rpc, request, serialize_exception as _serialize_exception, Response
from odoo import http, tools
import json
import logging
import operator
import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, PatternFill, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.worksheet import filters
from odoo.modules import get_module_path


class Style:
    def __init__(self, border=None, font=None, alignment=None, number_format=None):
        self.border = border
        self.font = font
        self.alignment = alignment
        self.number_format = number_format

class CustomExcelExport(ExcelExport):

    @http.route('/web/export/xlsx', type='http', auth="user")
    @serialize_exception
    def index(self, data, token):
        params = json.loads(data)
        model, fields, ids, domain, import_compat ,context= \
            operator.itemgetter('model', 'fields', 'ids', 'domain', 'import_compat','context')(params)
        if model == "stock.report.value.cost":
            return self.export_stock_report_value_cost(domain)
        if model == "employee.damage.value.report":
            return self.export_employee_damage_value_report(domain,context)
        if model =="project.efficiency.coefficient":
            return request.env['project.efficiency.coefficient'].export_project_efficiency_coefficient(domain,context)
        if model == "project.efficiency.detail":
            return request.env['project.efficiency.detail'].export_project_efficiency_detail(domain, context)
        if model == "project.efficiency.reward.line":
            return request.env['project.efficiency.reward.line'].export_project_efficiency_reward_line(domain, context)
        if model == "stock.supplies.rp":
            return request.env['stock.supplies.rp'].export_excel_report(domain, context)
        if model == "project.sum.value.labor":
            return request.env['project.sum.value.labor'].export_excel_report(domain, context)
        if model == "report.training.monthly.line":
            return request.env['report.training.monthly.line'].export_excel_report(domain, context)
        if model == "report.trainers":
            return request.env['sale.closed.incentive'].export_excel_report(domain, context)
        if model == "sale.closed.incentive":
            return request.env['sale.closed.incentive'].export_excel_report(domain, context)

        return self.base(data, token)

    def safe_write_cell(self,worksheet, row, col, value, style=None):
        """Safely write to a cell, handling merged cells"""
        try:
            cell = worksheet.cell(row=row, column=col)
            # Check if it's a merged cell
            for merged_range in worksheet.merged_cells.ranges:
                if merged_range.min_row <= row <= merged_range.max_row and \
                   merged_range.min_col <= col <= merged_range.max_col:
                    # Write to the top-left cell of the merged range
                    top_left_cell = worksheet.cell(row=merged_range.min_row, column=merged_range.min_col)
                    if hasattr(top_left_cell, 'value'):
                        top_left_cell.value = value
                        if style:
                            if hasattr(style, 'border'):
                                top_left_cell.border = style.border
                            if hasattr(style, 'font'):
                                top_left_cell.font = style.font
                            if hasattr(style, 'alignment'):
                                top_left_cell.alignment = style.alignment
                            if hasattr(style, 'number_format') and style.number_format:
                                top_left_cell.number_format = style.number_format
                    return top_left_cell
            cell.value = value
            if style:
                if hasattr(style, 'border'):
                    cell.border = style.border
                if hasattr(style, 'font'):
                    cell.font = style.font
                if hasattr(style, 'alignment'):
                    cell.alignment = style.alignment
                if hasattr(style, 'number_format') and style.number_format:
                    cell.number_format = style.number_format
            return cell
        except Exception:
            return None  # Skip if can't write

    def export_stock_report_value_cost(self, domain):
        module_path = get_module_path('xstock')
        excel_path = os.path.join(module_path, 'templates', 'Template_uniform_cost.xlsx')
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        ws = wb['Unifrom']
        
        # Define styles
        border = Border(
            left=Side(border_style='thin'),
            right=Side(border_style='thin'),
            top=Side(border_style='thin'),
            bottom=Side(border_style='thin')
        )
        normal_style = Style(border=border)
        bold_style = Style(border=border, font=Font(bold=True))
        center_style = Style(border=border, alignment=Alignment(horizontal='center'))
        # Temporarily remove number_format to test
        bold_number_style = Style(border=border, font=Font(bold=True))
        
        stock_report_value_cost_ids = request.env['stock.report.value.cost'].sudo().search(domain)
        
        start_row = 5
        
        for idx, record in enumerate(stock_report_value_cost_ids):
            row = start_row + idx

            self.safe_write_cell(ws, row, 1, idx + 1, center_style)
            
            cells_data = [
                (2, record.login),
                (3, record.name),
                (4, record.amount_bf)
            ]
            
            for col, value in cells_data:
                self.safe_write_cell(ws, row, col, value, normal_style)
            
            # Dữ liệu tháng
            monthly_values = [
                record.t1, record.t2, record.t3, record.t4, 
                record.t5, record.t6, record.t7, record.t8,
                record.t9, record.t10, record.t11, record.t12
            ]
            
            for month_idx, value in enumerate(monthly_values):
                self.safe_write_cell(ws, row, 5 + month_idx, value, normal_style)
            
            # Cột tổng
            self.safe_write_cell(ws, row, 17, record.total, normal_style)
        
        if stock_report_value_cost_ids:
            summary_row = start_row + len(stock_report_value_cost_ids)
            
            self.safe_write_cell(ws, summary_row, 3, "TỔNG CỘNG", bold_style)
            
            for col in range(1, 3):
                self.safe_write_cell(ws, summary_row, col, "", bold_style)
            
            for col in range(4, 18):
                total_formula = f"=SUM({ws.cell(row=start_row, column=col).coordinate}:{ws.cell(row=summary_row-1, column=col).coordinate})"
                cell = self.safe_write_cell(ws, summary_row, col, total_formula, bold_style)
                # Apply number format directly to cell if it's a number column (not remark)
                if cell and col < 18:  # Not remark column
                    try:
                        cell.number_format = '#,##0'
                    except:
                        pass  # Skip if formatting fails
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"BaoCaoChiPhiNhanVien.xlsx"
        
        # Return response
        response = Response(
            output.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename))
            ]
        )
        
        return response

    def export_employee_damage_value_report(self,domain,context):
        module_path = get_module_path('xstock')
        model_main = context.get('active_model')
        if model_main == 'employee.damage.value.popup':
            excel_path = os.path.join(module_path, 'templates', 'template_damage.xlsx')
            wb = openpyxl.load_workbook(excel_path, data_only=False)
            ws = wb['Damage']
            filename = f"BaoCaoGiaTriTonThat.xlsx"
        else:
            excel_path = os.path.join(module_path, 'templates', 'template_borrowed.xlsx')
            wb = openpyxl.load_workbook(excel_path, data_only=False)
            ws = wb['Borrowed']
            filename = f"BaoCaoCongCuNhanVienQuanLi.xlsx"

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        bold_font = Font(bold=True)

        normal_style = Style(border=thin_border)
        bold_style = Style(border=thin_border, font=bold_font)
        center_style = Style(border=thin_border, alignment=Alignment(horizontal='center'))
        # Number format: format 0 as '-'
        num_format = '#,##0;-#,##0;"-"'
        normal_number_style = Style(border=thin_border, number_format=num_format)
        bold_number_style = Style(border=thin_border, font=bold_font, number_format=num_format)

        damage_report_ids = request.env['employee.damage.value.report'].sudo().search(domain)
        
        start_row = 5

        for idx, record in enumerate(damage_report_ids):
            row = start_row + idx

            self.safe_write_cell(ws, row, 1, record.index, center_style)
            
            self.safe_write_cell(ws, row, 2, record.login, normal_style)
            self.safe_write_cell(ws, row, 3, record.name, normal_style)
            self.safe_write_cell(ws, row, 4, record.amount_bf, normal_number_style)
            
            monthly_values = [
                record.t1, record.t2, record.t3, record.t4, 
                record.t5, record.t6, record.t7, record.t8,
                record.t9, record.t10, record.t11, record.t12
            ]
            
            for month_idx, value in enumerate(monthly_values):
                self.safe_write_cell(ws, row, 5 + month_idx, value, normal_number_style)
            
            self.safe_write_cell(ws, row, 17, record.this_year, normal_number_style)
            
            self.safe_write_cell(ws, row, 18, record.total, normal_number_style)
            
            self.safe_write_cell(ws, row, 19, record.remark or '', normal_style)

        if damage_report_ids:
            summary_row = start_row + len(damage_report_ids)
            self.safe_write_cell(ws, summary_row, 3, "TỔNG CỘNG", bold_style)
            
            for col in range(1, 3):
                self.safe_write_cell(ws, summary_row, col, "", bold_style)
            
            for col in range(4, 19):
                total_formula = f"=SUM({ws.cell(row=start_row, column=col).coordinate}:{ws.cell(row=summary_row-1, column=col).coordinate})"
                cell = self.safe_write_cell(ws, summary_row, col, total_formula, bold_style)
                if cell and col < 19:
                    try:
                        cell.number_format = '#,##0;-#,##0;"-"'
                    except:
                        pass
            
            self.safe_write_cell(ws, summary_row, 19, "", bold_style)

        # Save to BytesIO
        output = BytesIO()
        try:
            wb.save(output)
            output.seek(0)
        except Exception as e:
            _logger = logging.getLogger(__name__)
            _logger.error(f"Error saving workbook in export_employee_damage_value_report: {str(e)}")

        
        # Return response
        response = Response(
            output.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', content_disposition(filename))
            ]
        )
        
        return response
