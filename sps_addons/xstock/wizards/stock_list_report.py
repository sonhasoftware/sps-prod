# -*- coding: utf-8 -*-

from datetime import timedelta, date, datetime
from copy import copy
import base64

import os
from io import BytesIO
import openpyxl
from openpyxl.styles import NamedStyle, Font, Border, Side, Alignment
from openpyxl.writer.excel import save_virtual_workbook
from openpyxl.worksheet import filters
from odoo import api, fields, models
from openpyxl.styles import Font, Border, Alignment


class StockListReport(models.TransientModel):
    _name = 'stock.list.report'
    _description = 'Báo cáo tồn kho vật tư'

    year = fields.Integer('Năm', required=True, default=fields.date.today().year)
    stock = fields.Many2one('stock.location', domain=[('usage', '=', 'internal')])
    is_order = fields.Boolean('Cần đặt hàng' , default=False)

    def apply_highlight_style(self,cell):
        from openpyxl.styles import Side
        bd1 = Side(style='thin', color="000000")
        cell.font = Font(name='Verdana', size=12, bold=True)
        cell.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd1)
        cell.alignment = Alignment(horizontal='right', vertical='center', wrapText=True)
        cell.number_format = '#,##0'  

    def action_report_stock(self):
        dir_path = os.path.dirname(os.path.realpath(__file__))
        wb = openpyxl.load_workbook(dir_path + '%s..%stemplates%sstock_material_report.xlsx' % (os.sep, os.sep, os.sep))
        ws = wb['Sheet1']
        current_year = self.year
        stock = []
        if self.stock:
            stock.append(self.stock.id)
        else:
            rec = self.env['stock.location'].search([('usage', '=', 'internal')])
            for r in rec:
                stock.append(r['id'])

        if self.is_order:
            sql = '''with dt(ma_sp,ma_hang,ten_sp,phan_loai,hang_sx,model,don_vi,don_gia,x_manage_frequency,x_frequency_id,ton_dau_ky,nhap_trong_ky,xuat_trong_ky,nhap_tra_lai,kiem_ke_kho,ton_cuoi_ky,ton_kho_toi_thieu,ton_kho_toi_da,sl_thuc_te,slsd_t1,slsd_t2,slsd_t3,slsd_t4,slsd_t5,slsd_t6,slsd_t7,slsd_t8,slsd_t9,slsd_t10,slsd_t11,slsd_t12,gtsd_t1,gtsd_t2,gtsd_t3,gtsd_t4,gtsd_t5,gtsd_t6,gtsd_t7,gtsd_t8,gtsd_t9,gtsd_t10,gtsd_t11,gtsd_t12,gtsd_cac_nam_truoc,gt_ton_kho) as(
                       SELECT pp.id as ma_sp ,pt.default_code as ma_hang, pt.name as ten_sp,  
                           case when pt.x_supplies_type = 'supplies' then 'Vật tư tiêu chuẩn'
                            when pt.x_supplies_type = 'labor_protection' then 'Bảo hộ lao động'
                            when pt.x_supplies_type = 'stationery' then 'Văn phòng phẩm'
                            when pt.x_supplies_type = 'supplies_project' then 'Vật tư xuất cho dự án' end as phan_loai,
                           pt.x_product_company as hang_sx,
                           pt.x_code_brand as  model,
                           coalesce(it2.value,uu."name") as don_vi,
                           ip.value_float as don_gia,
                           pt.x_manage_frequency as x_manage_frequency,
                           sf.name as x_frequency_id,		 
                                sum(
                                    case
                                    when(sml.date::date < '{year}-01-01' and sl1.usage = 'internal') then -sml.qty_done
                                    when(sml.date::date < '{year}-01-01' and sl2.usage = 'internal') then sml.qty_done
                                    else 0 end ) as ton_dau_ky,		
                               sum(
                                   case 
                                   when(sm.date::date between '{year}-01-01' and '{year}-12-31' and sl2.usage = 'internal' and 
                                     sm.picking_id is not null and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_8')) then sml.qty_done
                                   ELSE 0 end )as nhap_trong_ky,
                               sum(
                                   case 
                                   when(sm.picking_id is not null and sm.date::date between '{year}-01-01' and '{year}-12-31' and sl1.usage = 'internal') then sml.qty_done
                                   ELSE 0 end )as xuat_trong_ky,
                               sum(
                                   case 
                                   when(sm.date::date between '{year}-01-01' and '{year}-12-31' and sl2.usage = 'internal' and 
                                   sm.picking_id is not null and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_8')) then sml.qty_done
                                   ELSE 0 end )as nhap_tra_lai,
                               sum(
                                   case
                                   when(sml.date::date between '{year}-01-01' and '{year}-12-31' and sl1.usage = 'internal' and sm.inventory_id is not null) then -sml.qty_done
                                   when(sml.date::date between '{year}-01-01' and '{year}-12-31' and sl2.usage = 'internal' and sm.inventory_id is not null) then sml.qty_done
                                   else 0 end ) as kiem_ke_kho,
                               sum(
                                   case
                                   when(sml.date::date < '{year}-12-31' and sl1.usage = 'internal') then -sml.qty_done
                                   when(sml.date::date < '{year}-12-31' and sl2.usage = 'internal') then sml.qty_done
                                   else 0 end ) as ton_cuoi_ky,
                               pt.min_inventory as ton_kho_toi_thieu, pt.max_inventory as ton_kho_toi_da,
                               case when sum(pol.product_qty - pol.qty_received) > 0 then sum(pol.product_qty - pol.qty_received) else 0 end as SL_thuc_te,														
                                                                                                                                                                                          sum(case
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t1,
                               sum(case
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t2,
                               sum(case
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t3,
                               sum(case
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t4,
                               sum(case
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t5,
                               sum(case
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t6,
                               sum(case
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t7,
                               sum(case
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t8,
                               sum(case
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t9,
                               sum(case
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t10,
                               sum(case
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t11,
                               sum(case
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t12,
                               sum(case
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t1,
                               sum(case
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t2,
                               sum(case
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t3,
                               sum(case
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t4,
                               sum(case
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t5,
                               sum(case
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t6,
                               sum(case
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t7,
                               sum(case
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t8,
                               sum(case
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t9,
                               sum(case
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t10,
                               sum(case
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t11,
                               sum(case
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t12,
                               sum(
                                   case 
                                   when(sm.date::date < '{year}-01-01' and sl1.usage = 'internal' and 
                                   sm.picking_id is not null ) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   ELSE 0 end
                                  ) as gtsd_cac_nam_truoc,                              
                               sum(
                                   case 
                                   when(sm.date::date <= '{year}-12-31' and sl2.usage = 'internal' and 
                                   sm.picking_id is not null) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   ELSE 0 end) 
                                   - sum (
                                           case 
                                           when(sm.date::date <= '{year}-12-31' and sl1.usage = 'internal' and 
                                   sm.picking_id is not null) then sml.qty_done*coalesce(svl.unit_cost,0)
                                           ELSE 0 end
                                           ) as gt_ton_kho
                   FROM product_template pt
                           LEFT JOIN uom_uom uu  ON uu.id = pt.uom_id
                           LEFT JOIN sale_frequency sf  ON sf.id = pt.x_frequency_id
                           LEFT JOIN product_product pp  ON pt.id = pp.product_tmpl_id
                           LEFT JOIN stock_move sm ON sm.product_id = pp.id
                           left join stock_valuation_layer svl on svl.stock_move_id = sm.id 
                           LEFT JOIN stock_move_line sml ON sml.move_id = sm.id 
                           LEFT JOIN stock_picking sp  ON sm.picking_id = sp.id
                           LEFT JOIN stock_picking_type spt ON sp.picking_type_id = spt.id
                           left join stock_inventory si on sm.inventory_id=si.id
                           left join stock_location sl1 on sm.location_id=sl1.id
                           left join stock_location sl2 on sm.location_dest_id=sl2.id
                           left join ir_translation it on it."name" = 'stock.picking.type,name' and it.lang = 'vi_VN' and it.src = spt."name"
                           left join ir_translation it2 on it2."name" = 'uom.uom,name' and it2.lang = 'vi_VN' and it2.src = uu."name"
                           LEFT JOIN purchase_order_line pol on sm.purchase_line_id = pol.id
                           left join purchase_order po on po.id = pol.order_id 
                           left join ir_property ip on ip.name = 'standard_price' and substring(ip.res_id,17,100)::int = pp.id
                   where sm.state = 'done' 
                   and pt.x_type = 'product'
                   and pt.active = true
                   and pt.x_product_type = 'supplies'
                   and sl1.id != sl2.id 
                   and (sl1.usage = 'internal' or sl2.usage = 'internal') 
                   and (sl1.id in {stock} or sl2.id in {stock})
                   and (sp.state = 'done' or si.state = 'done')
                   and not (sl1.usage = 'internal' and sl2.usage = 'internal')
                   GROUP BY pp.id,pt.default_code,pt.name,pt.x_manage_frequency,sf.name,pt.x_supplies_type, pt.x_product_company,pt.x_code_brand, uu.name,ip.value_float, pt.min_inventory, pt.max_inventory,it2.value)
                SELECT * from dt
                WHERE  (dt.ton_dau_ky+dt.nhap_trong_ky-dt.xuat_trong_ky+dt.nhap_tra_lai+dt.kiem_ke_kho) < dt.ton_kho_toi_thieu
               '''.format(year=current_year, stock=tuple(stock + [0, 0]))

        else:
            sql = '''
                      with chi_tiet as (SELECT pp.id as id_sp, ip.value_float as don_gia,     
                                sum(
                                    case
                                    when(sml.date::date < '{year}-01-01' and sl1.usage = 'internal') then -sml.qty_done
                                    when(sml.date::date < '{year}-01-01' and sl2.usage = 'internal') then sml.qty_done
                                    else 0 end ) as ton_dau_ky,    
                               sum(
                                   case 
                                   when(sm.date::date between '{year}-01-01' and '{year}-12-31' and sl2.usage = 'internal' and 
                                     sm.picking_id is not null and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_8')) then sml.qty_done
                                   ELSE 0 end )as nhap_trong_ky,
                               sum(
                                   case 
                                   when(sm.picking_id is not null and sm.date::date between '{year}-01-01' and '{year}-12-31' and sl1.usage = 'internal') then sml.qty_done
                                   ELSE 0 end )as xuat_trong_ky,
                               sum(
                                   case 
                                   when(sm.date::date between '{year}-01-01' and '{year}-12-31' and sl2.usage = 'internal' and 
                                   sm.picking_id is not null and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_8')) then sml.qty_done
                                   ELSE 0 end )as nhap_tra_lai,
                               sum(
                                   case
                                   when(sml.date::date between '{year}-01-01' and '{year}-12-31' and sl1.usage = 'internal' and sm.inventory_id is not null) then -sml.qty_done
                                   when(sml.date::date between '{year}-01-01' and '{year}-12-31' and sl2.usage = 'internal' and sm.inventory_id is not null) then sml.qty_done
                                   else 0 end ) as kiem_ke_kho,
                               sum(
                                   case
                                   when(sml.date::date < '{year}-12-31' and sl1.usage = 'internal') then -sml.qty_done
                                   when(sml.date::date < '{year}-12-31' and sl2.usage = 'internal') then sml.qty_done
                                   else 0 end ) as ton_cuoi_ky,
                               pt.min_inventory as ton_kho_toi_thieu, pt.max_inventory as ton_kho_toi_da,
                               case when sum(pol.product_qty - pol.qty_received) > 0 then sum(pol.product_qty - pol.qty_received) else 0 end as sl_thuc_te,                            
                                                                                                                                                                                          sum(case
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t1,
                               sum(case
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t2,
                               sum(case
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t3,
                               sum(case
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t4,
                               sum(case
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t5,
                               sum(case
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t6,
                               sum(case
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t7,
                               sum(case
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t8,
                               sum(case
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t9,
                               sum(case
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t10,
                               sum(case
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t11,
                               sum(case
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done
                                   else 0 end) as slsd_t12,
                               sum(case
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 1 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t1,
                               sum(case
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 2 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t2,
                               sum(case
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 3 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t3,
                               sum(case
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 4 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t4,
                               sum(case
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 5 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t5,
                               sum(case
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 6 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t6,
                               sum(case
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 7 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t7,
                               sum(case
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 8 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t8,
                               sum(case
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 9 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t9,
                               sum(case
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 10 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t10,
                               sum(case
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 11 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t11,
                               sum(case
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl1.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id not in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   when (extract(month from sm.date) = 12 and extract(year from sm.date) = {year} and sl2.usage = 'internal'
                                         and sm.picking_id is not null
                                         and sp.picking_type_id in (SELECT spt.id FROM stock_picking_type spt WHERE x_type = 'type_2')) then -sml.qty_done*coalesce(svl.unit_cost,0)
                                   else 0 end) as gtsd_t12,
                               sum(
                                   case 
                                   when(sm.date::date < '{year}-01-01' and sl1.usage = 'internal' and 
                                   sm.picking_id is not null ) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   ELSE 0 end
                                  ) as gtsd_cac_nam_truoc,                              
                               sum(
                                   case 
                                   when(sm.date::date <= '{year}-12-31' and sl2.usage = 'internal' and 
                                   sm.picking_id is not null) then sml.qty_done*coalesce(svl.unit_cost,0)
                                   ELSE 0 end) 
                                   - sum (
                                           case 
                                           when(sm.date::date <= '{year}-12-31' and sl1.usage = 'internal' and 
                                   sm.picking_id is not null) then sml.qty_done*coalesce(svl.unit_cost,0)
                                           ELSE 0 end
                                           ) as gt_ton_kho
                               FROM product_template pt
                               LEFT JOIN uom_uom uu  ON uu.id = pt.uom_id
                               LEFT JOIN product_product pp  ON pt.id = pp.product_tmpl_id
                               LEFT JOIN stock_move sm ON sm.product_id = pp.id
                               left join stock_valuation_layer svl on svl.stock_move_id = sm.id 
                               LEFT JOIN stock_move_line sml ON sml.move_id = sm.id 
                               LEFT JOIN stock_picking sp  ON sm.picking_id = sp.id
                               LEFT JOIN stock_picking_type spt ON sp.picking_type_id = spt.id
                               left join stock_inventory si on sm.inventory_id=si.id
                               left join stock_location sl1 on sm.location_id=sl1.id
                               left join stock_location sl2 on sm.location_dest_id=sl2.id
                               left join ir_translation it on it."name" = 'stock.picking.type,name' and it.lang = 'vi_VN' and it.src = spt."name"
                               LEFT JOIN purchase_order_line pol on sm.purchase_line_id = pol.id
                               left join purchase_order po on po.id = pol.order_id 
                               left join ir_property ip on ip.name = 'standard_price' and substring(ip.res_id,17,100)::int = pp.id
                               where sm.state = 'done'
                               and pt.active = true
                               and pt.x_type = 'product'
                               and pt.x_product_type = 'supplies'
                               and sl1.id != sl2.id 
                               and (sl1.usage = 'internal' or sl2.usage = 'internal') 
                               and (sl1.id in (select id from stock_location where usage = 'internal') or sl2.id in (select id from stock_location where usage = 'internal'))
                               and (sp.state = 'done' or si.state = 'done')
                               and not (sl1.usage = 'internal' and sl2.usage = 'internal')
                               GROUP BY pp.id, ip.value_float, pt.min_inventory, pt.max_inventory)
                select 
                    a.default_code as ma_hang, 
                    b.name as ten_sp, 
                    case 
                        when b.x_supplies_type = 'supplies' then 'Vật tư tiêu chuẩn'
                        when b.x_supplies_type = 'labor_protection' then 'Bảo hộ lao động'
                        when b.x_supplies_type = 'stationery' then 'Văn phòng phẩm'
                        when b.x_supplies_type = 'supplies_project' then 'Vật tư xuất cho dự án' 
                    end as phan_loai,
                    b.x_product_company as hang_sx,
                    b.x_manage_frequency as x_manage_frequency, 
                    sf.name as x_frequency_id,
                    b.x_code_brand as model,
                    coalesce(d.value, c.name) as don_vi,
                    coalesce(e.don_gia, 0) as don_gia,
                    coalesce(e.ton_dau_ky, 0) as ton_dau_ky,
                    coalesce(e.nhap_trong_ky, 0) as nhap_trong_ky,
                    coalesce(e.xuat_trong_ky, 0) as xuat_trong_ky,
                    coalesce(e.nhap_tra_lai, 0) as nhap_tra_lai,
                    coalesce(e.kiem_ke_kho, 0) as kiem_ke_kho,
                    coalesce(e.ton_cuoi_ky, 0) as ton_cuoi_ky,
                    coalesce(e.ton_kho_toi_thieu, 0) as ton_kho_toi_thieu,
                    coalesce(e.ton_kho_toi_da, 0) as ton_kho_toi_da,
                    coalesce(e.sl_thuc_te, 0) as sl_thuc_te,
                    coalesce(e.slsd_t1, 0) as slsd_t1,
                    coalesce(e.slsd_t2, 0) as slsd_t2,
                    coalesce(e.slsd_t3, 0) as slsd_t3,
                    coalesce(e.slsd_t4, 0) as slsd_t4,
                    coalesce(e.slsd_t5, 0) as slsd_t5,
                    coalesce(e.slsd_t6, 0) as slsd_t6,
                    coalesce(e.slsd_t7, 0) as slsd_t7,
                    coalesce(e.slsd_t8, 0) as slsd_t8,
                    coalesce(e.slsd_t9, 0) as slsd_t9,
                    coalesce(e.slsd_t10, 0) as slsd_t10,
                    coalesce(e.slsd_t11, 0) as slsd_t11,
                    coalesce(e.slsd_t12, 0) as slsd_t12,
                    coalesce(e.gtsd_t1, 0) as gtsd_t1,
                    coalesce(e.gtsd_t2, 0) as gtsd_t2,
                    coalesce(e.gtsd_t3, 0) as gtsd_t3,
                    coalesce(e.gtsd_t4, 0) as gtsd_t4,
                    coalesce(e.gtsd_t5, 0) as gtsd_t5,
                    coalesce(e.gtsd_t6, 0) as gtsd_t6,
                    coalesce(e.gtsd_t7, 0) as gtsd_t7,
                    coalesce(e.gtsd_t8, 0) as gtsd_t8,
                    coalesce(e.gtsd_t9, 0) as gtsd_t9,
                    coalesce(e.gtsd_t10, 0) as gtsd_t10,
                    coalesce(e.gtsd_t11, 0) as gtsd_t11,
                    coalesce(e.gtsd_t12, 0) as gtsd_t12,
                    coalesce(e.gtsd_cac_nam_truoc, 0) as gtsd_cac_nam_truoc,
                    coalesce(e.gt_ton_kho, 0) as gt_ton_kho
                from product_product a
                join product_template b on a.product_tmpl_id = b.id
                LEFT JOIN uom_uom c ON c.id = b.uom_id
                LEFT JOIN sale_frequency sf ON sf.id = b.x_frequency_id
                left join ir_translation d on d.name = 'uom.uom,name' and d.lang = 'vi_VN' and d.src = c.name
                left join chi_tiet e on e.id_sp = a.id 
                where b.active = true
                    and b.x_type = 'product'
                    and b.x_product_type = 'supplies';
                '''.format(year=current_year)

        self._cr.execute(sql)
        recs = self._cr.dictfetchall()


        highlight = NamedStyle(name="highlight")
        bd1 = Side(style='thin', color="000000")
        bd2 = Side(style='dotted', color="000000")
        bd3 = Side(style='none')
        highlight.border = Border(left=bd1, top=bd2, right=bd1, bottom=bd2)

        style_sum1 = NamedStyle(name="style_sum1")
        style_sum1.PatternFill = Font(color='0099CC00')
        style_sum1.number_format = '#,##0'
        style_sum1.border = Border(left=bd1, top=bd1, right=bd1, bottom=bd2)


        # Khởi tạo các biến tổng cho từng cột (cột 7-47)
        totals = {}
        for col in range(7, 48):
            totals[col] = 0
        
        row = 7
        x = 1

        # Chụp format của dòng 7 (dòng dữ liệu mẫu trong template) theo TỪNG cột (1-47). Template
        # chỉ định dạng sẵn vài dòng đầu, nên các dòng dữ liệu phía sau bị trơ. Copy style dòng 7
        # sang mọi dòng ghi để định dạng đồng nhất (border, số #,##0, font, canh lề... như dòng 7).
        ref_styles = {c: copy(ws.cell(7, c)._style) for c in range(1, 48)}

        for r in recs:
            for c in range(1, 48):
                ws.cell(row, c)._style = copy(ref_styles[c])
            ws.cell(row, 1).value = x
            ws.cell(row, 2).value = r['ma_hang']
            # ws.cell(row, 2).value, ws.cell(row, 2).style = r['ma_hang'],highlight1
            ws.cell(row, 3).value = r['ten_sp']
            ws.cell(row, 4).value = r['phan_loai']
            ws.cell(row, 5).value = r['hang_sx']
            ws.cell(row, 6).value = r['model']
            ws.cell(row, 7).value = r['don_vi']
            ws.cell(row, 8).value = r['don_gia']
            ws.cell(row, 9).value = r['ton_dau_ky']
            ws.cell(row, 10).value = r['nhap_trong_ky']
            ws.cell(row, 11).value = r['xuat_trong_ky']
            ws.cell(row, 12).value = r['nhap_tra_lai']
            ws.cell(row, 13).value = r['kiem_ke_kho']
            ws.cell(row, 14).value = '=(I%s+J%s-K%s+L%s+M%s)' % (row, row, row, row, row)
            ws.cell(row, 15).value = r['ton_kho_toi_thieu']
            ws.cell(row, 16).value = r['ton_kho_toi_da']
            if r['x_manage_frequency']:
                ws.cell(row, 17).value = r['x_frequency_id']
            else:
                ws.cell(row, 17).value = ''
            ws.cell(row, 18).value = '=IF(N%s>=O%s,0,P%s-O%s)' % (row, row, row, row)
            ws.cell(row, 19).value = r['sl_thuc_te']
            ws.cell(row, 20).value = r['slsd_t1']
            ws.cell(row, 21).value = r['slsd_t2']
            ws.cell(row, 22).value = r['slsd_t3']
            ws.cell(row, 23).value = r['slsd_t4']
            ws.cell(row, 24).value = r['slsd_t5']
            ws.cell(row, 25).value = r['slsd_t6']
            ws.cell(row, 26).value = r['slsd_t7']
            ws.cell(row, 27).value = r['slsd_t8']
            ws.cell(row, 28).value = r['slsd_t9']
            ws.cell(row, 29).value = r['slsd_t10']
            ws.cell(row, 30).value = r['slsd_t11']
            ws.cell(row, 31).value = r['slsd_t12']
            ws.cell(row, 32).value = r['gtsd_t1']
            ws.cell(row, 33).value = r['gtsd_t2']
            ws.cell(row, 34).value = r['gtsd_t3']
            ws.cell(row, 35).value = r['gtsd_t4']
            ws.cell(row, 36).value = r['gtsd_t5']
            ws.cell(row, 37).value = r['gtsd_t6']
            ws.cell(row, 38).value = r['gtsd_t7']
            ws.cell(row, 39).value = r['gtsd_t8']
            ws.cell(row, 40).value = r['gtsd_t9']
            ws.cell(row, 41).value = r['gtsd_t10']
            ws.cell(row, 42).value = r['gtsd_t11']
            ws.cell(row, 43).value = r['gtsd_t12']
            ws.cell(row, 44).value = r['gtsd_cac_nam_truoc']
            ws.cell(row, 45).value = '=SUM(AF%s:AQ%s)' % (row, row)
            # ws.cell(row, 46).value = r['gt_ton_kho']
            ws.cell(row, 46).value = '=(H%s*N%s)' % (row, row)
            ws.cell(row, 47).value = '=(AR%s+AS%s)' % (row, row)


            # Cộng dồn các giá trị vào tổng (chỉ cộng các cột có giá trị số)
            totals[8] += r['don_gia'] or 0  # Đơn giá
            totals[9] += r['ton_dau_ky'] or 0  # Tồn đầu kỳ
            totals[10] += r['nhap_trong_ky'] or 0  # Nhập trong kỳ
            totals[11] += r['xuat_trong_ky'] or 0  # Xuất trong kỳ
            totals[12] += r['nhap_tra_lai'] or 0  # Nhập trả lại
            totals[13] += r['kiem_ke_kho'] or 0  # Kiểm kê kho
            totals[15] += r['ton_kho_toi_thieu'] or 0  # Tồn kho tối thiểu
            totals[16] += r['ton_kho_toi_da'] or 0  # Tồn kho tối đa
            totals[19] += r['sl_thuc_te'] or 0  # SL thực tế
            
            # Cộng dồn các cột slsd_t1 đến slsd_t12 (cột 20-31)
            totals[20] += r['slsd_t1'] or 0
            totals[21] += r['slsd_t2'] or 0
            totals[22] += r['slsd_t3'] or 0
            totals[23] += r['slsd_t4'] or 0
            totals[24] += r['slsd_t5'] or 0
            totals[25] += r['slsd_t6'] or 0
            totals[26] += r['slsd_t7'] or 0
            totals[27] += r['slsd_t8'] or 0
            totals[28] += r['slsd_t9'] or 0
            totals[29] += r['slsd_t10'] or 0
            totals[30] += r['slsd_t11'] or 0
            totals[31] += r['slsd_t12'] or 0
            
            # Cộng dồn các cột gtsd_t1 đến gtsd_t12 (cột 32-43)
            totals[32] += r['gtsd_t1'] or 0
            totals[33] += r['gtsd_t2'] or 0
            totals[34] += r['gtsd_t3'] or 0
            totals[35] += r['gtsd_t4'] or 0
            totals[36] += r['gtsd_t5'] or 0
            totals[37] += r['gtsd_t6'] or 0
            totals[38] += r['gtsd_t7'] or 0
            totals[39] += r['gtsd_t8'] or 0
            totals[40] += r['gtsd_t9'] or 0
            totals[41] += r['gtsd_t10'] or 0
            totals[42] += r['gtsd_t11'] or 0
            totals[43] += r['gtsd_t12'] or 0
            
            totals[44] += r['gtsd_cac_nam_truoc'] or 0  # GTSD các năm trước

            x += 1
            row += 1

        # Ghi các giá trị tổng vào hàng 6 từ cột 7
        total_row = 6
        last_data_row = row - 1  # Hàng cuối cùng có dữ liệu
        
        
        for col in [8, 9, 10, 11, 12, 13, 15, 16, 19]:
            cell = ws.cell(total_row, col)
            cell.value = totals[col]
            self.apply_highlight_style(cell)
        
        # Các cột công thức tính
        cell = ws.cell(total_row, 14)
        cell.value = f'=SUM(N7:N{last_data_row})'  # Tồn cuối kỳ
        self.apply_highlight_style(cell)
        
        cell = ws.cell(total_row, 18)
        cell.value = f'=SUM(R7:R{last_data_row})'  # Cần đặt hàng
        self.apply_highlight_style(cell)
        
        for col in range(20, 32):
            cell = ws.cell(total_row, col)
            cell.value = totals[col]
            self.apply_highlight_style(cell)
        
        for col in range(32, 44):
            cell = ws.cell(total_row, col)
            cell.value = totals[col]
            self.apply_highlight_style(cell)
        
        # GTSD các năm trước
        cell = ws.cell(total_row, 44)
        cell.value = totals[44]
        self.apply_highlight_style(cell)
        
        cell = ws.cell(total_row, 45)
        cell.value = f'=SUM(AS7:AS{last_data_row})'
        self.apply_highlight_style(cell)
        
        cell = ws.cell(total_row, 46)
        cell.value = f'=SUM(AT7:AT{last_data_row})'
        self.apply_highlight_style(cell)
        
        cell = ws.cell(total_row, 47)
        cell.value = f'=SUM(AU7:AU{last_data_row})'
        self.apply_highlight_style(cell)

        stream = BytesIO(save_virtual_workbook(wb))
        xls = stream.getvalue()

        attachment_id = self.env['ir.attachment'].create({
            'name': 'Báo cáo tồn kho vật tư.xlsx',
            'datas': base64.b64encode(xls),
            'type': 'binary',
        })
        # download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/' + str(attachment_id.id) + '?download=true',
            'target': 'new',
        }
