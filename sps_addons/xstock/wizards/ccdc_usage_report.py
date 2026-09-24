# -*- coding: utf-8 -*-
import base64
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import save_virtual_workbook

from odoo import _, api, fields, models
from odoo.exceptions import UserError


class CcdcUsageReport(models.TransientModel):
    _name = 'ccdc.usage.report'
    _description = 'Báo cáo giờ sử dụng CCDC'

    so_nam = fields.Integer(
        'Số năm (để 0 = tự lấy theo dữ liệu)',
        default=0,
        help='Số năm dùng để chia ra giờ/năm. Để 0 thì lấy số năm trải dữ liệu '
             '(năm đầu → năm cuối có lịch sử mượn) của từng sản phẩm.')

    import_file = fields.Binary('File Excel cập nhật')
    import_filename = fields.Char('Tên file')

    # Lịch sử mượn dựng lại từ stock_move_line bằng gaps-and-islands.
    # Số lượng sở hữu lấy từ tồn kho mọi location nội bộ (CCDC cho mượn vẫn nằm
    # trong kho nội bộ nên đây là số cái thật).
    _SQL = """
        WITH ccdc AS (
            -- Khớp 1:1 với ir.filter "Công cụ dụng cụ": type in (consu, product)
            -- + x_product_type = 'tools'. Raw SQL không lọc active nên gồm cả
            -- CCDC đã lưu trữ (active = false).
            SELECT pp.id AS product_id FROM product_product pp
            JOIN product_template pt ON pt.id = pp.product_tmpl_id
            WHERE pt.type IN ('consu', 'product') AND pt.x_product_type = 'tools'
        ),
        owned AS (
            SELECT sq.product_id, SUM(sq.quantity) AS so_luong_that
            FROM stock_quant sq
            JOIN stock_location sl ON sl.id = sq.location_id AND sl.usage = 'internal'
            JOIN ccdc USING (product_id)
            GROUP BY sq.product_id
        ),
        events AS (
            SELECT sml.product_id, sml.lot_id, sml.x_to_the_project_id AS project_id,
                   (sml.date + interval '7 hour') AS dt, sml.qty_done AS q
            FROM stock_move_line sml JOIN ccdc USING (product_id)
            WHERE sml.state = 'done' AND sml.x_picking_type IN ('type_4', 'type_5')
              AND sml.x_to_the_project_id IS NOT NULL AND sml.qty_done > 0
            UNION ALL
            SELECT sml.product_id, sml.lot_id, sml.x_from_the_project_id,
                   (sml.date + interval '7 hour'), -sml.qty_done
            FROM stock_move_line sml JOIN ccdc USING (product_id)
            WHERE sml.state = 'done' AND sml.x_picking_type IN ('type_5', 'type_6')
              AND sml.x_from_the_project_id IS NOT NULL AND sml.qty_done > 0
        ),
        tl AS (
            SELECT e.*, SUM(q) OVER w AS qty_held FROM events e
            WINDOW w AS (PARTITION BY product_id, COALESCE(lot_id, 0), project_id
                         ORDER BY dt, q DESC ROWS UNBOUNDED PRECEDING)
        ),
        flg AS (
            SELECT *, CASE WHEN (qty_held - q) <= 0 AND qty_held > 0 THEN 1 ELSE 0 END AS is_open
            FROM tl
        ),
        grp AS (
            SELECT *, SUM(is_open) OVER (PARTITION BY product_id, COALESCE(lot_id, 0), project_id
                         ORDER BY dt, q DESC ROWS UNBOUNDED PRECEDING) AS g
            FROM flg
        ),
        isl AS (
            SELECT product_id, MAX(qty_held) AS sl, MIN(dt) AS seg_start,
                   COALESCE(MIN(dt) FILTER (WHERE qty_held <= 0),
                            (now() AT TIME ZONE 'UTC') + interval '7 hour') AS seg_end
            FROM grp GROUP BY product_id, lot_id, project_id, g
        ),
        hours AS (
            SELECT product_id,
                   SUM(sl * ((seg_end::date - seg_start::date + 1) * 8)) AS tong_gio_8h,
                   SUM(sl * GREATEST(extract(epoch FROM (seg_end - seg_start)) / 3600.0, 0)) AS tong_gio_thuc,
                   MIN(seg_start)::date AS lan_dau, MAX(seg_end)::date AS lan_cuoi
            FROM isl GROUP BY product_id
        ),
        -- Ước lượng số cái đã đưa vào lưu thông, dùng làm mẫu số "giờ/cái"
        -- khi tồn kho hiện tại = 0 (CCDC đã thanh lý/lưu trữ).
        -- peak_total: số cái nhiều nhất cùng mượn đồng thời trên TẤT CẢ dự án.
        peak AS (
            SELECT product_id, MAX(run) AS peak_total FROM (
                SELECT product_id,
                       SUM(q) OVER (PARTITION BY product_id ORDER BY dt
                                    ROWS UNBOUNDED PRECEDING) AS run
                FROM events
            ) t GROUP BY product_id
        ),
        -- so_lot: số serial/lot khác nhau từng được cấp (chính xác với hàng định danh).
        lots AS (
            SELECT sml.product_id, COUNT(DISTINCT sml.lot_id) AS so_lot
            FROM stock_move_line sml JOIN ccdc USING (product_id)
            WHERE sml.state = 'done' AND sml.x_picking_type IN ('type_4', 'type_5')
              AND sml.x_to_the_project_id IS NOT NULL AND sml.qty_done > 0
              AND sml.lot_id IS NOT NULL
            GROUP BY sml.product_id
        )
        SELECT
            pt.default_code AS ma_sp, pt.name AS ten_sp,
            (NOT pt.active) AS da_luu_tru,
            COALESCE(o.so_luong_that, 0) AS so_luong_that,
            GREATEST(COALESCE(lo.so_lot, 0), COALESCE(pk.peak_total, 0)) AS so_cai_uoc_luong,
            h.tong_gio_8h, h.tong_gio_thuc,
            (extract(year FROM h.lan_cuoi) - extract(year FROM h.lan_dau) + 1)::int AS so_nam_dulieu,
            pt.x_hours_per_year AS gio_nam_cauhinh,
            pt.x_depreciation_years AS so_nam_kh
        FROM ccdc c
        JOIN product_product pp ON pp.id = c.product_id
        JOIN product_template pt ON pt.id = pp.product_tmpl_id
        LEFT JOIN hours h USING (product_id)
        LEFT JOIN owned o USING (product_id)
        LEFT JOIN peak pk USING (product_id)
        LEFT JOIN lots lo USING (product_id)
        ORDER BY pt.active DESC, pt.default_code
    """

    def action_export_xlsx(self):
        self.ensure_one()
        self._cr.execute(self._SQL)
        recs = self._cr.dictfetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'CCDC'

        headers = [
            ('Mã SP', 14),
            ('Tên sản phẩm', 45),
            ('Đã lưu trữ', 11),
            ('SL sở hữu', 11),
            ('SL tính giờ/cái', 13),
            ('Tổng giờ mượn (8h)', 16),
            ('Tổng giờ mượn (thực)', 16),
            ('Số năm dữ liệu', 12),
            ('Số năm tính', 11),
            ('Giờ/cái (8h)', 13),
            ('Giờ/cái (thực)', 13),
            ('Giờ/năm 1 cái (8h)', 16),
            ('Giờ/năm 1 cái (thực)', 16),
            ('Giờ/năm cấu hình', 14),
            ('Số năm KH', 11),
        ]

        thin = Side(style='thin', color='000000')
        border = Border(left=thin, top=thin, right=thin, bottom=thin)
        head_font = Font(bold=True, color='FFFFFF', size=10)
        head_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
        head_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        num_fmt = '#,##0'

        for col, (title, width) in enumerate(headers, start=1):
            c = ws.cell(1, col, title)
            c.font = head_font
            c.fill = head_fill
            c.alignment = head_align
            c.border = border
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = 'A2'

        row = 2
        for r in recs:
            sl = r['so_luong_that'] or 0
            # Mẫu số "giờ/cái": ưu tiên tồn kho hiện tại; nếu = 0 (đã thanh lý/
            # lưu trữ) thì fallback sang số cái ước lượng đã đưa vào lưu thông.
            divisor = sl if sl else (r['so_cai_uoc_luong'] or 0)
            g8 = r['tong_gio_8h'] or 0.0
            gt = r['tong_gio_thuc'] or 0.0
            # Số năm chia: ưu tiên ô nhập tay, fallback số năm trải dữ liệu.
            nam = self.so_nam if self.so_nam and self.so_nam > 0 else (r['so_nam_dulieu'] or 0)
            gio_cai_8h = g8 / divisor if divisor else 0.0
            gio_cai_thuc = gt / divisor if divisor else 0.0
            gio_nam_8h = gio_cai_8h / nam if nam else 0.0
            gio_nam_thuc = gio_cai_thuc / nam if nam else 0.0

            values = [
                r['ma_sp'] or '', r['ten_sp'] or '',
                'Đã lưu trữ' if r['da_luu_tru'] else '', sl, divisor,
                round(g8), round(gt), r['so_nam_dulieu'] or 0, nam,
                round(gio_cai_8h), round(gio_cai_thuc),
                round(gio_nam_8h), round(gio_nam_thuc),
                r['gio_nam_cauhinh'] or 0, r['so_nam_kh'] or 0,
            ]
            for col, val in enumerate(values, start=1):
                c = ws.cell(row, col, val)
                c.border = border
                if col >= 4:
                    c.number_format = num_fmt
            row += 1

        stream = BytesIO(save_virtual_workbook(wb))
        attachment = self.env['ir.attachment'].create({
            'name': 'Bao_cao_gio_su_dung_CCDC.xlsx',
            'datas': base64.b64encode(stream.getvalue()),
            'type': 'binary',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    # Cập nhật cấu hình CCDC từ file Excel đã chỉnh sửa.
    # Khớp sản phẩm theo cột "Mã SP" (default_code), ghi:
    #   "Giờ/năm cấu hình" -> x_hours_per_year
    #   "Số năm KH"        -> x_depreciation_years
    # Chỉ ghi vào CCDC (type in (consu, product), x_product_type='tools').
    def action_import_xlsx(self):
        self.ensure_one()
        if not self.import_file:
            raise UserError(_('Vui lòng tải lên file Excel trước khi cập nhật.'))
        try:
            wb = openpyxl.load_workbook(BytesIO(base64.b64decode(self.import_file)), data_only=True)
        except Exception as e:
            raise UserError(_('Không đọc được file Excel: %s') % e)
        ws = wb.active

        # Dò vị trí cột theo tiêu đề ở dòng 1.
        header = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(1, col).value
            if val is not None:
                header[str(val).strip()] = col
        col_ma = header.get('Mã SP')
        col_gio = header.get('Giờ/năm cấu hình')
        col_nam_kh = header.get('Số năm KH')
        if not col_ma or (not col_gio and not col_nam_kh):
            raise UserError(_(
                'File không đúng định dạng. Cần cột "Mã SP" và ít nhất một trong '
                '"Giờ/năm cấu hình" / "Số năm KH". Hãy xuất file từ nút "Xuất Excel" '
                'rồi chỉnh sửa trên đó.'))

        # active_test=False để cập nhật được cả CCDC đã lưu trữ.
        Template = self.env['product.template'].with_context(active_test=False)
        updated, skipped = 0, []
        for row in range(2, ws.max_row + 1):
            ma = ws.cell(row, col_ma).value
            if ma is None or str(ma).strip() == '':
                continue
            ma = str(ma).strip()
            tmpls = Template.search([
                ('default_code', '=', ma),
                ('type', 'in', ['consu', 'product']),
                ('x_product_type', '=', 'tools'),
            ])
            if not tmpls:
                skipped.append(ma)
                continue
            vals = {}
            if col_gio:
                v = ws.cell(row, col_gio).value
                if isinstance(v, (int, float)):
                    vals['x_hours_per_year'] = float(v)
            if col_nam_kh:
                v = ws.cell(row, col_nam_kh).value
                if isinstance(v, (int, float)):
                    vals['x_depreciation_years'] = float(v)
            if vals:
                tmpls.write(vals)
                updated += len(tmpls)

        msg = _('Đã cập nhật %s sản phẩm.') % updated
        if skipped:
            msg += _(' Bỏ qua %s mã không tìm thấy: %s') % (
                len(skipped), ', '.join(skipped[:20]) + ('...' if len(skipped) > 20 else ''))
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Cập nhật CCDC'),
                'message': msg,
                'type': 'success' if updated else 'warning',
                'sticky': bool(skipped),
            },
        }
